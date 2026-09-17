"""Tests for Table metadata: t.info, table_info(), and the created_at/
modified_at/sheet lifecycle."""

from __future__ import annotations

from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    ColumnType,
    SchemaRebuiltWarning,
    Table,
    TableInfo,
    TableNotFoundError,
    create_sheet,
    table_info,
)
from pyhandlexl._multi_table import SCHEMA_SHEET


@pytest.fixture
def data_sheet(book):
    create_sheet(book, "Data")
    return book


def _delete_schema_sheet(path) -> None:
    wb = load_workbook(path)
    del wb[SCHEMA_SHEET]
    wb.save(path)


class TestFreshTable:
    def test_not_yet_placed_has_no_sheet_or_dates(self):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T")
        info = t.info
        assert info.sheet is None
        assert info.created_at is None
        assert info.modified_at is None

    def test_size_reflects_in_memory_shape(self):
        t = Table(data=[[1, 2], [3, 4]], column_headers=["a", "b"], row_labels=["x", "y"], name="T")
        info = t.info
        assert info.n_rows == 2
        assert info.n_cols == 2

    def test_style_and_column_types_are_cross_referenced(self):
        t = Table(
            data=[[1]],
            column_headers=["a"],
            row_labels=["r"],
            name="T",
            column_types={"a": ColumnType.NUMBER},
        )
        assert t.info.style == t.style
        assert t.info.column_types == t.column_types


class TestCreate:
    def test_sets_sheet_and_both_dates(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T")
        before = datetime.now(timezone.utc)
        t.create(data_sheet, sheet="Data")
        after = datetime.now(timezone.utc)

        info = t.info
        assert info.sheet == "Data"
        assert before <= info.created_at <= after
        assert info.created_at == info.modified_at

    def test_size_matches_placed_table(self, data_sheet):
        Table(
            data=[[1, 2], [3, 4], [5, 6]],
            column_headers=["a", "b"],
            row_labels=["x", "y", "z"],
            name="T",
        ).create(data_sheet, sheet="Data")
        info = Table.read(data_sheet, "T").info
        assert info.n_rows == 3
        assert info.n_cols == 2


class TestWrite:
    def test_modified_at_advances_but_created_at_does_not(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T")
        t.create(data_sheet, sheet="Data")
        created = t.info.created_at

        t.add_row("r2", [2])
        t.write(data_sheet)

        assert t.info.created_at == created
        assert t.info.modified_at >= created

    def test_read_then_write_keeps_the_original_created_at(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T").create(
            data_sheet, sheet="Data"
        )
        original_created = Table.read(data_sheet, "T").info.created_at

        t = Table.read(data_sheet, "T")
        t.add_row("r2", [2])
        t.write(data_sheet)

        assert Table.read(data_sheet, "T").info.created_at == original_created

    def test_write_without_prior_read_still_syncs_created_at_from_disk(self, data_sheet):
        # a Table object that never went through .read() has no in-memory
        # created_at of its own — write() must pull it from the schema
        # rather than leaving/writing None.
        Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T").create(
            data_sheet, sheet="Data"
        )
        real_created = table_info(data_sheet, "T").created_at

        fresh = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T")
        assert fresh.info.created_at is None  # never created/read itself
        fresh.write(data_sheet)

        assert table_info(data_sheet, "T").created_at == real_created

    def test_size_updates_on_write(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T")
        t.create(data_sheet, sheet="Data")
        t.add_column("b", [99])
        t.write(data_sheet)
        assert t.info.n_cols == 2
        assert Table.read(data_sheet, "T").info.n_cols == 2


class TestShiftDoesNotCountAsModified:
    def test_shifted_neighbour_keeps_its_own_modified_at(self, data_sheet):
        sales = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="Sales")
        sales.create(data_sheet, sheet="Data")
        inventory = Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="Inventory")
        inventory.create(data_sheet, sheet="Data")

        inventory_modified_before = table_info(data_sheet, "Inventory").modified_at

        sales.add_column("b", [2])
        sales.write(data_sheet)  # grows Sales, shifts Inventory right

        assert table_info(data_sheet, "Inventory").modified_at == inventory_modified_before
        assert table_info(data_sheet, "Inventory").sheet == "Data"  # confirms it did shift/exist


class TestTableInfoFunction:
    def test_matches_table_read_info(self, data_sheet):
        Table(
            data=[[1, 2]],
            column_headers=["a", "b"],
            row_labels=["r"],
            name="T",
            column_types={"a": ColumnType.NUMBER},
        ).create(data_sheet, sheet="Data")
        assert table_info(data_sheet, "T") == Table.read(data_sheet, "T").info

    def test_unknown_name_raises(self, data_sheet):
        with pytest.raises(TableNotFoundError):
            table_info(data_sheet, "Ghost")

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            table_info(tmp_path / "nope.xlsx", "T")

    def test_returns_a_table_info_instance(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T").create(
            data_sheet, sheet="Data"
        )
        assert isinstance(table_info(data_sheet, "T"), TableInfo)


class TestRebuildLosesDates:
    def test_rebuilt_table_reports_none_for_both_dates(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["r"], name="T").create(
            data_sheet, sheet="Data"
        )
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            info = table_info(data_sheet, "T")
        assert info.created_at is None
        assert info.modified_at is None

    def test_size_and_sheet_are_still_exact_after_rebuild(self, data_sheet):
        Table(
            data=[[1, 2], [3, 4]], column_headers=["a", "b"], row_labels=["x", "y"], name="T"
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            info = table_info(data_sheet, "T")
        assert info.n_rows == 2
        assert info.n_cols == 2
        assert info.sheet == "Data"
