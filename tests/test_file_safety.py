"""What the library does to files: never destroy one, never quietly drop what's in it."""

from __future__ import annotations

import os
import stat
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from pyhandlexl import (
    CellTypeError,
    FileLockedError,
    FileReadOnlyError,
    InvalidFileError,
    SchemaRebuiltWarning,
    Table,
    append_rows,
    clear_all_sheet_data,
    create_sheet,
    create_workbook,
    delete_sheet,
    delete_table,
    delete_workbook,
    export_xl_to_csv,
    is_valid_xlsx,
    list_sheets,
    list_tables,
    read_sheet,
    rename_sheet,
    sheet_exists,
    sheet_kind,
    table_info,
    write_sheet,
)
from pyhandlexl._safety import _is_well_formed, atomic_save, safe_load

MACRO_PAYLOAD = b"FAKE-VBA-PROJECT-BYTES"


def _temp_files(directory: Path) -> list[str]:
    return sorted(p.name for p in directory.iterdir() if ".tmp." in p.name)


# ------------------------------------------------------------- create_workbook


class TestCreateWorkbook:
    @pytest.mark.parametrize("name", ["b.xlsx", "B.XLSX", "b.XlSx", "with space.xlsx", "é.xlsx"])
    def test_only_xlsx_names_are_accepted(self, tmp_path, name):
        create_workbook(tmp_path / name)
        assert is_valid_xlsx(tmp_path / name)

    @pytest.mark.parametrize(
        "name", ["b.xlsm", "b.xltx", "b.xltm", "b.xls", "b.csv", "b.txt", "b", "b.xlsx.bak", "b."]
    )
    def test_any_other_name_is_refused_and_nothing_is_created(self, tmp_path, name):
        with pytest.raises(ValueError, match=r"only makes \.xlsx"):
            create_workbook(tmp_path / name)
        assert list(tmp_path.iterdir()) == []

    def test_a_csv_named_workbook_can_no_longer_be_mistaken_for_a_csv(self, tmp_path):
        with pytest.raises(ValueError):
            create_workbook(tmp_path / "data.csv")

    def test_a_missing_directory_is_named_not_a_temp_file(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="does not exist") as caught:
            create_workbook(tmp_path / "nope" / "b.xlsx")
        assert ".tmp." not in str(caught.value)
        assert "nope" in str(caught.value)

    def test_an_existing_file_is_never_overwritten(self, tmp_path):
        target = tmp_path / "b.xlsx"
        target.write_bytes(b"precious")
        with pytest.raises(FileExistsError):
            create_workbook(target)
        assert target.read_bytes() == b"precious"

    def test_the_first_sheet_can_be_named(self, tmp_path):
        create_workbook(tmp_path / "b.xlsx", sheet="Start")
        assert list_sheets(tmp_path / "b.xlsx") == ["Start"]


# -------------------------------------------------------------- macros & templates


def _make_macro_workbook(path: Path) -> None:
    """A workbook the way Excel writes a macro-enabled one: macro content type + a vbaProject."""
    plain = path.with_suffix(".plain.xlsx")
    wb = Workbook()
    wb.active["A1"] = "data"
    wb.save(plain)
    with zipfile.ZipFile(plain) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                ).replace(
                    b"</Types>",
                    b'<Override PartName="/xl/vbaProject.bin" '
                    b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            if item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdVba" Target="vbaProject.bin" '
                    b'Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject"/>'
                    b"</Relationships>",
                )
            target.writestr(item, data)
        target.writestr("xl/vbaProject.bin", MACRO_PAYLOAD)
    plain.unlink()


def _macros_intact(path: Path) -> bool:
    with zipfile.ZipFile(path) as z:
        return (
            z.read("xl/vbaProject.bin") == MACRO_PAYLOAD
            if "xl/vbaProject.bin" in z.namelist()
            else False
        ) and b"macroEnabled" in z.read("[Content_Types].xml")


@pytest.fixture
def macro_book(tmp_path):
    path = tmp_path / "macro.xlsm"
    _make_macro_workbook(path)
    assert _macros_intact(path)
    return path


class TestMacroEnabledWorkbooks:
    """A write used to strip the macros silently, leaving a file Excel would refuse to open."""

    def test_reading_never_touches_it(self, macro_book):
        before = macro_book.read_bytes()
        read_sheet(macro_book)
        list_sheets(macro_book)
        assert macro_book.read_bytes() == before

    @pytest.mark.parametrize(
        "operation",
        [
            lambda p: write_sheet(p, [["new"]], sheet="S2"),
            lambda p: write_sheet(p, [["over"]]),
            lambda p: append_rows(p, [["more"]]),
            lambda p: create_sheet(p, "Extra"),
            lambda p: rename_sheet(p, "Sheet", "Renamed"),
            lambda p: clear_all_sheet_data(p, "Sheet"),
        ],
        ids=["write new sheet", "overwrite", "append", "create_sheet", "rename", "clear"],
    )
    def test_every_write_keeps_the_macros(self, macro_book, operation):
        operation(macro_book)
        assert _macros_intact(macro_book)
        assert is_valid_xlsx(macro_book)

    def test_tables_can_live_in_it_and_the_macros_survive(self, macro_book):
        create_sheet(macro_book, "Data")
        table = Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T")
        table.create(macro_book, sheet="Data")
        table.add_row("r2", [2])
        table.write(macro_book)
        assert Table.read(macro_book, "T").data.row_labels == ["r", "r2"]
        assert _macros_intact(macro_book)

    def test_deleting_a_sheet_or_table_keeps_the_macros(self, macro_book):
        create_sheet(macro_book, "Data")
        Table(column_headers=["c"], name="T").create(macro_book, sheet="Data")
        delete_table(macro_book, "T")
        delete_sheet(macro_book, "Data")
        assert _macros_intact(macro_book)

    def test_a_macro_template_is_kept_too(self, tmp_path):
        path = tmp_path / "macro.xltm"
        _make_macro_workbook(path)
        write_sheet(path, [["x"]], sheet="S2")
        assert path.suffix == ".xltm"
        with zipfile.ZipFile(path) as z:
            assert z.read("xl/vbaProject.bin") == MACRO_PAYLOAD

    def test_a_plain_xlsx_never_gains_macro_parts(self, tmp_path):
        path = tmp_path / "plain.xlsx"
        create_workbook(path)
        write_sheet(path, [["x"]])
        with zipfile.ZipFile(path) as z:
            assert "xl/vbaProject.bin" not in z.namelist()
            assert b"macroEnabled" not in z.read("[Content_Types].xml")

    def test_using_one_prints_nothing_to_stderr(self, macro_book):
        # openpyxl's macro archive used to raise "Exception ignored ... I/O operation on closed
        # file" from a __del__ at garbage collection: invisible to pytest's own output, so a
        # subprocess is the honest way to see what a user's terminal would.
        code = (
            "import gc, sys, pyhandlexl as p\n"
            "path = sys.argv[1]\n"
            "p.read_sheet(path)\n"
            "p.write_sheet(path, [['x']], sheet='S2')\n"
            "p.list_tables(path)\n"
            "gc.collect()\n"
        )
        result = subprocess.run(
            [sys.executable, "-c", code, str(macro_book)], capture_output=True, text=True
        )
        assert result.returncode == 0, result.stderr
        assert result.stderr == ""


class TestTemplates:
    def test_an_xltx_stays_a_template(self, tmp_path):
        path = tmp_path / "tmpl.xltx"
        wb = Workbook()
        wb.template = True
        wb.save(path)
        write_sheet(path, [["x"]], sheet="S2")
        with zipfile.ZipFile(path) as z:
            assert b"template" in z.read("[Content_Types].xml")


# ---------------------------------------------------------- whole-file validation


class TestNeverReplacingAGoodFileWithABadOne:
    def _poisoned_workbook(self) -> Workbook:
        """A workbook whose sheet XML can't be parsed, built around the value checks."""
        wb = Workbook()
        cell = wb.active["A1"]
        cell._value = "a\ud800b"  # bypass every check on purpose
        cell.data_type = "s"
        return wb

    def test_atomic_save_refuses_it_even_though_it_is_a_valid_zip(self, tmp_path):
        target = tmp_path / "b.xlsx"
        create_workbook(target)
        write_sheet(target, [["precious"]])
        before = target.read_bytes()
        with pytest.raises(InvalidFileError, match="failed validation"):
            atomic_save(self._poisoned_workbook(), target)
        assert target.read_bytes() == before
        assert read_sheet(target) == [["precious"]]
        assert _temp_files(tmp_path) == []

    def test_is_valid_xlsx_alone_would_have_let_it_through(self, tmp_path):
        # why the stronger check exists: this is the gap it closes
        probe = tmp_path / "probe.xlsx"
        self._poisoned_workbook().save(probe)
        assert is_valid_xlsx(probe) is True
        assert _is_well_formed(probe) is False

    def test_well_formed_accepts_a_normal_workbook(self, book):
        write_sheet(book, [["é", "😀", 1, None]])
        assert _is_well_formed(book) is True

    def test_well_formed_rejects_a_broken_part_and_a_non_zip(self, tmp_path, book):
        broken = tmp_path / "broken.xlsx"
        with zipfile.ZipFile(book) as source, zipfile.ZipFile(broken, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"</worksheet>", b"")
                target.writestr(item, data)
        assert _is_well_formed(broken) is False
        garbage = tmp_path / "garbage.xlsx"
        garbage.write_bytes(b"not a zip")
        assert _is_well_formed(garbage) is False
        assert _is_well_formed(tmp_path / "missing.xlsx") is False


# ------------------------------------------------------------- bad input files


class TestUnusableFiles:
    @pytest.mark.parametrize(
        ("content", "what"),
        [(b"", "zero bytes"), (b"this is not a zip", "garbage"), (b"PK\x03\x04junk", "bad zip")],
    )
    def test_an_unreadable_file_is_reported_and_left_exactly_as_it_was(
        self, tmp_path, content, what
    ):
        path = tmp_path / "bad.xlsx"
        path.write_bytes(content)
        for call in (
            lambda: list_sheets(path),
            lambda: read_sheet(path),
            lambda: write_sheet(path, [["x"]]),
            lambda: append_rows(path, [["x"]]),
            lambda: create_sheet(path, "S"),
            lambda: list_tables(path),
            lambda: Table.read(path, "T"),
            lambda: sheet_exists(path, "S"),
        ):
            with pytest.raises(InvalidFileError):
                call()
        assert path.read_bytes() == content, what
        assert _temp_files(tmp_path) == []

    def test_a_zip_that_is_not_a_workbook(self, tmp_path):
        path = tmp_path / "other.xlsx"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("hello.txt", "hi")
        before = path.read_bytes()
        with pytest.raises(InvalidFileError):
            write_sheet(path, [["x"]])
        assert path.read_bytes() == before

    def test_a_directory_is_not_a_workbook(self, tmp_path):
        with pytest.raises(InvalidFileError):
            list_sheets(tmp_path)

    def test_a_missing_file_is_a_file_not_found_everywhere(self, tmp_path):
        missing = tmp_path / "missing.xlsx"
        for call in (
            lambda: list_sheets(missing),
            lambda: read_sheet(missing),
            lambda: write_sheet(missing, [["x"]]),
            lambda: append_rows(missing, [["x"]]),
            lambda: create_sheet(missing, "S"),
            lambda: list_tables(missing),
            lambda: table_info(missing, "T"),
            lambda: sheet_kind(missing, "S"),
            lambda: Table.read(missing, "T"),
            lambda: Table(column_headers=["a"], name="T").create(missing, sheet="S"),
            lambda: delete_workbook(missing),
        ):
            with pytest.raises(FileNotFoundError):
                call()
        assert list(tmp_path.iterdir()) == []  # and nothing was created as a side effect


# ------------------------------------------------ replacing a file keeps it what it was

posix_only = pytest.mark.skipif(
    sys.platform == "win32", reason="POSIX permission and symlink semantics"
)
not_root = pytest.mark.skipif(
    sys.platform != "win32" and hasattr(os, "geteuid") and os.geteuid() == 0,
    reason="root can write to read-only files",
)


class TestReplacingAFileKeepsItWhatItWas:
    """A save writes a new file and swaps it in; that must not change what the file *is*."""

    @posix_only
    def test_permissions_are_kept(self, book):
        os.chmod(book, 0o640)
        write_sheet(book, [["x"]])
        assert stat.S_IMODE(book.stat().st_mode) == 0o640

    @posix_only
    def test_a_symlinked_workbook_stays_a_symlink_and_its_target_is_updated(self, tmp_path):
        real = tmp_path / "real.xlsx"
        create_workbook(real)
        link = tmp_path / "link.xlsx"
        link.symlink_to(real)
        write_sheet(link, [["through the link"]])
        assert link.is_symlink() and link.resolve() == real.resolve()
        assert read_sheet(real) == [["through the link"]]
        assert _temp_files(tmp_path) == []

    @posix_only
    def test_a_workbook_in_a_symlinked_directory(self, tmp_path):
        real_dir = tmp_path / "real"
        real_dir.mkdir()
        (tmp_path / "linked").symlink_to(real_dir)
        create_workbook(real_dir / "b.xlsx")
        write_sheet(tmp_path / "linked" / "b.xlsx", [["ok"]])
        assert read_sheet(real_dir / "b.xlsx") == [["ok"]]

    @posix_only
    def test_a_csv_keeps_its_permissions_and_its_symlink(self, tmp_path):
        real = tmp_path / "real.csv"
        real.write_bytes(b"a\n")
        os.chmod(real, 0o640)
        link = tmp_path / "link.csv"
        link.symlink_to(real)
        write_sheet(link, [["b"]])
        assert link.is_symlink()
        assert stat.S_IMODE(real.stat().st_mode) == 0o640
        assert read_sheet(real) == [["b"]]

    @not_root
    def test_a_read_only_workbook_is_refused_the_same_way_on_every_platform(self, book):
        write_sheet(book, [["keep"]])
        create_sheet(book, "Blank")
        before = book.read_bytes()
        os.chmod(book, stat.S_IREAD)
        try:
            for call in (
                lambda: write_sheet(book, [["nope"]]),
                lambda: append_rows(book, [["nope"]]),
                lambda: create_sheet(book, "S2"),
                lambda: Table(column_headers=["a"], name="T").create(book, sheet="Blank"),
            ):
                with pytest.raises(FileReadOnlyError, match="read-only") as caught:
                    call()
                assert isinstance(caught.value, FileLockedError)  # what 0.9.5 raised
        finally:
            os.chmod(book, stat.S_IWRITE | stat.S_IREAD)
        assert book.read_bytes() == before
        assert _temp_files(book.parent) == []

    @not_root
    def test_a_read_only_csv_is_refused_too(self, tmp_path):
        path = tmp_path / "data.csv"
        path.write_bytes(b"keep\n")
        os.chmod(path, stat.S_IREAD)
        try:
            with pytest.raises(FileReadOnlyError, match="read-only"):
                write_sheet(path, [["nope"]])
            with pytest.raises(PermissionError):  # an append opens the file itself
                append_rows(path, [["nope"]])
        finally:
            os.chmod(path, stat.S_IWRITE | stat.S_IREAD)
        assert path.read_bytes() == b"keep\n"
        assert _temp_files(tmp_path) == []

    @not_root
    def test_read_only_files_can_still_be_read_and_converted(self, book, tmp_path):
        write_sheet(book, [["v"]])
        os.chmod(book, stat.S_IREAD)
        try:
            assert read_sheet(book) == [["v"]]
            export_xl_to_csv(book, tmp_path / "out.csv")
            assert list_tables(book) == []
        finally:
            os.chmod(book, stat.S_IWRITE | stat.S_IREAD)


# --------------------------------------------------------------------- paths


class TestPathForms:
    def test_str_and_path_and_pathlike_all_work(self, book):
        class Like:
            def __fspath__(self) -> str:
                return str(book)

        for form in (str(book), book, Like()):
            assert list_sheets(form) == ["Sheet"]

    def test_a_relative_path(self, book, monkeypatch):
        monkeypatch.chdir(book.parent)
        write_sheet(book.name, [["rel"]])
        assert read_sheet(book.name) == [["rel"]]

    def test_spaces_and_unicode_in_the_path(self, tmp_path):
        folder = tmp_path / "sp ace ünï 日本"
        folder.mkdir()
        path = folder / "a b é.xlsx"
        create_workbook(path)
        create_sheet(path, "S")
        Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T").create(path, sheet="S")
        assert Table.read(path, "T").read_row("r") == [1]
        assert _temp_files(folder) == []

    def test_a_bytes_path_is_a_type_error(self, book):
        with pytest.raises(TypeError):
            list_sheets(str(book).encode())  # type: ignore[arg-type]

    def test_dotdot_and_dot_segments(self, book):
        twisted = book.parent / "." / "sub" / ".." / book.name
        (book.parent / "sub").mkdir()
        write_sheet(twisted, [["x"]])
        assert read_sheet(book) == [["x"]]


# ---------------------------------------------------- nothing else gets touched


def _rich_workbook(path: Path) -> None:
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    for i in range(1, 6):
        data.append([i, i * 2])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=5))
    data.add_chart(chart, "D2")
    other = wb.create_sheet("Other")
    other["A1"] = "merged"
    other.merge_cells("A1:B2")
    other.column_dimensions["A"].width = 44
    other.row_dimensions[3].height = 33
    other.freeze_panes = "A2"
    validation = DataValidation(type="list", formula1='"a,b"')
    other.add_data_validation(validation)
    validation.add("C1")
    other["D1"].comment = Comment("hello", "me")
    other["E1"].hyperlink = "https://example.com"
    other["F1"] = "=1+1"
    other.sheet_properties.tabColor = "00FF00"
    other.sheet_view.showGridLines = False
    wb.defined_names["myname"] = DefinedName("myname", attr_text="Other!$A$1")
    wb.save(path)


def _features(path: Path) -> dict:
    wb = load_workbook(path)
    data, other = wb["Data"], wb["Other"]
    return {
        "charts": len(data._charts),
        "merged": sorted(str(m) for m in other.merged_cells.ranges),
        "col width": other.column_dimensions["A"].width,
        "row height": other.row_dimensions[3].height,
        "freeze": other.freeze_panes,
        "validations": len(other.data_validations.dataValidation),
        "comment": other["D1"].comment.text if other["D1"].comment else None,
        "hyperlink": other["E1"].hyperlink.target if other["E1"].hyperlink else None,
        "formula": other["F1"].value,
        "tab color": str(other.sheet_properties.tabColor.rgb)
        if other.sheet_properties.tabColor
        else None,
        "gridlines": other.sheet_view.showGridLines,
        "names": sorted(wb.defined_names.keys()),
        "sheets": wb.sheetnames,
    }


class TestOtherContentSurvives:
    @pytest.fixture
    def rich(self, tmp_path):
        path = tmp_path / "rich.xlsx"
        _rich_workbook(path)
        return path

    @pytest.mark.parametrize(
        "operation",
        [
            lambda p: write_sheet(p, [["x"]], sheet="Scratch"),
            lambda p: append_rows(p, [["x"]], sheet="Scratch"),
            lambda p: create_sheet(p, "Scratch"),
            lambda p: Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T").create(
                p, sheet=(create_sheet(p, "Scratch"), "Scratch")[1]
            ),
        ],
        ids=["write_sheet", "append_rows", "create_sheet", "Table.create"],
    )
    def test_charts_merges_widths_validations_comments_links_formulas_names(self, rich, operation):
        before = _features(rich)
        operation(rich)
        after = _features(rich)
        after["sheets"] = [s for s in after["sheets"] if s not in ("Scratch", "_pyhandlexl_tables")]
        assert after == before

    def test_a_table_write_leaves_the_neighbouring_sheets_alone(self, rich):
        create_sheet(rich, "Scratch")
        table = Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T")
        table.create(rich, sheet="Scratch")
        before = _features(rich)
        table.add_row("r2", [2])
        table.write(rich)
        assert _features(rich) == before


class TestReadingNeverWrites:
    """A read of a healthy workbook must leave its bytes exactly as they were."""

    @pytest.fixture
    def healthy(self, book):
        create_sheet(book, "Data")
        create_sheet(book, "Grid")
        write_sheet(book, [["g", 1]], sheet="Grid")
        Table(data=[[1, 2]], row_labels=["r"], column_headers=["a", "b"], name="T").create(
            book, sheet="Data"
        )
        return book

    def test_every_read_function(self, healthy, tmp_path):
        before = healthy.read_bytes()
        read_sheet(healthy, "Grid")
        read_sheet(healthy, "Grid", pad=True, orientation="columns")
        list_sheets(healthy)
        sheet_exists(healthy, "Grid")
        sheet_exists(healthy, "Nope")
        sheet_kind(healthy, "Grid")
        sheet_kind(healthy, "Data")
        list_tables(healthy)
        table_info(healthy, "T")
        table = Table.read(healthy, "T")
        table.show()
        table.to_dict()
        _ = table.info, table.data, table.column_types
        export_xl_to_csv(healthy, tmp_path / "out.csv", sheet="Grid")
        assert healthy.read_bytes() == before
        assert _temp_files(healthy.parent) == []

    def test_failed_reads_leave_it_alone_too(self, healthy):
        before = healthy.read_bytes()
        from pyhandlexl import SheetNotFoundError, TableNotFoundError

        with pytest.raises(SheetNotFoundError):
            read_sheet(healthy, "Nope")
        with pytest.raises(TableNotFoundError):
            Table.read(healthy, "Nope")
        with pytest.raises(TableNotFoundError):
            table_info(healthy, "Nope")
        assert healthy.read_bytes() == before


class TestRefusedWritesLeaveTheFileUntouched:
    """Every way a write can be refused, checked on one workbook, byte for byte."""

    @pytest.fixture
    def setup(self, book):
        create_sheet(book, "Data")
        create_sheet(book, "Grid")
        write_sheet(book, [["g"]], sheet="Grid")
        Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T").create(
            book, sheet="Data"
        )
        return book

    def test_each_refusal(self, setup):
        from pyhandlexl import (
            DimensionError,
            SheetKindError,
            SheetNameError,
            SheetNotFoundError,
            TableExistsError,
            TableNotFoundError,
        )

        refusals = [
            (CellTypeError, lambda: write_sheet(setup, [[float("nan")]], sheet="Grid")),
            (CellTypeError, lambda: append_rows(setup, [[object()]], sheet="Grid")),
            (SheetNameError, lambda: create_sheet(setup, "a/b")),
            (ValueError, lambda: create_sheet(setup, "Data")),
            (ValueError, lambda: create_sheet(setup, "DATA")),
            (SheetKindError, lambda: write_sheet(setup, [["x"]], sheet="Data")),
            (SheetKindError, lambda: append_rows(setup, [["x"]], sheet="Data")),
            (
                SheetKindError,
                lambda: Table(column_headers=["c"], name="U").create(setup, sheet="Grid"),
            ),
            (SheetKindError, lambda: write_sheet(setup, [["x"]], sheet="_pyhandlexl_tables")),
            (SheetKindError, lambda: delete_sheet(setup, "_pyhandlexl_tables")),
            (SheetKindError, lambda: rename_sheet(setup, "Grid", "_PYHANDLEXL_TABLES")),
            (
                TableExistsError,
                lambda: Table(column_headers=["c"], name="T").create(setup, sheet="Data"),
            ),
            (TableNotFoundError, lambda: delete_table(setup, "Nope")),
            (SheetNotFoundError, lambda: delete_sheet(setup, "Nope")),
            (SheetNotFoundError, lambda: rename_sheet(setup, "Nope", "X")),
            (
                SheetNotFoundError,
                lambda: Table(column_headers=["c"], name="V").create(setup, sheet="Nope"),
            ),
            (DimensionError, lambda: write_sheet(setup, [[1] * 16_385], sheet="Grid")),
            (ValueError, lambda: rename_sheet(setup, "Grid", "Data")),
            (TypeError, lambda: write_sheet(setup, ["notarow"], sheet="Grid")),
        ]
        before = setup.read_bytes()
        for expected, call in refusals:
            with pytest.raises(expected):
                call()
            assert setup.read_bytes() == before, f"{expected.__name__} refusal changed the file"
        assert _temp_files(setup.parent) == []

    def test_a_table_write_after_the_table_was_deleted(self, setup):
        from pyhandlexl import TableNotFoundError

        table = Table.read(setup, "T")
        delete_table(setup, "T")
        before = setup.read_bytes()
        table.add_row("x", [1])
        with pytest.raises(TableNotFoundError):
            table.write(setup)
        assert setup.read_bytes() == before


class TestRebuildWarningIsNotSwallowedByFilters:
    def test_it_is_a_user_warning_subclass(self):
        assert issubclass(SchemaRebuiltWarning, UserWarning)


class TestSafeLoad:
    def test_returns_a_workbook_that_can_be_closed(self, book):
        wb = safe_load(book)
        try:
            assert wb.sheetnames == ["Sheet"]
        finally:
            wb.close()
