"""Sheet names that would collide with the schema sheet, or with each other, are refused.

Excel treats sheet names case-insensitively and openpyxl silently renames a would-be
duplicate ("data" -> "data1"), so an unguarded call reports success but touches — or
creates — a different sheet from the one asked for.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    SheetKindError,
    Table,
    append_rows,
    clear_all_sheet_data,
    create_sheet,
    delete_sheet,
    list_sheets,
    list_tables,
    read_sheet,
    rename_sheet,
    sheet_kind,
    write_sheet,
)
from pyhandlexl import _multi_table as mt
from pyhandlexl._multi_table import SCHEMA_SHEET

LOOKALIKES = [
    "_PYHANDLEXL_TABLES",
    "_Pyhandlexl_Tables",
    "_pyhandlexl_Tables",
    "_pyhandlexl_tables ",
    " _pyhandlexl_tables",
]
# Whitespace that is a control character makes the name invalid outright (SheetNameError),
# before the reserved-name check is even reached — still refused, just by the earlier rule.
CONTROL_LOOKALIKES = ["\t_PyHandleXL_Tables\t", "_pyhandlexl_tables\n"]


@pytest.fixture
def path(book):
    create_sheet(book, "Data")
    Table(data=[[1]], row_labels=["r"], column_headers=["v"], name="T").create(book, sheet="Data")
    return book


def _sheets(path) -> list[str]:
    return load_workbook(path).sheetnames


class TestTheHelpers:
    @pytest.mark.parametrize("name", [SCHEMA_SHEET, *LOOKALIKES, *CONTROL_LOOKALIKES])
    def test_the_reserved_name_and_its_lookalikes_are_recognised(self, name):
        assert mt.is_schema_name(name)

    @pytest.mark.parametrize("name", CONTROL_LOOKALIKES)
    def test_a_lookalike_with_control_characters_is_refused_as_an_invalid_name(self, path, name):
        from pyhandlexl import SheetNameError

        before = _sheets(path)
        with pytest.raises(SheetNameError):
            create_sheet(path, name)
        assert _sheets(path) == before

    @pytest.mark.parametrize(
        "name", ["Data", "_pyhandlexl_table", "_pyhandlexl_tables1", "x_pyhandlexl_tables", ""]
    )
    def test_other_names_are_not(self, name):
        assert not mt.is_schema_name(name)

    def test_case_clash_finds_only_a_different_capitalisation(self):
        assert mt.case_clash(["Sheet", "Data"], "DATA") == "Data"
        assert mt.case_clash(["Sheet", "Data"], "Data") is None  # the very same name
        assert mt.case_clash(["Sheet", "Data"], "Other") is None

    def test_the_error_names_the_offender_and_the_reserved_name(self):
        exact = str(mt.reserved_error(SCHEMA_SHEET, "cannot be created"))
        lookalike = str(mt.reserved_error("_PYHANDLEXL_TABLES", "cannot be created"))
        assert exact == "'_pyhandlexl_tables' is reserved and cannot be created"
        assert "'_PYHANDLEXL_TABLES'" in lookalike and "'_pyhandlexl_tables'" in lookalike
        assert "too close" in lookalike


class TestNoCallCanCreateOrUseALookalike:
    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_create_sheet_refuses_it_and_creates_nothing(self, path, name):
        before = _sheets(path)
        with pytest.raises(SheetKindError, match="reserved"):
            create_sheet(path, name)
        assert _sheets(path) == before

    def test_create_sheet_refuses_the_exact_name_even_when_no_schema_exists_yet(self, book):
        with pytest.raises(SheetKindError, match="reserved"):
            create_sheet(book, SCHEMA_SHEET)
        assert _sheets(book) == ["Sheet"]

    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_write_sheet_refuses_it(self, path, name):
        before = _sheets(path)
        with pytest.raises(SheetKindError, match="reserved"):
            write_sheet(path, [["x"]], sheet=name)
        assert _sheets(path) == before

    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_append_rows_refuses_it(self, path, name):
        before = _sheets(path)
        with pytest.raises(SheetKindError, match="reserved"):
            append_rows(path, [["x"]], sheet=name)
        assert _sheets(path) == before

    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_rename_sheet_refuses_it_as_a_new_name(self, path, name):
        before = _sheets(path)
        with pytest.raises(SheetKindError, match="reserved"):
            rename_sheet(path, "Data", name)
        assert _sheets(path) == before
        assert Table.read(path, "T").read_row("r") == [1]  # the table didn't lose its sheet

    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_table_create_refuses_it(self, path, name):
        with pytest.raises(SheetKindError, match="reserved"):
            Table(column_headers=["a"], name="N").create(path, sheet=name)

    @pytest.mark.parametrize("name", LOOKALIKES)
    def test_the_other_sheet_operations_refuse_it_too(self, path, name):
        for call in (
            lambda: clear_all_sheet_data(path, name),
            lambda: delete_sheet(path, name),
            lambda: sheet_kind(path, name),
        ):
            with pytest.raises(SheetKindError, match="reserved"):
                call()

    def test_import_csv_to_xl_refuses_it(self, path, tmp_path):
        from pyhandlexl import import_csv_to_xl

        source = tmp_path / "in.csv"
        source.write_text("a,b\n1,2\n")
        with pytest.raises(SheetKindError, match="reserved"):
            import_csv_to_xl(source, path, sheet="_PYHANDLEXL_TABLES")

    def test_a_refused_lookalike_leaves_the_real_schema_working(self, path):
        with pytest.raises(SheetKindError):
            create_sheet(path, "_PYHANDLEXL_TABLES")
        assert list_tables(path) == ["T"]
        assert SCHEMA_SHEET in _sheets(path)
        assert [n for n in _sheets(path) if n.lower() == SCHEMA_SHEET] == [SCHEMA_SHEET]


class TestOrdinaryCaseClashesAreRefusedToo:
    def test_create_sheet_refuses_a_different_capitalisation(self, path):
        with pytest.raises(ValueError, match="case-insensitive"):
            create_sheet(path, "DATA")
        assert "DATA1" not in _sheets(path) and "DATA" not in _sheets(path)

    def test_the_exact_duplicate_message_is_unchanged(self, path):
        with pytest.raises(ValueError, match=r"sheet 'Data' already exists$"):
            create_sheet(path, "Data")

    def test_write_sheet_will_not_quietly_write_to_a_new_data1_sheet(self, path):
        create_sheet(path, "Log")
        with pytest.raises(ValueError, match="case-insensitive"):
            write_sheet(path, [["x"]], sheet="log")
        assert "log1" not in _sheets(path)
        assert read_sheet(path, "Log") == []

    def test_append_rows_refuses_it_too(self, path):
        create_sheet(path, "Log")
        with pytest.raises(ValueError, match="case-insensitive"):
            append_rows(path, [["x"]], sheet="LOG")
        assert list_sheets(path) == ["Sheet", "Data", "Log"]

    def test_rename_sheet_refuses_a_name_another_sheet_has_in_other_letters(self, path):
        create_sheet(path, "Log")
        with pytest.raises(ValueError, match="case-insensitive"):
            rename_sheet(path, "Sheet", "LOG")
        assert list_sheets(path) == ["Sheet", "Data", "Log"]

    def test_writing_to_an_existing_sheet_by_its_exact_name_still_works(self, path):
        create_sheet(path, "Log")
        write_sheet(path, [["x"]], sheet="Log")
        append_rows(path, [["y"]], sheet="Log")
        assert read_sheet(path, "Log") == [["x"], ["y"]]


class TestRenamingOnlyTheCapitalisation:
    def test_it_works_and_keeps_the_requested_name(self, path):
        create_sheet(path, "log")
        write_sheet(path, [["keep me"]], sheet="log")
        rename_sheet(path, "log", "LOG")
        assert list_sheets(path) == ["Sheet", "Data", "LOG"]  # not "LOG1"
        assert read_sheet(path, "LOG") == [["keep me"]]

    def test_tables_on_the_sheet_follow_it(self, path):
        rename_sheet(path, "Data", "DATA")
        assert list_sheets(path) == ["Sheet", "DATA"]
        assert Table.read(path, "T").read_row("r") == [1]
        assert Table.read(path, "T").info.sheet == "DATA"

    def test_renaming_to_the_same_name_is_still_a_no_op(self, path):
        rename_sheet(path, "Data", "Data")
        assert list_sheets(path) == ["Sheet", "Data"]


class TestAHandMadeCollisionFailsLoudly:
    def test_saving_the_schema_refuses_rather_than_naming_it_tables1(self, book):
        wb = load_workbook(book)
        wb.create_sheet("_PYHANDLEXL_TABLES")  # only a hand-edited file can hold this
        wb.save(book)
        create_sheet(book, "Data")
        with pytest.raises(SheetKindError, match="collides with the reserved"):
            Table(data=[[1]], row_labels=["r"], column_headers=["v"], name="T").create(
                book, sheet="Data"
            )
        assert SCHEMA_SHEET not in _sheets(book)
        assert "_pyhandlexl_tables1" not in _sheets(book)
