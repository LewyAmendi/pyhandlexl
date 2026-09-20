"""A write repaints only the cells whose look changes — and ends up with exactly the look a
full repaint would give.

The reference is the same code with ``stable_extent`` forced to ``(0, 0)``, which clears and
repaints the whole table on every write, as every version up to 0.9.6 did.
"""

from __future__ import annotations

import random
from copy import copy
from datetime import date, datetime, time, timedelta
from unittest import mock

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from pyhandlexl import Table, TableStyle, create_sheet
from pyhandlexl import _multi_table as mt
from pyhandlexl._multi_table import TableEntry, stable_extent

CUSTOM = TableStyle(header_fill="00FF00", band_fill="FFFF00", border_color="123456")
STYLES = [TableStyle.DEFAULT, TableStyle.MINIMAL, TableStyle.NONE, CUSTOM]
VALUES = [
    1,
    2.5,
    -3,
    True,
    "text",
    "a\nb",
    None,
    datetime(2026, 3, 4, 5, 6, 7),
    date(2026, 3, 4),
    time(1, 2, 3),
    timedelta(hours=5),
    timedelta(days=2, hours=1),
]


def full_repaint():
    """The reference: every write clears and repaints the whole table."""
    return mock.patch.object(mt, "stable_extent", lambda *args: (0, 0))


def look(cell) -> tuple:
    return (
        cell.value,
        type(cell.value),
        copy(cell.font),
        copy(cell.fill),
        copy(cell.border),
        copy(cell.alignment),
        cell.number_format,
        copy(cell.protection),
        cell.has_style,
    )


def sheet_looks(path, sheet="D") -> dict[tuple[int, int], tuple]:
    ws = load_workbook(path)[sheet]
    return {
        (r, c): look(ws.cell(row=r, column=c))
        for r in range(1, ws.max_row + 4)
        for c in range(1, ws.max_column + 4)
    }


def entry(n_rows, n_cols, style=TableStyle.DEFAULT) -> TableEntry:
    return TableEntry("T", "D", 1, 1, n_rows, n_cols, style, [], created_at=None, modified_at=None)


class TestStableExtent:
    def test_nothing_changes_so_everything_is_stable(self):
        assert stable_extent(entry(4, 3), 4, 3, TableStyle.DEFAULT) == (5, 4)

    def test_a_new_style_repaints_everything(self):
        assert stable_extent(entry(4, 3), 4, 3, TableStyle.NONE) == (0, 0)
        assert stable_extent(entry(4, 3), 9, 9, TableStyle.NONE) == (0, 0)

    @pytest.mark.parametrize(
        ("old", "new", "expected"),
        [
            ((4, 3), (5, 3), (4, 4)),  # a row added: the old last row (local 4) is repainted
            ((4, 3), (4, 4), (5, 3)),  # a column added: the old last column is repainted
            ((4, 3), (3, 3), (3, 4)),  # a row dropped: the new last row is repainted
            ((4, 3), (4, 2), (5, 2)),
            ((4, 3), (7, 5), (4, 3)),
            ((4, 3), (1, 1), (1, 1)),
            ((0, 1), (2, 1), (0, 2)),  # an empty table's header row has the bottom border
            ((2, 1), (0, 1), (0, 2)),
        ],
    )
    def test_growing_or_shrinking_touches_only_the_last_row_and_column(self, old, new, expected):
        assert stable_extent(entry(*old), *new, TableStyle.DEFAULT) == expected


class TestSameAsAFullRepaint:
    @pytest.mark.parametrize("seed", range(30))
    def test_random_edits_to_two_side_by_side_tables(self, tmp_path, seed):
        """Rows and columns added and dropped, values of every type set, the style changed —
        with a neighbour that gets shifted right — must leave every cell exactly as a full
        repaint would."""
        paths = {"fast": tmp_path / "fast.xlsx", "full": tmp_path / "full.xlsx"}
        rng = random.Random(seed)
        script = self.script(rng)
        for kind, path in paths.items():
            from pyhandlexl import create_workbook

            create_workbook(path)
            create_sheet(path, "D")
            for name, shape in (("T", (3, 2)), ("N", (2, 2))):
                Table(
                    data=[[i * 10 + j for j in range(shape[1])] for i in range(shape[0])],
                    row_labels=[f"{name}r{i}" for i in range(shape[0])],
                    column_headers=[f"{name}c{j}" for j in range(shape[1])],
                    name=name,
                ).create(path, sheet="D")
            self.replay(path, script, mock_full=(kind == "full"))
        assert sheet_looks(paths["fast"]) == sheet_looks(paths["full"])

    @staticmethod
    def script(rng: random.Random) -> list[tuple]:
        steps: list[tuple] = []
        for n in range(10):
            op = rng.choice(["add_row", "add_row", "drop_row", "add_column", "drop_column"])
            op = rng.choice([op, "set", "set", "insert_row", "style", "set_row"])
            steps.append(
                (
                    op,
                    n,
                    rng.randrange(1000),
                    rng.choice(["T", "T", "N"]),
                    [rng.choice(VALUES) for _ in range(12)],
                    rng.choice(STYLES),
                )
            )
        return steps

    @staticmethod
    def replay(path, script, *, mock_full: bool) -> None:
        for op, n, pick, name, values, style in script:
            t = Table.read(path, name)
            width = len(t.data.column_headers)
            height = len(t.data.row_labels)
            if op == "add_row":
                t.add_row(f"x{n}", values[:width])
            elif op == "insert_row":
                t.insert_row(pick % (height + 1) + 1, f"i{n}", values[:width])
            elif op == "drop_row" and height > 1:
                t.drop_row(t.data.row_labels[pick % height])
            elif op == "add_column":
                t.add_column(f"k{n}", values[:height])
            elif op == "drop_column" and width > 1:
                t.drop_column(t.data.column_headers[pick % width])
            elif op == "set" and height:
                t.set_cell(
                    row=t.data.row_labels[pick % height],
                    column=t.data.column_headers[pick % width],
                    value=values[0],
                )
            elif op == "set_row" and height:
                t.set_row(t.data.row_labels[pick % height], values[:width])
            elif op == "style":
                t.style = style
            if mock_full:
                with full_repaint():
                    t.write(path)
            else:
                t.write(path)


class TestNumberFormatsDontLeak:
    def test_a_number_written_over_a_date_reads_back_as_a_number(self, book):
        create_sheet(book, "D")
        Table(
            data=[[datetime(2026, 1, 1), time(1, 0)], [timedelta(hours=1), date(2026, 2, 2)]],
            row_labels=["a", "b"],
            column_headers=["x", "y"],
            name="T",
        ).create(book, sheet="D")
        t = Table.read(book, "T")
        for row in ("a", "b"):
            for column in ("x", "y"):
                t.set_cell(row=row, column=column, value=7)
        t.write(book)
        back = Table.read(book, "T")
        assert back.data.rows == [[7, 7], [7, 7]]
        ws = load_workbook(book)["D"]
        assert {ws.cell(row=r, column=c).number_format for r in (3, 4) for c in (2, 3)} == {
            "General"
        }

    def test_a_date_written_over_a_number_reads_back_as_a_date(self, book):
        create_sheet(book, "D")
        Table(data=[[1, 2]], row_labels=["a"], column_headers=["x", "y"], name="T").create(
            book, sheet="D"
        )
        t = Table.read(book, "T")
        t.set_row("a", [datetime(2026, 1, 1, 12), timedelta(hours=30)])
        t.write(book)
        assert Table.read(book, "T").data.rows == [[datetime(2026, 1, 1, 12), timedelta(hours=30)]]

    def test_dates_that_stay_dates_keep_their_format_across_writes(self, book):
        create_sheet(book, "D")
        Table(
            data=[[date(2026, 1, 1), time(2, 0)]],
            row_labels=["a"],
            column_headers=["x", "y"],
            name="T",
        ).create(book, sheet="D")
        for n in range(3):
            t = Table.read(book, "T")
            t.add_row(f"b{n}", [date(2026, 3, 3), time(3, 0)])
            t.write(book)
        assert Table.read(book, "T").data.rows == [
            [datetime(2026, 1, 1), time(2, 0)],
            [datetime(2026, 3, 3), time(3, 0)],
            [datetime(2026, 3, 3), time(3, 0)],
            [datetime(2026, 3, 3), time(3, 0)],
        ]


class TestOnlyWhatChangedIsRepainted:
    @pytest.fixture
    def table(self, book):
        create_sheet(book, "D")
        Table(
            data=[[i * 3 + j for j in range(3)] for i in range(6)],
            row_labels=[f"r{i}" for i in range(6)],
            column_headers=["a", "b", "c"],
            name="T",
        ).create(book, sheet="D")
        return book

    @staticmethod
    def hand_format(book, ref: str) -> None:
        wb = load_workbook(book)
        wb["D"][ref].fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        wb.save(book)

    @staticmethod
    def fill_of(book, ref: str) -> str | None:
        return load_workbook(book)["D"][ref].fill.start_color.rgb

    def test_editing_a_value_paints_no_cell(self, table):
        t = Table.read(table, "T")
        t.set_cell(row="r2", column="b", value=-1)
        with (
            mock.patch.object(mt, "_paint_data") as data,
            mock.patch.object(mt, "_paint_header_like") as head,
        ):
            t.write(table)
        assert data.call_count == 0 and head.call_count == 0

    def test_formatting_you_added_by_hand_survives_an_ordinary_write(self, table):
        self.hand_format(table, "C4")
        t = Table.read(table, "T")
        t.set_cell(row="r1", column="b", value=-1)
        t.write(table)
        assert self.fill_of(table, "C4") == "00FF0000"
        assert Table.read(table, "T").read_cell(row="r1", column="b") == -1

    def test_adding_a_row_leaves_the_middle_of_the_table_alone(self, table):
        self.hand_format(table, "C4")  # a middle row
        t = Table.read(table, "T")
        t.add_row("new", [1, 2, 3])
        t.write(table)
        assert self.fill_of(table, "C4") == "00FF0000"

    def test_the_old_last_row_is_repainted_when_the_table_grows(self, table):
        self.hand_format(table, "C8")  # the last data row
        t = Table.read(table, "T")
        t.add_row("new", [1, 2, 3])
        t.write(table)
        assert self.fill_of(table, "C8") != "00FF0000"

    def test_changing_the_style_repaints_everything(self, table):
        self.hand_format(table, "C4")
        t = Table.read(table, "T")
        t.style = TableStyle.MINIMAL
        t.write(table)
        assert self.fill_of(table, "C4") != "00FF0000"

    def test_a_table_shifted_by_a_growing_neighbour_keeps_its_look(self, book):
        create_sheet(book, "D")
        for name in ("A", "B"):
            Table(
                data=[[1, 2], [3, 4], [5, 6]],
                row_labels=["x", "y", "z"],
                column_headers=["p", "q"],
                name=name,
            ).create(book, sheet="D")
        before = sheet_looks(book)
        a = Table.read(book, "A")
        a.add_column("r", [7, 8, 9])
        a.write(book)
        moved = sheet_looks(book)
        # B moved one column right, and every cell of it looks as it did
        for r in range(1, 6):
            for c in range(5, 8):
                assert moved[(r, c + 1)][2:] == before[(r, c)][2:], (r, c)
