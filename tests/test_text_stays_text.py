"""Text that merely starts with ``=`` must stay text where the library promises text.

openpyxl turns any string starting with ``=`` into a formula. That is wrong for a row label,
a column header, the corner, a table's name (all documented as always ``str``) and for a
field imported from a CSV file (documented as plain text): Excel would calculate it — and
show ``#NAME?``, or refuse the file, if it isn't a valid formula.
"""

from __future__ import annotations

import re
import zipfile
from copy import copy

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    Table,
    create_sheet,
    delete_table,
    import_csv_to_xl,
    list_tables,
    read_sheet,
    write_sheet,
)

STRAY = ["=1+1", '=HYPERLINK("http://example.com","x")', "=== NOTES ===", "= 5 apples", "=)", "=="]


def formula_cells(path, part="xl/worksheets/sheet1.xml") -> list[str]:
    with zipfile.ZipFile(path) as z:
        text = z.read(part).decode()
    return re.findall(r'<c r="(\w+)"[^>]*><f>', text)


def sheet(path, name):
    return load_workbook(path)[name]


class TestCsvImportIsPlainText:
    @pytest.fixture
    def csv_file(self, tmp_path):
        path = tmp_path / "in.csv"
        rows = [["plain", "007", *STRAY[:3]], [STRAY[3], STRAY[4], STRAY[5], "x", ""]]
        import csv

        with path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)
        return path

    def test_nothing_is_stored_as_a_formula(self, book, csv_file):
        import_csv_to_xl(csv_file, book, sheet="Sheet")
        assert formula_cells(book) == []
        ws = sheet(book, "Sheet")
        assert {c.data_type for row in ws.iter_rows() for c in row if c.value is not None} == {"s"}

    def test_every_field_reads_back_as_it_was_written(self, book, csv_file):
        import_csv_to_xl(csv_file, book, sheet="Sheet")
        assert read_sheet(book) == [
            ["plain", "007", "=1+1", '=HYPERLINK("http://example.com","x")', "=== NOTES ==="],
            ["= 5 apples", "=)", "==", "x"],
        ]

    def test_the_text_keeps_excels_quote_prefix_so_editing_it_stays_text(self, book, csv_file):
        import_csv_to_xl(csv_file, book, sheet="Sheet")
        ws = sheet(book, "Sheet")
        assert ws["C1"].quotePrefix and ws["A2"].quotePrefix
        assert not ws["A1"].quotePrefix and not ws["B1"].quotePrefix  # ordinary fields

    def test_importing_over_an_existing_sheet_still_works(self, book, csv_file):
        write_sheet(book, [["old", "=1+1"]])
        import_csv_to_xl(csv_file, book, sheet="Sheet")
        assert formula_cells(book) == []

    def test_the_grid_functions_store_it_as_text_too(self, book):
        write_sheet(book, [["=1+1"]])
        assert read_sheet(book) == [["=1+1"]]
        assert formula_cells(book) == []


class TestTableTextStaysText:
    @pytest.fixture
    def table(self, book):
        create_sheet(book, "D")
        Table(
            data=[[1, 2], [3, 4]],
            row_labels=["=lbl1", "=== r2 ==="],
            column_headers=["=hdr", "= 5 apples"],
            corner="=corner",
            name="=NAME",
        ).create(book, sheet="D")
        return book

    def test_the_marker_corner_headers_and_labels_are_text_in_the_file(self, table):
        assert formula_cells(table, "xl/worksheets/sheet2.xml") == []
        ws = sheet(table, "D")
        for ref in ("B1", "A2", "B2", "C2", "A3", "A4"):  # name, corner, headers, labels
            assert ws[ref].data_type == "s", ref
            assert ws[ref].quotePrefix, ref
        assert not ws["B3"].quotePrefix  # data is untouched

    def test_they_read_back_exactly(self, table):
        t = Table.read(table, "=NAME")
        assert t.data.corner == "=corner"
        assert t.data.row_labels == ["=lbl1", "=== r2 ==="]
        assert t.data.column_headers == ["=hdr", "= 5 apples"]
        assert t.data.rows == [[1, 2], [3, 4]]

    def test_addressing_by_such_a_label_works(self, table):
        t = Table.read(table, "=NAME")
        assert t.read_cell(row="=lbl1", column="=hdr") == 1
        t.set_cell(row="=== r2 ===", column="= 5 apples", value=9)
        t.write(table)
        assert Table.read(table, "=NAME").read_row("=== r2 ===") == [3, 9]

    def test_they_stay_text_through_edits_and_rewrites(self, table):
        t = Table.read(table, "=NAME")
        t.add_row("=new", [5, 6])
        t.add_column("=col", [7, 8, 9])
        t.rename_row("=lbl1", "=renamed")
        t.set_corner("=other")
        t.write(table)
        assert formula_cells(table, "xl/worksheets/sheet2.xml") == []
        t = Table.read(table, "=NAME")
        assert t.data.row_labels == ["=renamed", "=== r2 ===", "=new"]
        assert t.data.column_headers == ["=hdr", "= 5 apples", "=col"]
        assert t.data.corner == "=other"

    def test_a_label_that_stops_starting_with_an_equals_sign_loses_the_prefix(self, table):
        t = Table.read(table, "=NAME")
        t.rename_row("=lbl1", "lbl1")
        t.write(table)
        ws = sheet(table, "D")
        assert ws["A3"].value == "lbl1" and not ws["A3"].quotePrefix
        assert ws["A4"].quotePrefix  # the one that still starts with "="

    def test_a_label_that_starts_with_an_equals_sign_gains_the_prefix(self, table):
        t = Table.read(table, "=NAME")
        t.rename_row("=lbl1", "lbl1")
        t.write(table)
        t = Table.read(table, "=NAME")
        t.rename_row("lbl1", "=lbl1")
        t.write(table)  # the same cell, rewritten without being repainted
        ws = sheet(table, "D")
        assert ws["A3"].value == "=lbl1" and ws["A3"].data_type == "s" and ws["A3"].quotePrefix

    def test_they_look_like_every_other_label(self, table):
        ws = sheet(table, "D")
        plain, stray = copy(ws["A3"]), copy(ws["A4"])
        create_sheet(table, "E")
        Table(data=[[1]], row_labels=["r"], column_headers=["h"], name="P").create(table, "E")
        control = sheet(table, "E")["A3"]
        for cell in (plain, stray):
            assert copy(cell.font) == copy(control.font)
            assert copy(cell.fill) == copy(control.fill)
            assert copy(cell.border).left == copy(control.border).left

    def test_a_neighbour_growing_shifts_a_table_without_turning_its_text_into_formulas(self, book):
        create_sheet(book, "D")
        Table(data=[[1]], row_labels=["a"], column_headers=["x"], name="A").create(book, "D")
        Table(data=[[2]], row_labels=["=b"], column_headers=["=y"], corner="=c", name="=B").create(
            book, "D"
        )
        a = Table.read(book, "A")
        a.add_column("z", [5])
        a.write(book)
        assert formula_cells(book, "xl/worksheets/sheet2.xml") == []
        moved = Table.read(book, "=B")
        assert moved.data.row_labels == ["=b"] and moved.data.column_headers == ["=y"]
        assert moved.data.corner == "=c"
        ws = sheet(book, "D")
        assert ws["E3"].quotePrefix  # the label, now one column further right

    def test_editing_only_a_value_keeps_the_text_cells_as_they_were(self, table):
        t = Table.read(table, "=NAME")
        t.set_cell(row="=lbl1", column="=hdr", value=42)
        t.write(table)  # nothing is repainted, the cells are only rewritten
        ws = sheet(table, "D")
        assert ws["A3"].quotePrefix and ws["B2"].quotePrefix
        assert ws["B3"].value == 42 and not ws["B3"].quotePrefix

    def test_only_the_cells_that_need_it_have_the_quote_prefix(self, book):
        # plain and "=" text alternate, so a style copied from one cell to the next can't
        # hide a prefix given to the wrong cell (or one that went missing)
        create_sheet(book, "D")
        labels = ["a", "=b", "c", "=d", "e", "f", "=g"]
        headers = ["x", "=y", "z", "=w"]
        Table(
            data=[[i] * 4 for i in range(7)], row_labels=labels, column_headers=headers, name="T"
        ).create(book, "D")
        for _ in range(2):  # a fresh create, then an ordinary rewrite of the same table
            ws = sheet(book, "D")
            assert [ws.cell(row=r, column=1).quotePrefix for r in range(3, 10)] == [
                v.startswith("=") for v in labels
            ]
            assert [ws.cell(row=2, column=c).quotePrefix for c in range(2, 6)] == [
                v.startswith("=") for v in headers
            ]
            assert not any(
                ws.cell(row=r, column=c).quotePrefix for r in range(3, 10) for c in range(2, 6)
            )
            t = Table.read(book, "T")
            t.set_cell(row="c", column="x", value=99)
            t.write(book)

    def test_the_schema_sheet_is_text_too(self, table):
        assert list_tables(table) == ["=NAME"]
        names = [
            i for i in zipfile.ZipFile(table).namelist() if i.startswith("xl/worksheets/sheet")
        ]
        for part in names:
            assert formula_cells(table, part) == [], part
        delete_table(table, "=NAME")
        assert list_tables(table) == []


class TestASheetNamedWithAnEqualsSign:
    def test_a_table_on_such_a_sheet_round_trips(self, book):
        create_sheet(book, "=S")
        Table(data=[[1]], row_labels=["r"], column_headers=["h"], name="T").create(book, "=S")
        t = Table.read(book, "T")
        t.add_row("r2", [2])
        t.write(book)
        assert Table.read(book, "T").data.rows == [[1], [2]]
        assert Table.read(book, "T").info.sheet == "=S"
        for part in (
            "xl/worksheets/sheet1.xml",
            "xl/worksheets/sheet2.xml",
            "xl/worksheets/sheet3.xml",
        ):
            assert formula_cells(book, part) == [], part
