"""pyhandlexl does not support formulas.

It never writes one: a string is text, even if it starts with ``=``. And when it reads a
workbook that holds real formulas it says so (``FormulaWarning``), because reading a range
and writing it back stores each formula's text in its place.
"""

from __future__ import annotations

import re
import warnings
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from pyhandlexl import (
    FormulaWarning,
    Table,
    append_rows,
    create_sheet,
    export_xl_to_csv,
    read_sheet,
    write_sheet,
)

STRAY = [
    "=1+1",
    '=HYPERLINK("http://example.com","x")',
    "=== NOTES ===",
    "= 5 apples",
    "=)",
    "=SUM(",
]


def formulas_in_file(path) -> int:
    with zipfile.ZipFile(path) as z:
        return sum(
            len(re.findall(r"<f>", z.read(name).decode()))
            for name in z.namelist()
            if name.startswith("xl/worksheets/sheet")
        )


@pytest.fixture
def formula_book(tmp_path):
    """A workbook as Excel might have produced it: real formulas on a grid sheet and in a table."""
    path = tmp_path / "authored.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Grid"
    ws["A1"], ws["B1"], ws["C1"] = 1, 2, "=A1+B1"
    ws["A2"] = "=SUM(A1:B1)"
    ws["A4"].value = "=== not a formula ==="  # openpyxl makes a formula of this ...
    ws["A4"].data_type = "s"  # ... so store it as the text it is
    wb.save(path)
    create_sheet(path, "D")
    Table(
        data=[[1, 2, None], [3, 4, None]],
        row_labels=["r1", "r2"],
        column_headers=["a", "b", "total"],
        name="T",
    ).create(path, sheet="D")
    wb = load_workbook(path)
    wb["D"]["D3"] = "=B3+C3"
    wb["D"]["D4"] = "=B4+C4"
    wb.save(path)
    return path


class TestNothingWritesAFormula:
    def test_grid_cells(self, book):
        write_sheet(book, [STRAY[:3], STRAY[3:]])
        append_rows(book, [STRAY])
        assert formulas_in_file(book) == 0
        assert read_sheet(book) == [STRAY[:3], STRAY[3:], STRAY]

    def test_table_cells_of_every_kind(self, book):
        create_sheet(book, "D")
        Table(
            data=[[STRAY[0], STRAY[1]], [STRAY[2], 5]],
            row_labels=[STRAY[3], "r2"],
            column_headers=[STRAY[4], "b"],
            corner=STRAY[5],
            name="T",
        ).create(book, sheet="D")
        t = Table.read(book, "T")
        t.set_cell(row="r2", column="b", value="=A1*2")
        t.add_row("r3", ["=B3", "=SUM(B3:B4)"])
        t.write(book)
        assert formulas_in_file(book) == 0
        back = Table.read(book, "T")
        assert back.data.rows == [[STRAY[0], STRAY[1]], [STRAY[2], "=A1*2"], ["=B3", "=SUM(B3:B4)"]]
        assert back.data.corner == STRAY[5]

    def test_every_such_cell_is_a_string_with_a_quote_prefix(self, book):
        write_sheet(book, [["=1+1", "plain"]])
        ws = load_workbook(book).active
        assert ws["A1"].data_type == "s" and ws["A1"].quotePrefix
        assert ws["B1"].data_type == "s" and not ws["B1"].quotePrefix

    def test_our_own_text_is_never_reported_as_a_formula(self, book):
        write_sheet(book, [STRAY])
        create_sheet(book, "D")
        Table(data=[[STRAY[0]]], row_labels=[STRAY[1]], column_headers=[STRAY[2]], name="T").create(
            book, sheet="D"
        )
        with warnings.catch_warnings():
            warnings.simplefilter("error")  # a FormulaWarning would fail the test
            read_sheet(book, "Sheet")
            Table.read(book, "T")
            export_xl_to_csv(book, book.with_suffix(".csv"), sheet="Sheet")


class TestReadingRealFormulasWarns:
    def test_read_sheet_reads_the_text_and_warns(self, formula_book):
        with pytest.warns(FormulaWarning, match=r"sheet 'Grid' holds 2 formula cells \(C1, A2\)"):
            rows = read_sheet(formula_book, "Grid")
        assert rows == [[1, 2, "=A1+B1"], ["=SUM(A1:B1)"], [], ["=== not a formula ==="]]

    def test_table_read_reads_the_text_and_warns(self, formula_book):
        with pytest.warns(FormulaWarning, match=r"table 'T' holds 2 formula cells \(D3, D4\)"):
            t = Table.read(formula_book, "T")
        assert t.data.rows == [[1, 2, "=B3+C3"], [3, 4, "=B4+C4"]]

    def test_a_formula_outside_the_table_is_not_the_tables_business(self, formula_book):
        wb = load_workbook(formula_book)
        wb["D"]["H1"] = "=1+1"  # on the same sheet, well clear of the table
        wb.save(formula_book)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            Table.read(formula_book, "T")
        assert [w.message.args[0][:22] for w in caught if w.category is FormulaWarning] == [
            "table 'T' holds 2 form"
        ]

    def test_the_message_says_what_happens_and_what_to_do(self, formula_book):
        with pytest.warns(FormulaWarning) as caught:
            read_sheet(formula_book, "Grid")
        message = str(caught[0].message)
        assert "doesn't support formulas" in message
        assert "as their text" in message and "replacing the formula" in message
        assert "openpyxl" in message

    def test_many_formulas_are_summarised(self, tmp_path):
        path = tmp_path / "many.xlsx"
        wb = Workbook()
        for row in range(1, 13):
            wb.active.cell(row=row, column=1, value="=1+1")
        wb.save(path)
        with pytest.warns(
            FormulaWarning, match=r"12 formula cells \(A1, A2, A3, A4, A5 and 7 more\)"
        ):
            read_sheet(path)

    def test_a_sheet_without_formulas_is_silent(self, book):
        write_sheet(book, [[1, "x", None]])
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            read_sheet(book)

    def test_writing_it_back_stores_the_text_and_says_so_in_the_file(self, formula_book):
        with pytest.warns(FormulaWarning):
            rows = read_sheet(formula_book, "Grid")
        write_sheet(formula_book, rows, sheet="Grid")
        assert formulas_in_file(formula_book) == 2  # only the two in the table are left
        cell = load_workbook(formula_book)["Grid"]["C1"]
        assert cell.data_type == "s" and cell.value == "=A1+B1"

    def test_formulas_on_other_sheets_are_left_alone(self, formula_book):
        # untouched by a write elsewhere (Excel or LibreOffice recalculates them on opening)
        write_sheet(formula_book, [["x"]], sheet="Other")
        wb = load_workbook(formula_book)
        assert wb["Grid"]["C1"].value == "=A1+B1" and wb["Grid"]["C1"].data_type == "f"
        assert wb["D"]["D3"].data_type == "f"


class TestTheWarningNamesTheCallersLine:
    def test_read_sheet(self, formula_book):
        with pytest.warns(FormulaWarning) as caught:
            read_sheet(formula_book, "Grid")
        assert caught[0].filename == __file__

    def test_table_read(self, formula_book):
        with pytest.warns(FormulaWarning) as caught:
            Table.read(formula_book, "T")
        assert caught[0].filename == __file__

    def test_export_xl_to_csv(self, formula_book, tmp_path):
        with pytest.warns(FormulaWarning) as caught:
            export_xl_to_csv(formula_book, tmp_path / "out.csv", sheet="Grid")
        assert [w.filename for w in caught] == [__file__]
        assert (tmp_path / "out.csv").read_text().splitlines()[0] == "1,2,=A1+B1"


def test_a_csv_file_is_text_and_never_warns(tmp_path):
    path = tmp_path / "d.csv"
    path.write_text("=1+1,x\n", encoding="utf-8")
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        assert read_sheet(path) == [["=1+1", "x"]]
