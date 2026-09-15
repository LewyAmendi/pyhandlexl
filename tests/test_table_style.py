"""Tests for Table styling: painting geometry, auto-extend, persistence, shifts."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pyhandlexl import Table, TableStyle, create_sheet, delete_table, read_sheet


@pytest.fixture
def data_sheet(book):
    create_sheet(book, "Data")
    return book


def _cell(book, row, col, sheet="Data"):
    return load_workbook(book)[sheet].cell(row=row, column=col)


class TestDefaultsOnByDefault:
    def test_new_table_gets_default_style(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T")
        assert t.style == TableStyle.DEFAULT

    def test_create_paints_the_default_look(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        header = _cell(data_sheet, 2, 1)  # corner cell
        assert header.font.bold is True
        assert header.fill.fill_type == "solid"

    def test_explicit_style_overrides_default(self, data_sheet):
        t = Table(
            data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style=TableStyle.NONE
        )
        t.create(data_sheet, sheet="Data")
        header = _cell(data_sheet, 2, 1)
        assert header.font.bold is False
        assert header.border.top is None


class TestPaintGeometry:
    """A 2 data-row x 2 data-column table: header row 2, data rows 3-4, label
    col A, data cols B-C — local (r, c) with r in 0..2, c in 0..2."""

    @pytest.fixture
    def painted(self, data_sheet):
        Table(
            data=[[1, 2], [3, 4]],
            column_headers=["h1", "h2"],
            row_labels=["r1", "r2"],
            corner="#",
            name="T",
        ).create(data_sheet, sheet="Data")
        return data_sheet

    def test_header_row_is_bold_and_filled(self, painted):
        for col in (1, 2, 3):
            cell = _cell(painted, 2, col)
            assert cell.font.bold is True
            assert cell.font.color.rgb == "00000000"
            assert cell.fill.fill_type == "solid"

    def test_label_column_is_bold_and_filled(self, painted):
        for row in (3, 4):
            cell = _cell(painted, row, 1)
            assert cell.font.bold is True
            assert cell.fill.fill_type == "solid"

    def test_data_cells_are_not_bold(self, painted):
        for row in (3, 4):
            for col in (2, 3):
                assert _cell(painted, row, col).font.bold is False

    def test_data_rows_alternate_fill(self, painted):
        # local r=1 (row 3, first data row) unbanded, r=2 (row 4) banded
        assert _cell(painted, 3, 2).fill.fill_type is None
        assert _cell(painted, 4, 2).fill.fill_type == "solid"
        assert _cell(painted, 4, 2).fill.fgColor.rgb == "00F2F2F2"

    def test_outer_border_is_thick(self, painted):
        assert _cell(painted, 2, 1).border.top.style == "thick"
        assert _cell(painted, 2, 1).border.left.style == "thick"
        assert _cell(painted, 2, 3).border.top.style == "thick"
        assert _cell(painted, 2, 3).border.right.style == "thick"
        assert _cell(painted, 4, 1).border.bottom.style == "thick"
        assert _cell(painted, 4, 1).border.left.style == "thick"
        assert _cell(painted, 4, 3).border.bottom.style == "thick"
        assert _cell(painted, 4, 3).border.right.style == "thick"

    def test_header_data_boundary_is_medium(self, painted):
        assert _cell(painted, 2, 2).border.bottom.style == "medium"
        assert _cell(painted, 3, 2).border.top.style == "medium"

    def test_label_data_boundary_is_medium(self, painted):
        assert _cell(painted, 3, 1).border.right.style == "medium"
        assert _cell(painted, 3, 2).border.left.style == "medium"

    def test_inner_grid_is_thin(self, painted):
        assert _cell(painted, 3, 2).border.bottom.style == "thin"
        assert _cell(painted, 4, 2).border.top.style == "thin"
        assert _cell(painted, 3, 2).border.right.style == "thin"
        assert _cell(painted, 3, 3).border.left.style == "thin"

    def test_marker_row_is_untouched(self, painted):
        marker = _cell(painted, 1, 1)
        assert marker.font.bold is False
        # an untouched cell's default border is Side(style=None), not a bare
        # None — both mean "no visible line", so accept either
        assert marker.border.top is None or marker.border.top.style is None


class TestAutoExtend:
    def test_add_column_extends_styling(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T")
        t.create(data_sheet, sheet="Data")
        t.add_column("b", [2])
        t.write(data_sheet)

        new_header = _cell(data_sheet, 2, 3)  # column "b"
        assert new_header.font.bold is True
        assert new_header.fill.fill_type == "solid"
        assert new_header.border.right.style == "thick"  # new outer edge

        old_header = _cell(data_sheet, 2, 2)  # column "a", no longer the edge
        assert old_header.border.right.style == "thin"

    def test_add_row_extends_styling(self, data_sheet):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T")
        t.create(data_sheet, sheet="Data")
        t.add_row("y", [2])
        t.write(data_sheet)

        new_label = _cell(data_sheet, 4, 1)  # row "y"
        assert new_label.font.bold is True
        assert new_label.border.bottom.style == "thick"  # new outer edge

        old_label_row = _cell(data_sheet, 3, 1)  # row "x", no longer the edge
        assert old_label_row.border.bottom.style == "thin"

    def test_banding_recomputed_after_row_removed(self, data_sheet):
        t = Table(
            data=[[1], [2], [3]],
            column_headers=["a"],
            row_labels=["x", "y", "z"],
            name="T",
        )
        t.create(data_sheet, sheet="Data")
        t.drop_row("x")
        t.write(data_sheet)
        # "y" is now the first data row (local r=1, unbanded)
        assert _cell(data_sheet, 3, 2).fill.fill_type is None


class TestShrinkLeavesNoGhostFormatting:
    """A painted cell that falls outside a table's new, smaller size must be
    truly cleared — not just value-less but still carrying its old style,
    which would otherwise keep inflating the sheet's saved dimensions
    (openpyxl/Excel both treat a styled-but-empty cell as "in use")."""

    def test_dropped_column_leaves_no_residual_width(self, data_sheet):
        t = Table(data=[[1, 2]], column_headers=["a", "b"], row_labels=["x"], name="T")
        t.create(data_sheet, sheet="Data")
        t.drop_column("b")
        t.write(data_sheet)
        assert read_sheet(data_sheet, "Data") == [["TABLE NAME", "T"], [None, "a"], ["x", 1]]

    def test_deleted_table_leaves_a_genuinely_empty_sheet(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        delete_table(data_sheet, "T")
        assert read_sheet(data_sheet, "Data") == []


class TestPersistence:
    def test_read_restores_style(self, data_sheet):
        Table(
            data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style=TableStyle.MINIMAL
        ).create(data_sheet, sheet="Data")
        t = Table.read(data_sheet, "T")
        assert t.style == TableStyle.MINIMAL

    def test_changing_style_and_writing_repaints_and_persists(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        t = Table.read(data_sheet, "T")
        t.style = TableStyle.NONE
        t.write(data_sheet)

        assert _cell(data_sheet, 2, 1).font.bold is False
        assert Table.read(data_sheet, "T").style == TableStyle.NONE

    def test_style_setter_rejects_non_table_style(self):
        t = Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T")
        with pytest.raises(TypeError):
            t.style = "fancy"

    def test_constructor_rejects_non_table_style(self):
        with pytest.raises(TypeError):
            Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style="fancy")

    def test_style_is_excluded_from_equality(self, data_sheet):
        a = Table(
            data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style=TableStyle.DEFAULT
        )
        b = Table(
            data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style=TableStyle.NONE
        )
        assert a == b


class TestShiftPreservesStyle:
    def test_shifted_neighbour_keeps_its_style(self, data_sheet):
        sales = Table(
            data=[[100, 200]], column_headers=["North", "South"], row_labels=["Q1"], name="Sales"
        )
        sales.create(data_sheet, sheet="Data")
        inventory = Table(
            data=[[10]],
            column_headers=["Units"],
            row_labels=["A"],
            name="Inventory",
            style=TableStyle.MINIMAL,
        )
        inventory.create(data_sheet, sheet="Data")

        sales.add_column("East", [300])
        sales.write(data_sheet)

        moved = Table.read(data_sheet, "Inventory")
        assert moved.style == TableStyle.MINIMAL
        assert moved.data.rows == [[10]]

        wb = load_workbook(data_sheet)
        ws = wb["Data"]
        # Inventory's header ("Units") is now one column further right
        header_cells = [c for c in ws[2] if c.value == "Units"]
        assert len(header_cells) == 1
        assert header_cells[0].font.bold is True
