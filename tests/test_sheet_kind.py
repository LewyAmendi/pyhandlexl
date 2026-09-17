"""Tests for the table-vs-grid sheet invariant: sheet_kind, SheetKindError,
clear_all_sheet_data, and the schema staying in sync through delete/rename."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    SheetKindError,
    SheetNotFoundError,
    Table,
    TableNotFoundError,
    append_rows,
    clear_all_sheet_data,
    create_sheet,
    delete_sheet,
    list_sheets,
    list_tables,
    rename_sheet,
    sheet_kind,
    write_sheet,
)
from pyhandlexl._multi_table import SCHEMA_SHEET


def _sales(name="Sales"):
    return Table(data=[[1, 2]], column_headers=["a", "b"], row_labels=["r"], name=name)


class TestSheetKind:
    def test_fresh_sheet_is_empty(self, book):
        create_sheet(book, "Data")
        assert sheet_kind(book, "Data") == "empty"

    def test_grid_write_claims_it(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a", "b"]], sheet="Data")
        assert sheet_kind(book, "Data") == "grid"

    def test_table_create_claims_it(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        assert sheet_kind(book, "Data") == "table"

    def test_unknown_sheet_raises(self, book):
        with pytest.raises(SheetNotFoundError):
            sheet_kind(book, "Ghost")

    def test_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            sheet_kind(book, SCHEMA_SHEET)


class TestGridRefusesTableSheet:
    def test_write_sheet_on_table_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            write_sheet(book, [["x"]], sheet="Data")

    def test_append_rows_on_table_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            append_rows(book, [["x"]], sheet="Data")

    def test_write_sheet_does_not_corrupt_the_table_on_failure(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            write_sheet(book, [["x"]], sheet="Data")
        assert Table.read(book, "Sales").data.rows == [[1, 2]]


class TestTableRefusesGridSheet:
    def test_table_create_on_grid_sheet_raises(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a", "b"]], sheet="Data")
        with pytest.raises(SheetKindError):
            _sales().create(book, sheet="Data")

    def test_failure_does_not_leave_a_partial_table(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a", "b"]], sheet="Data")
        with pytest.raises(SheetKindError):
            _sales().create(book, sheet="Data")
        assert list_tables(book) == []
        assert sheet_kind(book, "Data") == "grid"  # still a plain grid sheet, untouched


class TestSchemaSheetIsReserved:
    def test_write_sheet_on_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            write_sheet(book, [["x"]], sheet=SCHEMA_SHEET)

    def test_append_rows_on_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            append_rows(book, [["x"]], sheet=SCHEMA_SHEET)

    def test_table_create_on_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="Other").create(
                book, sheet=SCHEMA_SHEET
            )

    def test_clear_all_sheet_data_on_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            clear_all_sheet_data(book, SCHEMA_SHEET)

    def test_delete_sheet_on_schema_sheet_raises(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        with pytest.raises(SheetKindError):
            delete_sheet(book, SCHEMA_SHEET)
        assert SCHEMA_SHEET in list(load_workbook(book).sheetnames)

    def test_delete_sheet_on_schema_sheet_raises_even_with_no_tables(self, book):
        # reserved and blocked outright, regardless of whether it currently
        # exists — same precedent as write_sheet/Table.create/
        # clear_all_sheet_data checking the reserved name before existence.
        with pytest.raises(SheetKindError):
            delete_sheet(book, SCHEMA_SHEET)


class TestClearAllSheetData:
    def test_unknown_sheet_raises(self, book):
        with pytest.raises(SheetNotFoundError):
            clear_all_sheet_data(book, "Ghost")

    def test_clears_grid_sheet_back_to_empty(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a", "b"]], sheet="Data")
        clear_all_sheet_data(book, "Data")
        assert sheet_kind(book, "Data") == "empty"

    def test_grid_sheet_can_then_hold_a_table(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a", "b"]], sheet="Data")
        clear_all_sheet_data(book, "Data")
        _sales().create(book, sheet="Data")
        assert sheet_kind(book, "Data") == "table"

    def test_clears_table_sheet_back_to_empty(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        clear_all_sheet_data(book, "Data")
        assert sheet_kind(book, "Data") == "empty"
        assert list_tables(book) == []

    def test_table_sheet_can_then_hold_a_grid(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        clear_all_sheet_data(book, "Data")
        write_sheet(book, [["x", "y"]], sheet="Data")
        assert sheet_kind(book, "Data") == "grid"

    def test_cleared_table_is_no_longer_readable(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        clear_all_sheet_data(book, "Data")
        with pytest.raises(TableNotFoundError):
            Table.read(book, "Sales")

    def test_only_the_named_sheet_is_cleared(self, book):
        create_sheet(book, "One")
        create_sheet(book, "Two")
        _sales("A").create(book, sheet="One")
        _sales("B").create(book, sheet="Two")
        clear_all_sheet_data(book, "One")
        assert list_tables(book) == ["B"]
        assert Table.read(book, "B").data.rows == [[1, 2]]

    def test_position_among_sheets_is_preserved(self, book):
        create_sheet(book, "First")
        create_sheet(book, "Data")
        create_sheet(book, "Last")
        write_sheet(book, [["a"]], sheet="Data")
        clear_all_sheet_data(book, "Data")
        assert list_sheets(book) == ["Sheet", "First", "Data", "Last"]

    def test_ghost_formatting_is_gone_after_clearing_a_table_sheet(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        clear_all_sheet_data(book, "Data")
        write_sheet(book, [["a"]], sheet="Data")
        ws = load_workbook(book)["Data"]
        assert ws.max_row == 1
        assert ws.max_column == 1


class TestDeleteSheetPurgesSchema:
    def test_deleting_a_table_sheet_forgets_its_tables(self, book):
        create_sheet(book, "One")
        create_sheet(book, "Two")
        _sales("A").create(book, sheet="One")
        _sales("B").create(book, sheet="Two")
        delete_sheet(book, "One")
        assert list_tables(book) == ["B"]
        with pytest.raises(TableNotFoundError):
            Table.read(book, "A")

    def test_recreating_the_sheet_afterwards_is_a_fresh_empty_one(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        delete_sheet(book, "Data")
        create_sheet(book, "Data")
        assert sheet_kind(book, "Data") == "empty"

    def test_deleting_a_grid_sheet_is_unaffected(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a"]], sheet="Data")
        delete_sheet(book, "Data")
        create_sheet(book, "Data")
        assert sheet_kind(book, "Data") == "empty"


class TestRenameSheetUpdatesSchema:
    def test_renaming_a_table_sheet_keeps_tables_readable(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        rename_sheet(book, "Data", "Renamed")
        assert Table.read(book, "Sales").data.rows == [[1, 2]]
        assert sheet_kind(book, "Renamed") == "table"

    def test_renaming_a_grid_sheet_is_unaffected(self, book):
        create_sheet(book, "Data")
        write_sheet(book, [["a"]], sheet="Data")
        rename_sheet(book, "Data", "Renamed")
        assert sheet_kind(book, "Renamed") == "grid"

    def test_renaming_to_the_same_name_is_a_noop(self, book):
        create_sheet(book, "Data")
        _sales().create(book, sheet="Data")
        rename_sheet(book, "Data", "Data")
        assert Table.read(book, "Sales").data.rows == [[1, 2]]
