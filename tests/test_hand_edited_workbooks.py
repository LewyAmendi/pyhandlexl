"""Workbooks somebody edited by hand in Excel: what still reads, what fails, and how clearly."""

from __future__ import annotations

import datetime as dt
import shutil
import warnings

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    ColumnType,
    ColumnTypeError,
    FormulaWarning,
    SchemaRebuiltWarning,
    Table,
    TableNotFoundError,
    create_sheet,
    list_tables,
    sheet_kind,
    table_info,
    write_sheet,
)
from pyhandlexl._multi_table import SCHEMA_SHEET


@pytest.fixture
def path(book):
    create_sheet(book, "D")
    Table(
        data=[[1, 2], [3, 4]],
        row_labels=["r1", "r2"],
        column_headers=["a", "b"],
        corner="c",
        name="T",
    ).create(book, sheet="D")
    return book


# The table sits at: A1 marker | B1 name; A2 corner, B2:C2 headers; A3:A4 labels, B3:C4 data.


def edit(path, mutate) -> None:
    wb = load_workbook(path)
    mutate(wb["D"])
    wb.save(path)


def drop_schema(path) -> None:
    wb = load_workbook(path)
    del wb[SCHEMA_SHEET]
    wb.save(path)


class TestValuesPeopleTypeIn:
    def test_a_numeric_header_or_label_is_read_as_text(self, path):
        def go(ws):
            ws["B2"] = 2024
            ws["A3"] = 7
            ws["C2"] = 3.5

        edit(path, go)
        table = Table.read(path, "T")
        assert table.data.column_headers == ["2024", "3.5"]
        assert table.data.row_labels == ["7", "r2"]

    def test_a_date_header_is_read_as_its_text(self, path):
        edit(path, lambda ws: ws.__setitem__("B2", dt.datetime(2026, 1, 1)))
        assert Table.read(path, "T").data.column_headers == ["2026-01-01 00:00:00", "b"]

    def test_a_blank_data_cell_is_none(self, path):
        edit(path, lambda ws: ws.__setitem__("B3", None))
        assert Table.read(path, "T").data.rows == [[None, 2], [3, 4]]

    def test_a_formula_reads_as_its_text_with_a_warning_and_becomes_text_when_written_back(
        self, path
    ):
        edit(path, lambda ws: ws.__setitem__("B3", "=C3*2"))
        with pytest.warns(FormulaWarning, match="table 'T' holds 1 formula cell"):
            table = Table.read(path, "T")
        assert table.read_cell(row="r1", column="a") == "=C3*2"
        table.set_cell(row="r2", column="a", value=99)
        table.write(path)
        cell = load_workbook(path)["D"]["B3"]
        assert cell.value == "=C3*2" and cell.data_type == "s"  # pyhandlexl has no formulas

    def test_an_error_value_reads_as_its_text(self, path):
        edit(path, lambda ws: ws.__setitem__("B3", "#DIV/0!"))
        assert Table.read(path, "T").read_cell(row="r1", column="a") == "#DIV/0!"

    def test_duplicate_labels_and_headers_are_tolerated_on_read(self, path):
        def go(ws):
            ws["A4"] = "r1"
            ws["C2"] = "a"

        edit(path, go)
        table = Table.read(path, "T")
        assert table.data.row_labels == ["r1", "r1"] and table.data.column_headers == ["a", "a"]
        assert table.read_row("r1") == [1, 2]  # the first match


class TestBlankHeadersAndLabelsAreReportedByCell:
    def test_a_blank_column_header(self, path):
        edit(path, lambda ws: ws.__setitem__("C2", None))
        with pytest.raises(ValueError, match=r"header in cell C2 is blank") as caught:
            Table.read(path, "T")
        assert "'T'" in str(caught.value) and "'D'" in str(caught.value)

    def test_a_blank_row_label(self, path):
        edit(path, lambda ws: ws.__setitem__("A4", None))
        with pytest.raises(ValueError, match=r"label in cell A4 is blank"):
            Table.read(path, "T")

    def test_the_first_blank_one_is_named(self, path):
        def go(ws):
            ws["B2"] = None
            ws["C2"] = None

        edit(path, go)
        with pytest.raises(ValueError, match="cell B2"):
            Table.read(path, "T")

    def test_a_row_deleted_in_excel_leaves_a_blank_label_that_is_reported(self, path):
        edit(path, lambda ws: ws.delete_rows(3))
        with pytest.raises(ValueError, match="blank"):
            Table.read(path, "T")

    def test_a_blank_header_does_not_stop_other_tables_or_the_workbook(self, path):
        Table(data=[[9]], row_labels=["x"], column_headers=["y"], name="U").create(path, sheet="D")
        edit(path, lambda ws: ws.__setitem__("C2", None))
        with pytest.raises(ValueError):
            Table.read(path, "T")
        assert Table.read(path, "U").read_row("x") == [9]
        assert sorted(list_tables(path)) == ["T", "U"]


class TestMarkersThatWereMovedOrDamaged:
    def test_the_marker_name_edited_means_not_found(self, path):
        edit(path, lambda ws: ws.__setitem__("B1", "Renamed"))
        with pytest.raises(TableNotFoundError, match="marker is missing"):
            Table.read(path, "T")

    def test_the_marker_deleted_means_not_found(self, path):
        edit(path, lambda ws: ws.__setitem__("A1", None))
        with pytest.raises(TableNotFoundError):
            Table.read(path, "T")

    @pytest.mark.parametrize(
        "mutate",
        [
            lambda ws: ws.insert_rows(1),
            lambda ws: ws.insert_rows(1, amount=5),
            lambda ws: ws.insert_cols(1),
            lambda ws: ws.insert_cols(1, amount=4),
        ],
        ids=["one row above", "five rows above", "one column left", "four columns left"],
    )
    def test_rows_or_columns_inserted_above_or_left_are_healed_and_saved(self, path, mutate):
        edit(path, mutate)
        table = Table.read(path, "T")
        assert table.data.rows == [[1, 2], [3, 4]] and table.data.row_labels == ["r1", "r2"]
        table.set_cell(row="r1", column="a", value=100)
        table.write(path)
        assert Table.read(path, "T").read_cell(row="r1", column="a") == 100
        # and the healed position was persisted, so reading again needs no scan
        assert table_info(path, "T").n_rows == 2

    def test_a_column_deleted_through_the_marker_means_not_found(self, path):
        edit(path, lambda ws: ws.delete_cols(2))
        with pytest.raises(TableNotFoundError):
            Table.read(path, "T")

    def test_stray_cells_beside_or_below_a_table_are_not_part_of_it(self, path):
        def go(ws):
            ws["E3"] = "stray"
            ws["B9"] = "stray"
            ws["A7"] = "stray"

        edit(path, go)
        assert Table.read(path, "T").data.rows == [[1, 2], [3, 4]]
        drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert Table.read(path, "T").data.rows == [[1, 2], [3, 4]]

    def test_a_moved_table_is_found_again_by_its_marker(self, path):
        Table(data=[[9]], row_labels=["x"], column_headers=["y"], name="U").create(path, sheet="D")
        edit(path, lambda ws: ws.move_range("E1:F3", rows=20))  # U (at E:F) moves down 20 rows
        drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert sorted(list_tables(path)) == ["T", "U"]
        assert Table.read(path, "U").read_row("x") == [9]
        assert Table.read(path, "T").data.rows == [[1, 2], [3, 4]]


class TestTheWordsTableNameInsideData:
    """A cell reading TABLE NAME must not be mistaken for a table when the schema is rebuilt."""

    def _make(self, path, **kwargs):
        kwargs.setdefault("name", "Tricky")
        Table(**kwargs).create(path, sheet="D")
        drop_schema(path)

    def test_a_row_label_and_a_data_value_that_say_it(self, path):
        self._make(
            path,
            data=[["TABLE NAME"], ["x"]],
            row_labels=["TABLE NAME", "r2"],
            column_headers=["v"],
        )
        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(path)
        assert sorted(names) == ["T", "Tricky"]  # no phantom "TABLE NAME" table
        tricky = Table.read(path, "Tricky")
        assert tricky.data.rows == [["TABLE NAME"], ["x"]]
        assert tricky.data.row_labels == ["TABLE NAME", "r2"]

    def test_a_column_header_and_a_corner_that_say_it(self, path):
        self._make(
            path,
            data=[[1, 2]],
            row_labels=["r"],
            column_headers=["TABLE NAME", "other"],
            corner="TABLE NAME",
        )
        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(path)
        assert sorted(names) == ["T", "Tricky"]
        back = Table.read(path, "Tricky")
        assert back.data.column_headers == ["TABLE NAME", "other"]
        assert back.data.rows == [[1, 2]]

    def test_the_real_tables_widths_are_not_cut_short_by_the_imposter(self, path):
        self._make(
            path,
            data=[["TABLE NAME", 1, 2, 3]],
            row_labels=["TABLE NAME"],
            column_headers=["a", "b", "c", "d"],
        )
        with pytest.warns(SchemaRebuiltWarning):
            list_tables(path)
        assert table_info(path, "Tricky").n_cols == 4

    def test_a_table_can_be_named_table_name(self, path):
        Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="TABLE NAME").create(
            path, sheet="D"
        )
        assert Table.read(path, "TABLE NAME").read_row("r") == [1]
        drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert sorted(list_tables(path)) == ["T", "TABLE NAME"]

    def test_known_limit_a_grid_sheet_that_starts_with_the_words_looks_like_a_table(self, path):
        # A grid sheet whose very first cells read "TABLE NAME" | <text> is indistinguishable
        # from a marker once the schema is gone. Documented in the README; pinned here.
        write_sheet(path, [["TABLE NAME", "hello"], ["a", "b"]], sheet="Grid")
        drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert "hello" in list_tables(path)
        assert sheet_kind(path, "Grid") == "table"


WIDE = 5_600  # the readable form of the type list needs ~34,000 characters here


@pytest.fixture(scope="module")
def _wide_master(tmp_path_factory):
    """One workbook holding a WIDE-column table, built once for the whole module."""
    from pyhandlexl import create_workbook

    master = tmp_path_factory.mktemp("wide") / "master.xlsx"
    create_workbook(master)
    create_sheet(master, "W")
    Table(
        data=[list(range(WIDE))],
        row_labels=["r"],
        column_headers=[f"c{i}" for i in range(WIDE)],
        name="Wide",
    ).create(master, sheet="W")
    return master


@pytest.fixture
def wide_book(_wide_master, tmp_path):
    copy = tmp_path / "wide.xlsx"
    shutil.copy(_wide_master, copy)
    return copy


class TestVeryWideTables:
    """A table's column types are stored in one cell; too many columns used to overflow it."""

    WIDE = WIDE

    def _wide(self, path, n=WIDE, types=None):
        create_sheet(path, "W")
        headers = [f"c{i}" for i in range(n)]
        Table(
            data=[list(range(n))],
            row_labels=["r"],
            column_headers=headers,
            name="Wide",
            column_types=types,
        ).create(path, sheet="W")
        return headers

    def test_creating_and_reading_a_table_wider_than_one_cell_can_describe(self, wide_book):
        book = wide_book
        table = Table.read(book, "Wide")
        assert len(table.data.column_headers) == self.WIDE
        assert table.read_cell(row="r", column="c5599") == 5599
        assert list_tables(book) == ["Wide"]

    def test_every_other_table_and_operation_still_works_beside_it(self, wide_book):
        book = wide_book
        create_sheet(book, "S")
        Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="Small").create(
            book, sheet="S"
        )
        assert sorted(list_tables(book)) == ["Small", "Wide"]
        table_info(book, "Wide")
        write_sheet(book, [["g"]], sheet="Grid")
        assert sheet_kind(book, "S") == "table"

    def test_the_column_types_survive_the_compact_form(self, book):
        headers = [f"c{i}" for i in range(self.WIDE)]
        types = {
            headers[0]: ColumnType.NUMBER,
            headers[1]: ColumnType.TEXT,
            headers[2]: ColumnType.BOOLEAN,
            headers[3]: ColumnType.DATE,
            headers[4]: ColumnType.TIME,
            headers[5]: ColumnType.DURATION,
            headers[-1]: ColumnType.NUMBER,
        }
        create_sheet(book, "W")
        data = [[None] * self.WIDE]
        data[0][0], data[0][1], data[0][2] = 1, "x", True
        data[0][3], data[0][4], data[0][5] = dt.datetime(2026, 1, 1), dt.time(1, 0), dt.timedelta(1)
        data[0][-1] = 7
        Table(
            data=data, row_labels=["r"], column_headers=headers, name="Wide", column_types=types
        ).create(book, sheet="W")
        back = Table.read(book, "Wide").column_types
        for header, expected in types.items():
            assert back[header] == expected
        assert back["c100"] == ColumnType.ANY
        # and a restriction is still enforced after the round trip — immediately, on the edit
        table = Table.read(book, "Wide")
        with pytest.raises(ColumnTypeError):
            table.set_cell(row="r", column="c0", value="not a number")

    def test_the_schema_cell_never_exceeds_what_a_cell_can_hold(self, wide_book):
        book = wide_book
        cells = [c.value for row in load_workbook(book)[SCHEMA_SHEET].iter_rows() for c in row]
        assert all(len(v) <= 32_767 for v in cells if isinstance(v, str))

    def test_writing_a_wide_table_back_keeps_working(self, wide_book):
        book = wide_book
        table = Table.read(book, "Wide")
        table.add_row("r2", [1] * self.WIDE)
        table.write(book)
        assert Table.read(book, "Wide").data.row_labels == ["r", "r2"]

    def test_the_widest_table_excel_allows(self, book):
        self._wide(book, n=16_383)  # 16,383 data columns + the label column = 16,384
        assert table_info(book, "Wide").n_cols == 16_383
        assert Table.read(book, "Wide").read_cell(row="r", column="c16382") == 16_382

    def test_a_wide_table_survives_a_schema_rebuild(self, wide_book):
        book = wide_book
        drop_schema(book)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", SchemaRebuiltWarning)
            assert list_tables(book) == ["Wide"]
            assert table_info(book, "Wide").n_cols == self.WIDE


class TestColumnTypeDriftFromHandEditing:
    """A restriction doesn't stop an edit made directly in Excel — pyhandlexl never saw it —
    so a read has to succeed regardless; only the next write() enforces it again."""

    def test_reading_a_value_that_no_longer_fits_its_column_type(self, book):
        create_sheet(book, "D")
        Table(
            data=[[1], [2]],
            row_labels=["r1", "r2"],
            column_headers=["a"],
            name="T",
            column_types={"a": ColumnType.NUMBER},
        ).create(book, sheet="D")
        edit(book, lambda ws: ws.__setitem__("B3", "not a number"))  # column a, row r1
        t = Table.read(book, "T")
        assert t.read_cell(row="r1", column="a") == "not a number"
        assert t.column_types == {"a": ColumnType.NUMBER}
        with pytest.raises(ColumnTypeError):
            t.write(book)  # the restriction bites again the moment it's written back
