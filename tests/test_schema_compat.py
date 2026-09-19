"""Workbooks written by older pyhandlexl versions, and reads that repair the schema."""

from __future__ import annotations

import warnings

import pytest
from openpyxl import load_workbook

import pyhandlexl._multi_table as mt
from pyhandlexl import (
    ColumnType,
    SchemaRebuiltWarning,
    Table,
    TableStyle,
    create_sheet,
    list_tables,
    table_info,
)

FULL_WIDTH = (
    10  # name, sheet, anchor_row, anchor_col, n_rows, n_cols, style, column_types, +2 dates
)


@pytest.fixture
def path(book):
    create_sheet(book, "Data")
    Table(
        data=[[1, "x"], [2, "y"]],
        row_labels=["r1", "r2"],
        column_headers=["n", "s"],
        name="T",
        style=TableStyle.MINIMAL,
        column_types={"n": ColumnType.NUMBER, "s": ColumnType.TEXT},
    ).create(book, sheet="Data")
    return book


def truncate_schema(path, width):
    """Rewrite the schema sheet the way an older version would have: only *width* columns."""
    wb = load_workbook(path)
    ws = wb[mt.SCHEMA_SHEET]
    rows = [[cell.value for cell in row][:width] for row in ws.iter_rows()]
    ws.delete_rows(1, ws.max_row)
    for row in rows:
        ws.append(row)
    wb.save(path)


def schema_width(path):
    ws = load_workbook(path)[mt.SCHEMA_SHEET]
    return max(len([c for c in row if c is not None]) for row in ws.iter_rows(values_only=True))


class TestOlderSchemaFormats:
    @pytest.mark.parametrize("width", [6, 7, 8, 9])
    def test_they_load_instead_of_raising(self, path, width):
        truncate_schema(path, width)
        assert list_tables(path) == ["T"]
        t = Table.read(path, "T")
        assert t.data.row_labels == ["r1", "r2"]
        assert t.read_row("r2") == [2, "y"]

    def test_seven_columns_have_style_but_no_types_or_dates(self, path):
        truncate_schema(path, 7)
        t = Table.read(path, "T")
        assert t.style == TableStyle.MINIMAL
        assert t.column_types == {"n": ColumnType.ANY, "s": ColumnType.ANY}
        info = table_info(path, "T")
        assert info.created_at is None and info.modified_at is None

    def test_eight_columns_have_types_but_no_dates(self, path):
        truncate_schema(path, 8)
        t = Table.read(path, "T")
        assert t.column_types == {"n": ColumnType.NUMBER, "s": ColumnType.TEXT}
        assert table_info(path, "T").created_at is None

    def test_six_columns_fall_back_to_the_default_style(self, path):
        truncate_schema(path, 6)
        assert Table.read(path, "T").style == TableStyle.DEFAULT

    def test_reading_an_old_schema_does_not_modify_the_file(self, path):
        truncate_schema(path, 7)
        before = path.read_bytes()
        Table.read(path, "T")
        list_tables(path)
        table_info(path, "T")
        assert path.read_bytes() == before

    @pytest.mark.parametrize("width", [6, 7, 8, 9])
    def test_the_next_write_upgrades_the_schema_in_place(self, path, width):
        truncate_schema(path, width)
        t = Table.read(path, "T")
        t.set_cell(row="r1", column="n", value=100)
        t.write(path)
        assert schema_width(path) == FULL_WIDTH
        again = Table.read(path, "T")
        assert again.read_cell(row="r1", column="n") == 100
        assert again.style == t.style
        assert table_info(path, "T").modified_at is not None

    def test_a_write_from_an_old_schema_still_merges_with_another_writer(self, path):
        truncate_schema(path, 7)
        alice, bob = Table.read(path, "T"), Table.read(path, "T")
        alice.add_row("alice", [1, "a"])
        bob.add_row("bob", [2, "b"])
        alice.write(path)
        bob.write(path)
        assert Table.read(path, "T").data.row_labels == ["r1", "r2", "alice", "bob"]

    def test_several_tables_in_an_old_schema(self, path):
        Table(data=[[1]], row_labels=["k"], column_headers=["v"], name="U").create(
            path, sheet="Data"
        )
        truncate_schema(path, 7)
        assert sorted(list_tables(path)) == ["T", "U"]
        assert Table.read(path, "U").read_cell(row="k", column="v") == 1

    def test_blank_style_and_type_cells_are_treated_like_missing_ones(self, path):
        wb = load_workbook(path)
        ws = wb[mt.SCHEMA_SHEET]
        ws["G2"], ws["H2"] = None, None
        wb.save(path)
        t = Table.read(path, "T")
        assert t.style == TableStyle.DEFAULT
        assert set(t.column_types.values()) == {ColumnType.ANY}


class TestReadsPersistTheirRepairs:
    """A read that rebuilds the schema or heals a table's position saves that repair."""

    def _drop_schema(self, path):
        wb = load_workbook(path)
        del wb[mt.SCHEMA_SHEET]
        wb.save(path)

    def test_a_read_of_a_healthy_workbook_leaves_the_file_byte_for_byte_alone(self, path):
        before = path.read_bytes()
        Table.read(path, "T")
        list_tables(path)
        table_info(path, "T")
        assert path.read_bytes() == before

    def test_a_missing_schema_is_rebuilt_and_saved_by_a_read(self, path):
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            Table.read(path, "T")
        assert mt.SCHEMA_SHEET in load_workbook(path).sheetnames

    def test_one_rebuild_emits_exactly_one_warning(self, path):
        self._drop_schema(path)
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            Table.read(path, "T")
        assert [w.category for w in caught] == [SchemaRebuiltWarning]

    def test_the_second_read_finds_a_schema_and_stays_quiet(self, path):
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            Table.read(path, "T")
        settled = path.read_bytes()
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            Table.read(path, "T")
        assert path.read_bytes() == settled

    def test_list_tables_and_table_info_persist_a_rebuild_too(self, path):
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert list_tables(path) == ["T"]
        assert mt.SCHEMA_SHEET in load_workbook(path).sheetnames
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            assert table_info(path, "T").n_rows == 2
        assert mt.SCHEMA_SHEET in load_workbook(path).sheetnames

    def test_a_relocated_table_position_is_healed_and_saved(self, path):
        wb = load_workbook(path)
        wb[mt.SCHEMA_SHEET]["D2"] = 30  # the schema now claims the table is far to the right
        wb.save(path)
        assert Table.read(path, "T").read_row("r1") == [1, "x"]
        assert load_workbook(path)[mt.SCHEMA_SHEET]["D2"].value != 30

    def test_a_workbook_with_no_tables_is_not_given_a_schema_sheet_by_a_read(self, book):
        assert list_tables(book) == []
        assert mt.SCHEMA_SHEET not in load_workbook(book).sheetnames

    def test_the_rebuilt_schema_reads_back_the_same_table(self, path):
        expected = Table.read(path, "T")
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            rebuilt = Table.read(path, "T")
        assert rebuilt == expected

    def test_a_rebuilt_table_can_still_merge_with_another_writer(self, path):
        self._drop_schema(path)
        with pytest.warns(SchemaRebuiltWarning):
            alice = Table.read(path, "T")
        bob = Table.read(path, "T")
        alice.add_row("alice", [1, "a"])
        bob.add_row("bob", [2, "b"])
        alice.write(path)
        bob.write(path)
        assert Table.read(path, "T").data.row_labels == ["r1", "r2", "alice", "bob"]
