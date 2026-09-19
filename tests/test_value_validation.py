"""Values and names Excel can't hold must be refused up front — never written, never mangled.

Each of these used to be accepted silently and then either lost data (a long string cut
short, ``nan`` turned blank, a date shifted) or, worst, produced a workbook that could no
longer be opened at all.
"""

from __future__ import annotations

import datetime as dt
from typing import ClassVar

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    CellTypeError,
    SheetNameError,
    Table,
    append_rows,
    check_cell_value,
    check_sheet_name,
    create_sheet,
    import_csv_to_xl,
    list_sheets,
    read_sheet,
    rename_sheet,
    write_sheet,
)

# Strings XML (and so .xlsx) cannot hold: any one of these in any cell breaks the whole file.
UNSTORABLE_STRINGS = [
    "a\x00b",
    "a\x01b",
    "a\x08b",
    "a\x0bb",
    "a\x0cb",
    "a\x0eb",
    "a\x1fb",
    "a\ud800b",  # a lone surrogate
    "a\udfffb",
    "a￾b",  # non-characters
    "a￿b",
    "x" * 32_768,  # one past what a cell holds
]

# Odd but perfectly storable strings.
STORABLE_STRINGS = [
    "",
    " ",
    "  padded  ",
    "tab\there",
    "line1\nline2",
    "\x7f",  # DEL is legal XML
    "\x85",  # NEL
    "\N{LINE SEPARATOR}\N{PARAGRAPH SEPARATOR}",
    "\N{ZERO WIDTH NO-BREAK SPACE}",  # a BOM in the middle of text
    "é ü ñ",
    "日本語",
    "😀 \U0001f468‍\U0001f469‍\U0001f467",
    "\U0010ffff",  # the highest code point XML allows
    "x" * 32_767,  # exactly what a cell holds
    "a,b;c\"d'e",
    "TRUE",
    "007",
    "1e3",
    "<tag>&amp;</tag>",
    "]]>",
    "'leading apostrophe",
]

UNSTORABLE_VALUES = [
    float("nan"),
    float("inf"),
    float("-inf"),
    10**400,
    -(10**400),
    dt.datetime(1899, 12, 30),
    dt.datetime(1899, 12, 31, 23, 59, 59),
    dt.datetime.min,
    dt.datetime.max,
    dt.datetime(9999, 12, 31, 23, 59, 59, 999999),
    dt.date(1899, 12, 31),
    dt.date.min,
    dt.time(1, 2, tzinfo=dt.timezone.utc),
    dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc),
    b"bytes",
    bytearray(b"x"),
    1j,
    [1],
    {"a": 1},
    object(),
]

STORABLE_VALUES = [
    None,
    True,
    False,
    0,
    -1,
    2**53,
    2**60,
    -0.0,
    1e308,
    5e-324,
    1e-10,
    dt.datetime(1900, 1, 1),
    dt.datetime(1900, 3, 1),
    dt.datetime(2026, 2, 28, 23, 59, 59),
    dt.datetime(9999, 12, 31, 23, 59, 59, 999000),
    dt.date(1900, 1, 1),
    dt.date(9999, 12, 31),
    dt.time(0, 0),
    dt.time(23, 59, 59, 999000),
    dt.timedelta(0),
    dt.timedelta(days=-1),
    dt.timedelta(days=40_000),
]


# Each test that goes through a whole write builds a workbook, so those use a representative
# few (a control character, a lone surrogate, a non-character, an over-long string; nan, inf,
# a huge int, an out-of-range date, a timezone, a foreign type). The unit tests above and
# below use every one.
API_BAD_STRINGS = [UNSTORABLE_STRINGS[i] for i in (0, 7, 9, 11)]
API_BAD_VALUES = [
    float("nan"),
    float("inf"),
    10**400,
    dt.datetime.min,
    dt.time(1, 2, tzinfo=dt.timezone.utc),
    object(),
]
API_BAD = API_BAD_STRINGS + API_BAD_VALUES


def _short(value) -> str:
    """A short test id — a 32,768-character string as an id would bury the report."""
    text = repr(value)
    return text if len(text) <= 40 else f"{text[:20]}...({len(text)} chars)"


def _hidden_temp_files(directory) -> list[str]:
    return sorted(p.name for p in directory.iterdir() if ".tmp." in p.name)


@pytest.fixture
def workbook(book):
    create_sheet(book, "Data")
    create_sheet(book, "Keep")
    write_sheet(book, [["precious", 1]], sheet="Keep")
    return book


def _untouched(path, before: bytes) -> None:
    assert path.read_bytes() == before, "a refused write must not change the file at all"
    assert list_sheets(path) == ["Sheet", "Data", "Keep"]  # and it is still readable
    assert read_sheet(path, "Keep") == [["precious", 1]]
    assert _hidden_temp_files(path.parent) == []


class TestTheValueRules:
    @pytest.mark.parametrize("value", UNSTORABLE_STRINGS + UNSTORABLE_VALUES, ids=_short)
    def test_check_cell_value_refuses_it(self, value):
        with pytest.raises(CellTypeError):
            check_cell_value(value)

    @pytest.mark.parametrize("value", STORABLE_STRINGS + STORABLE_VALUES, ids=_short)
    def test_check_cell_value_accepts_it(self, value):
        check_cell_value(value)

    def test_the_message_names_the_offending_character(self):
        with pytest.raises(CellTypeError, match="U\\+D800"):
            check_cell_value("a\ud800b")
        with pytest.raises(CellTypeError, match="32,768"):
            check_cell_value("x" * 32_768)

    def test_a_str_subclass_is_checked_like_any_string(self):
        class Name(str):
            pass

        check_cell_value(Name("fine"))
        with pytest.raises(CellTypeError):
            check_cell_value(Name("a\x00b"))


class TestRefusedBeforeAnythingIsWritten:
    """Through every public write path: refused, file byte-identical, still loadable."""

    @pytest.mark.parametrize("value", API_BAD, ids=_short)
    def test_write_sheet(self, workbook, value):
        before = workbook.read_bytes()
        with pytest.raises(CellTypeError):
            write_sheet(workbook, [["ok", value]], sheet="Data")
        _untouched(workbook, before)

    @pytest.mark.parametrize("value", API_BAD, ids=_short)
    def test_append_rows(self, workbook, value):
        before = workbook.read_bytes()
        with pytest.raises(CellTypeError):
            append_rows(workbook, [["ok"], [value]], sheet="Data")
        _untouched(workbook, before)

    @pytest.mark.parametrize("value", API_BAD, ids=_short)
    def test_a_data_value_in_a_table(self, workbook, value):
        before = workbook.read_bytes()
        table = Table(data=[[value]], row_labels=["r"], column_headers=["c"], name="T")
        with pytest.raises(CellTypeError):
            table.create(workbook, sheet="Data")
        _untouched(workbook, before)

    @pytest.mark.parametrize("bad", API_BAD_STRINGS, ids=_short)
    @pytest.mark.parametrize("part", ["name", "header", "label", "corner"])
    def test_every_string_part_of_a_table(self, workbook, part, bad):
        before = workbook.read_bytes()
        kwargs = {"name": "T", "column_headers": ["c"], "data": [[1]], "row_labels": ["r"]}
        if part == "name":
            kwargs["name"] = bad
        elif part == "header":
            kwargs["column_headers"] = [bad]
        elif part == "label":
            kwargs["row_labels"] = [bad]
        else:
            kwargs["corner"] = bad
        with pytest.raises(CellTypeError):
            Table(**kwargs).create(workbook, sheet="Data")
        _untouched(workbook, before)

    @pytest.mark.parametrize("value", [API_BAD_STRINGS[0], API_BAD_STRINGS[1], float("nan")])
    def test_a_later_write_of_an_existing_table(self, workbook, value):
        Table(data=[[1]], row_labels=["r"], column_headers=["c"], name="T").create(
            workbook, sheet="Data"
        )
        table = Table.read(workbook, "T")
        before = workbook.read_bytes()
        table.set_cell(row="r", column="c", value=value)
        table.add_row("new", [1])
        with pytest.raises(CellTypeError):
            table.write(workbook)
        assert workbook.read_bytes() == before
        assert table.data.row_labels == ["r", "new"]  # the object keeps the caller's edits
        assert _hidden_temp_files(workbook.parent) == []

    def test_import_csv_to_xl_refuses_a_poisoned_field(self, workbook, tmp_path):
        source = tmp_path / "in.csv"
        source.write_bytes(b"a,b\n1,x\x0by\n")
        before = workbook.read_bytes()
        with pytest.raises(CellTypeError):
            import_csv_to_xl(source, workbook, sheet="Data")
        _untouched(workbook, before)


class TestStorableValuesRoundTrip:
    @pytest.mark.parametrize("value", STORABLE_STRINGS, ids=_short)
    def test_a_string_comes_back_exactly(self, workbook, value):
        write_sheet(workbook, [["marker", value]], sheet="Data")
        row = read_sheet(workbook, "Data")[0]
        # an empty string is an empty cell (documented); everything else is exact
        assert row[1:] == ([value] if value != "" else [])
        assert list_sheets(workbook) == ["Sheet", "Data", "Keep"]

    @pytest.mark.parametrize("value", STORABLE_STRINGS, ids=_short)
    def test_in_every_part_of_a_table(self, workbook, value):
        if value == "":
            pytest.skip("an empty label or header is not allowed at all")
        Table(
            data=[[1]], row_labels=[value], column_headers=[value], corner=value, name="T"
        ).create(workbook, sheet="Data")
        back = Table.read(workbook, "T")
        assert back.data.row_labels == [value]
        assert back.data.column_headers == [value]
        assert back.data.corner == value

    @pytest.mark.parametrize("value", [v for v in STORABLE_VALUES if v is not None], ids=_short)
    def test_a_storable_value_survives_a_write(self, workbook, value):
        write_sheet(workbook, [[value]], sheet="Data")
        back = read_sheet(workbook, "Data")[0][0]
        if isinstance(value, dt.datetime):
            assert back == value.replace(microsecond=round(value.microsecond / 1000) * 1000) or (
                value.microsecond >= 999500
            )
        elif isinstance(value, dt.date):
            assert back == dt.datetime(value.year, value.month, value.day)
        elif isinstance(value, float) and value == 0:
            assert back == 0
        elif isinstance(value, int) and not isinstance(value, bool) and abs(value) > 2**53:
            assert back == pytest.approx(value)
        else:
            assert back == value and type(back) is type(value)

    def test_a_string_of_exactly_the_limit_is_kept_whole(self, workbook):
        write_sheet(workbook, [["x" * 32_767]], sheet="Data")
        assert len(read_sheet(workbook, "Data")[0][0]) == 32_767


class TestCarriageReturns:
    """Text is stored with newline line breaks on every platform (see the README notes)."""

    @pytest.mark.parametrize(
        ("written", "read_back"),
        [("cr\r\nlf", "cr\nlf"), ("lone\rcr", "lone\ncr"), ("\r\r", "\n\n"), ("lf\n", "lf\n")],
    )
    def test_a_carriage_return_is_stored_as_a_newline(self, workbook, written, read_back):
        write_sheet(workbook, [[written]], sheet="Data")
        assert read_sheet(workbook, "Data") == [[read_back]]


class TestPrecision:
    """What Excel keeps of a time: milliseconds, rounded — and now documented."""

    @pytest.mark.parametrize(
        ("micro", "expected_micro"),
        [
            (0, 0),
            (1, 0),
            (499, 0),
            (500, 0),
            (501, 1000),
            (1000, 1000),
            (123456, 123000),
            (123500, 124000),
        ],
    )
    def test_a_datetime_is_rounded_to_the_millisecond(self, workbook, micro, expected_micro):
        write_sheet(workbook, [[dt.datetime(2026, 1, 1, 12, 0, 0, micro)]], sheet="Data")
        assert read_sheet(workbook, "Data")[0][0].microsecond == expected_micro

    def test_rounding_up_carries_into_the_next_second(self, workbook):
        write_sheet(workbook, [[dt.datetime(2026, 1, 1, 12, 0, 0, 999600)]], sheet="Data")
        assert read_sheet(workbook, "Data")[0][0] == dt.datetime(2026, 1, 1, 12, 0, 1)

    def test_a_timedelta_smaller_than_a_millisecond_is_lost(self, workbook):
        write_sheet(workbook, [[dt.timedelta(microseconds=1)]], sheet="Data")
        assert read_sheet(workbook, "Data")[0][0] == dt.timedelta(0)

    def test_the_millisecond_rounding_is_not_mistaken_for_a_concurrent_edit(self, workbook):
        Table(
            data=[[dt.datetime(2026, 1, 1)]], row_labels=["r"], column_headers=["when"], name="T"
        ).create(workbook, sheet="Data")
        table = Table.read(workbook, "T")
        table.set_cell(row="r", column="when", value=dt.datetime(2026, 5, 5, 5, 5, 5, 123456))
        table.write(workbook)
        table.add_row("r2", [dt.datetime(2026, 1, 1)])
        table.write(workbook)  # nobody else wrote in between
        assert table.last_merge is None


class TestSheetNames:
    BAD: ClassVar[list[str]] = [
        "a\x00b",
        "a\x01b",
        "a\x1fb",
        "a\nb",
        "a\tb",
        "a\rb",
        "a\ud800b",
        "a￾b",
        "a￿b",
        "'lead",
        "trail'",
        "'both'",
        "",
        "x" * 32,
        "a:b",
        "a/b",
        "a\\b",
        "a?b",
        "a*b",
        "a[b",
        "a]b",
        "History",
        "history",
        "HISTORY",
    ]
    GOOD: ClassVar[list[str]] = [
        "Sheet1",
        "x" * 31,
        "bob's",
        "with space",
        " padded ",
        "日本語",
        "😀",
        "é",
        "a\x7fb",
        "a.b",
        "a-b_c",
        "1",
        "a&b",
        "a<b>",
    ]

    @pytest.mark.parametrize("name", BAD, ids=_short)
    def test_a_bad_name_is_refused(self, name):
        with pytest.raises(SheetNameError):
            check_sheet_name(name)

    @pytest.mark.parametrize("name", GOOD, ids=_short)
    def test_a_good_name_is_accepted(self, name):
        check_sheet_name(name)

    @pytest.mark.parametrize("name", BAD, ids=_short)
    def test_no_call_can_create_a_sheet_with_it(self, workbook, name):
        before = workbook.read_bytes()
        for call in (
            lambda: create_sheet(workbook, name),
            lambda: write_sheet(workbook, [["x"]], sheet=name),
            lambda: append_rows(workbook, [["x"]], sheet=name),
            lambda: rename_sheet(workbook, "Data", name),
        ):
            with pytest.raises(SheetNameError):
                call()
        _untouched(workbook, before)

    @pytest.mark.parametrize("name", GOOD, ids=_short)
    def test_a_good_name_can_be_created_used_renamed_and_survives(self, workbook, name):
        create_sheet(workbook, name)
        write_sheet(workbook, [["v"]], sheet=name)
        assert read_sheet(workbook, name) == [["v"]]
        assert name in load_workbook(workbook).sheetnames

    def test_a_non_string_name_is_refused(self):
        for bad in (None, 5, b"x", ["a"]):
            with pytest.raises(SheetNameError):
                check_sheet_name(bad)  # type: ignore[arg-type]
