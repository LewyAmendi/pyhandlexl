"""The fast painter must produce exactly what painting every cell individually would.

``paint_table`` paints each distinct look once and copies its style record onto the other
cells that share it. These tests keep the plain one-cell-at-a-time painter as the reference
and compare the two, cell by cell, across styles and shapes.
"""

from __future__ import annotations

from copy import copy
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from pyhandlexl import Table, TableStyle, create_sheet
from pyhandlexl import _multi_table as mt
from pyhandlexl._multi_table import TableEntry, paint_table

STYLES = {
    "default": TableStyle.DEFAULT,
    "minimal": TableStyle.MINIMAL,
    "none": TableStyle.NONE,
    "no border": TableStyle(border_color=""),
    "no banding": TableStyle(band_fill=""),
    "no header fill": TableStyle(header_fill=""),
    "custom": TableStyle(
        header_font_name="Arial",
        header_font_size=14,
        header_font_color="FF0000",
        header_bold=False,
        header_fill="00FF00",
        data_font_name="Courier New",
        data_font_size=9,
        data_font_color="0000FF",
        band_fill="FFFF00",
        border_color="123456",
    ),
}
SHAPES = [(0, 1), (1, 1), (2, 2), (3, 1), (5, 3), (12, 7), (1, 6), (30, 2)]
ANCHORS = [(1, 1), (3, 4)]


def paint_reference(ws, entry: TableEntry) -> None:
    """The original algorithm: every cell painted individually."""
    style = entry.style
    header_row = entry.anchor_row + 1
    for c in range(entry.width):
        cell = ws.cell(row=header_row, column=entry.anchor_col + c)
        mt._paint_header_like(cell, style)
        mt._paint_border(cell, 0, c, entry.n_rows, entry.n_cols, style)
    for r in range(1, entry.n_rows + 1):
        row = header_row + r
        label = ws.cell(row=row, column=entry.anchor_col)
        mt._paint_header_like(label, style)
        mt._paint_border(label, r, 0, entry.n_rows, entry.n_cols, style)
        banded = bool(style.band_fill) and r % 2 == 0
        for c in range(1, entry.n_cols + 1):
            cell = ws.cell(row=row, column=entry.anchor_col + c)
            mt._paint_data(cell, style, banded)
            mt._paint_border(cell, r, c, entry.n_rows, entry.n_cols, style)


def entry_for(style, shape, anchor) -> TableEntry:
    n_rows, n_cols = shape
    return TableEntry(
        "T", "S", anchor[0], anchor[1], n_rows, n_cols, style, [], created_at=None, modified_at=None
    )


def look(cell) -> tuple:
    """Everything visual about a cell, as real objects (a cell's own attributes are proxies
    that don't compare by value across workbooks)."""
    return (
        copy(cell.font),
        copy(cell.fill),
        copy(cell.border),
        copy(cell.alignment),
        cell.number_format,
        copy(cell.protection),
        cell.has_style,
    )


def assert_same_painting(fast_ws, slow_ws, entry: TableEntry, margin: int = 2) -> None:
    rows = range(1, entry.anchor_row + entry.height + margin)
    cols = range(1, entry.anchor_col + entry.width + margin)
    for r in rows:
        for c in cols:
            assert look(fast_ws.cell(row=r, column=c)) == look(slow_ws.cell(row=r, column=c)), (
                r,
                c,
            )


class TestSameAsPaintingEveryCell:
    @pytest.mark.parametrize("style_name", list(STYLES))
    @pytest.mark.parametrize("shape", SHAPES, ids=lambda s: f"{s[0]}x{s[1]}")
    @pytest.mark.parametrize("anchor", ANCHORS, ids=lambda a: f"at{a[0]},{a[1]}")
    def test_fresh_cells(self, style_name, shape, anchor):
        entry = entry_for(STYLES[style_name], shape, anchor)
        fast, slow = Workbook().active, Workbook().active
        paint_table(fast, entry)
        paint_reference(slow, entry)
        assert_same_painting(fast, slow, entry)

    @pytest.mark.parametrize("style_name", ["default", "custom", "none"])
    @pytest.mark.parametrize("shape", [(2, 2), (6, 4)], ids=lambda s: f"{s[0]}x{s[1]}")
    def test_cells_that_were_painted_and_then_cleared(self, style_name, shape):
        # exactly what a rewrite does: clear the old region, then paint the new one
        old = entry_for(TableStyle.DEFAULT, (8, 5), (1, 1))
        entry = entry_for(STYLES[style_name], shape, (1, 1))
        fast, slow = Workbook().active, Workbook().active
        for ws, painter in ((fast, paint_table), (slow, paint_reference)):
            paint_reference(ws, old)
            mt.clear_region(ws, 1, 1, old.height, old.width)
            painter(ws, entry)
        assert_same_painting(fast, slow, old)

    def test_repainting_with_a_different_style_replaces_the_old_look(self):
        first = entry_for(TableStyle.DEFAULT, (4, 3), (1, 1))
        second = entry_for(TableStyle.NONE, (4, 3), (1, 1))
        fast, slow = Workbook().active, Workbook().active
        for ws, painter in ((fast, paint_table), (slow, paint_reference)):
            painter(ws, first)
            mt.clear_region(ws, 1, 1, first.height, first.width)
            painter(ws, second)
        assert_same_painting(fast, slow, first)

    def test_the_result_survives_saving_and_reloading(self, tmp_path):
        entry = entry_for(STYLES["custom"], (12, 7), (2, 3))
        fast, slow = Workbook(), Workbook()
        paint_table(fast.active, entry)
        paint_reference(slow.active, entry)
        fast.save(tmp_path / "fast.xlsx")
        slow.save(tmp_path / "slow.xlsx")
        assert_same_painting(
            load_workbook(tmp_path / "fast.xlsx").active,
            load_workbook(tmp_path / "slow.xlsx").active,
            entry,
        )


class TestTheStyleTableStaysSmall:
    def test_painting_a_large_table_paints_only_a_handful_of_cells_the_slow_way(self, monkeypatch):
        calls = {"data": 0, "head": 0}
        real_data, real_head = mt._paint_data, mt._paint_header_like

        def counting_data(cell, style, banded):
            calls["data"] += 1
            return real_data(cell, style, banded)

        def counting_head(cell, style):
            calls["head"] += 1
            return real_head(cell, style)

        monkeypatch.setattr(mt, "_paint_data", counting_data)
        monkeypatch.setattr(mt, "_paint_header_like", counting_head)
        paint_table(Workbook().active, entry_for(TableStyle.DEFAULT, (2000, 10), (1, 1)))
        # 22,000 cells, but only the distinct looks (edge/interior x banded) are built afresh
        assert calls["data"] <= 40
        assert calls["head"] <= 20

    def test_the_saved_workbook_holds_a_few_dozen_styles_not_thousands(self, book):
        create_sheet(book, "D")
        Table(
            data=[[i * j for j in range(10)] for i in range(500)],
            row_labels=[f"r{i}" for i in range(500)],
            column_headers=[f"c{j}" for j in range(10)],
            name="Big",
        ).create(book, sheet="D")
        wb = load_workbook(book)
        assert len(wb._cell_styles) <= 60  # cell formats
        assert len(wb._borders) <= 40 and len(wb._fills) <= 6 and len(wb._fonts) <= 6


class TestStillReadsBackItsOwnStyle:
    @pytest.mark.parametrize("name", ["default", "minimal", "none", "custom"])
    def test_a_rebuilt_schema_recovers_the_style_from_fast_painted_cells(self, book, name):
        import warnings

        from pyhandlexl import SchemaRebuiltWarning
        from pyhandlexl._multi_table import SCHEMA_SHEET

        create_sheet(book, "D")
        Table(
            data=[[1, 2], [3, 4], [5, 6]],
            row_labels=["a", "b", "c"],
            column_headers=["x", "y"],
            name="T",
            style=STYLES[name],
        ).create(book, sheet="D")
        wb = load_workbook(book)
        del wb[SCHEMA_SHEET]
        wb.save(book)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", SchemaRebuiltWarning)
            assert Table.read(book, "T").style == STYLES[name]


TEMPORAL = [
    datetime(2026, 3, 4, 5, 6, 7),
    date(2026, 3, 4),
    time(1, 2, 3),
    timedelta(hours=5),
    timedelta(days=2, hours=1),
    1,
    2.5,
    True,
    "text",
    None,
]


def fill_values(ws, entry: TableEntry) -> None:
    """Write a value into every cell of the table's region, before it is painted — the order
    a real write uses, and the moment openpyxl gives a date or a duration its number format."""
    n = 0
    for r in range(entry.anchor_row, entry.anchor_row + entry.height):
        for c in range(entry.anchor_col, entry.anchor_col + entry.width):
            ws.cell(row=r, column=c).value = TEMPORAL[n % len(TEMPORAL)]
            n += 1


class TestValuesKeepTheirOwnNumberFormat:
    @pytest.mark.parametrize("style_name", list(STYLES))
    @pytest.mark.parametrize("shape", [(1, 1), (4, 3), (12, 7)], ids=lambda s: f"{s[0]}x{s[1]}")
    def test_painting_over_written_values_matches_painting_every_cell(self, style_name, shape):
        entry = entry_for(STYLES[style_name], shape, (2, 2))
        fast, slow = Workbook().active, Workbook().active
        for ws, painter in ((fast, paint_table), (slow, paint_reference)):
            fill_values(ws, entry)
            painter(ws, entry)
        assert_same_painting(fast, slow, entry)

    def test_a_cell_keeps_the_format_its_value_gave_it(self):
        entry = entry_for(TableStyle.DEFAULT, (3, 3), (1, 1))
        ws = Workbook().active
        fill_values(ws, entry)
        before = [
            [ws.cell(row=r, column=c).number_format for c in range(1, 5)] for r in range(1, 5)
        ]
        paint_table(ws, entry)
        after = [[ws.cell(row=r, column=c).number_format for c in range(1, 5)] for r in range(1, 5)]
        assert after == before
        assert len({f for row in after for f in row}) > 1  # the fixture really mixes formats


class TestTemporalValuesRoundTrip:
    @pytest.mark.parametrize("style_name", ["default", "none", "custom"])
    def test_create_then_read_returns_the_same_types(self, book, style_name):
        create_sheet(book, "D")
        values = [
            datetime(2026, 3, 4, 5, 6, 7),
            date(2026, 3, 4),
            time(1, 2, 3),
            timedelta(hours=5),
            timedelta(days=2, hours=1),
        ]
        Table(
            data=[values[:3], [*values[3:], 1]],
            row_labels=["a", "b"],
            column_headers=["x", "y", "z"],
            name="T",
            style=STYLES[style_name],
        ).create(book, sheet="D")
        back = Table.read(book, "T")
        assert (
            back.data.rows
            == [
                [values[0], datetime(2026, 3, 4), values[2]],  # Excel has no date-only type
                [*values[3:], 1],
            ]
        )
        assert [type(v) for v in back.data.rows[0]] == [datetime, datetime, time]
        assert type(back.data.rows[1][0]) is timedelta

    def test_a_write_after_editing_keeps_the_types(self, book):
        create_sheet(book, "D")
        Table(
            data=[[timedelta(hours=1), time(1, 0)]],
            row_labels=["a"],
            column_headers=["x", "y"],
            name="T",
        ).create(book, sheet="D")
        t = Table.read(book, "T")
        t.add_row("b", [timedelta(minutes=90), time(2, 30)])
        t.write(book)
        back = Table.read(book, "T")
        assert back.data.rows == [
            [timedelta(hours=1), time(1, 0)],
            [timedelta(minutes=90), time(2, 30)],
        ]
