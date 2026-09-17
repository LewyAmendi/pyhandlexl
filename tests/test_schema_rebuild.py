"""Tests for automatic schema recovery when _pyhandlexl_tables is missing."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from pyhandlexl import (
    ColumnType,
    SchemaRebuiltWarning,
    Table,
    TableNotFoundError,
    TableStyle,
    create_sheet,
    list_tables,
)
from pyhandlexl import _multi_table as mt
from pyhandlexl._multi_table import SCHEMA_SHEET


@pytest.fixture
def data_sheet(book):
    create_sheet(book, "Data")
    return book


def _delete_schema_sheet(path) -> None:
    wb = load_workbook(path)
    del wb[SCHEMA_SHEET]
    wb.save(path)


class TestRebuildBasicGeometry:
    def test_single_table_recovers_exactly(self, data_sheet):
        Table(
            data=[[100, 200], [150, 250]],
            column_headers=["North", "South"],
            row_labels=["Q1", "Q2"],
            corner="Metric",
            name="Sales",
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "Sales")
        assert t.data.rows == [[100, 200], [150, 250]]
        assert t.data.column_headers == ["North", "South"]
        assert t.data.row_labels == ["Q1", "Q2"]
        assert t.data.corner == "Metric"

    def test_table_with_no_data_rows_yet(self, data_sheet):
        Table(column_headers=["a", "b"], name="Empty").create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "Empty")
        assert t.data.column_headers == ["a", "b"]
        assert t.data.rows == []


class TestRebuildMultipleTables:
    def test_two_tables_bound_each_other_exactly(self, data_sheet):
        sales = Table(
            data=[[100, 200]], column_headers=["North", "South"], row_labels=["Q1"], name="Sales"
        )
        sales.create(data_sheet, sheet="Data")
        inventory = Table(data=[[10]], column_headers=["Units"], row_labels=["A"], name="Inventory")
        inventory.create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(data_sheet)
        assert set(names) == {"Sales", "Inventory"}
        assert Table.read(data_sheet, "Sales").data.rows == [[100, 200]]
        assert Table.read(data_sheet, "Inventory").data.rows == [[10]]

    def test_three_tables_in_a_row_all_bound_correctly(self, data_sheet):
        for name, headers in [("A", ["a1", "a2"]), ("B", ["b1"]), ("C", ["c1", "c2", "c3"])]:
            Table(
                data=[[1] * len(headers)],
                column_headers=headers,
                row_labels=["r"],
                name=name,
            ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(data_sheet)
        assert set(names) == {"A", "B", "C"}
        assert Table.read(data_sheet, "A").data.column_headers == ["a1", "a2"]
        assert Table.read(data_sheet, "B").data.column_headers == ["b1"]
        assert Table.read(data_sheet, "C").data.column_headers == ["c1", "c2", "c3"]

    def test_tables_on_different_sheets(self, book):
        create_sheet(book, "One")
        create_sheet(book, "Two")
        Table(data=[[1]], column_headers=["x"], row_labels=["r"], name="First").create(
            book, sheet="One"
        )
        Table(data=[[2]], column_headers=["y"], row_labels=["s"], name="Second").create(
            book, sheet="Two"
        )
        _delete_schema_sheet(book)

        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(book)
        assert set(names) == {"First", "Second"}


def _plant_raw_table(
    path, sheet: str, name: str, corner: str, headers: list, rows: list[list]
) -> None:
    """Hand-write a table's cells directly, bypassing Table's own validation.

    Column headers/row labels can no longer be "" through the public API
    (see test_table.py's empty-header/label validation tests), but a
    hand-edited or pre-existing workbook can still contain one — these tests
    exercise the rebuild scanner's tolerance for that, independent of what
    the API itself allows.
    """
    wb = load_workbook(path)
    ws = wb[sheet]
    ws.cell(row=1, column=1, value="TABLE NAME")
    ws.cell(row=1, column=2, value=name)
    ws.cell(row=2, column=1, value=corner)
    for j, header in enumerate(headers):
        ws.cell(row=2, column=2 + j, value=header)
    for i, row in enumerate(rows):
        ws.cell(row=3 + i, column=1, value=row[0])
        for j, value in enumerate(row[1]):
            ws.cell(row=3 + i, column=2 + j, value=value)
    wb.save(path)


class TestRebuildEdgeCases:
    """Blank headers/labels can't be created via the API anymore, but the
    scanner must still tolerate one already sitting in a hand-edited or
    pre-existing sheet rather than mistaking it for the table's end."""

    def test_blank_header_in_the_middle_is_not_mistaken_for_the_end(self, data_sheet):
        _plant_raw_table(data_sheet, "Data", "T", "", ["a", "", "c"], [])

        with pytest.warns(SchemaRebuiltWarning):
            list_tables(data_sheet)
        entry = mt.load_schema(load_workbook(data_sheet))["T"]
        assert entry.n_cols == 3
        assert entry.n_rows == 0

    def test_blank_row_label_in_the_middle_is_not_mistaken_for_the_end(self, data_sheet):
        _plant_raw_table(data_sheet, "Data", "T", "", ["x"], [["r1", [1]], ["", [2]], ["r3", [3]]])

        with pytest.warns(SchemaRebuiltWarning):
            list_tables(data_sheet)
        entry = mt.load_schema(load_workbook(data_sheet))["T"]
        assert entry.n_cols == 1
        assert entry.n_rows == 3

    def test_blank_label_but_real_data_is_kept(self, data_sheet):
        _plant_raw_table(data_sheet, "Data", "T", "", ["a", "b"], [["", [1, 2]]])

        with pytest.warns(SchemaRebuiltWarning):
            list_tables(data_sheet)
        entry = mt.load_schema(load_workbook(data_sheet))["T"]
        assert entry.n_cols == 2
        assert entry.n_rows == 1


class TestRebuildStyle:
    def test_default_style_recovered(self, data_sheet):
        # >=2 data rows: banding can only be detected by comparing two rows,
        # so a single-row table would always (correctly) come back with
        # band_fill="" regardless of the original style — see
        # test_single_data_row_cannot_confirm_banding below.
        Table(data=[[1], [2]], column_headers=["a"], row_labels=["x", "y"], name="T").create(
            data_sheet, sheet="Data"
        )
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "T")
        assert t.style == TableStyle.DEFAULT

    def test_minimal_style_recovered(self, data_sheet):
        Table(
            data=[[1]], column_headers=["a"], row_labels=["x"], name="T", style=TableStyle.MINIMAL
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "T")
        assert t.style == TableStyle.MINIMAL

    def test_custom_style_recovered_exactly(self, data_sheet):
        brand = TableStyle(
            header_font_name="Georgia",
            header_font_size=14,
            header_font_color="FFFFFF",
            header_bold=True,
            header_fill="2E7D32",
            data_font_name="Georgia",
            data_font_size=10,
            data_font_color="1B1B1B",
            band_fill="E8F5E9",
            border_color="1B1B1B",
        )
        Table(
            data=[[1, 2], [3, 4], [5, 6]],
            column_headers=["a", "b"],
            row_labels=["r1", "r2", "r3"],
            name="Brand",
            style=brand,
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "Brand")
        assert t.style == brand

    def test_style_with_no_banding_detected_as_no_banding(self, data_sheet):
        no_band = TableStyle(band_fill="")
        Table(
            data=[[1], [2], [3]],
            column_headers=["a"],
            row_labels=["r1", "r2", "r3"],
            name="T",
            style=no_band,
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "T")
        assert t.style.band_fill == ""

    def test_single_data_row_cannot_confirm_banding(self, data_sheet):
        # a known, inherent limit of reading style back from cells: banding
        # can only be detected by comparing two data rows, so a table with
        # exactly one always comes back reporting no banding, even though
        # TableStyle.DEFAULT actually has some — there's nothing wrong to
        # fix here, it's a real gap in what's recoverable from the sheet.
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "T")
        assert t.style.band_fill == ""
        assert TableStyle.DEFAULT.band_fill != ""


class TestRebuildColumnTypes:
    def test_column_type_restrictions_cannot_be_recovered(self, data_sheet):
        # a known, inherent limit: nothing in a cell says "this column is
        # restricted," only what happens to already be in it, so a rebuild
        # always reports ColumnType.ANY for every column, even though this
        # table was originally restricted.
        Table(
            data=[[1]],
            column_headers=["a"],
            row_labels=["x"],
            name="T",
            column_types={"a": ColumnType.NUMBER},
        ).create(data_sheet, sheet="Data")
        _delete_schema_sheet(data_sheet)
        with pytest.warns(SchemaRebuiltWarning):
            t = Table.read(data_sheet, "T")
        assert t.column_types == {"a": ColumnType.ANY}


class TestRebuildConflicts:
    def test_duplicate_marker_excluded_but_others_recovered(self, book):
        create_sheet(book, "A")
        create_sheet(book, "B")
        Table(data=[[1]], column_headers=["x"], row_labels=["r"], name="Dup").create(
            book, sheet="A"
        )
        Table(data=[[3]], column_headers=["z"], row_labels=["t"], name="Unique").create(
            book, sheet="A"
        )
        # Simulate a hand copy-paste onto sheet B that plants a second "Dup"
        # marker — Table.create() itself would never allow this, since names
        # are unique per workbook.
        wb = load_workbook(book)
        ws_b = wb["B"]
        ws_b["A1"] = "TABLE NAME"
        ws_b["B1"] = "Dup"
        ws_b["A2"] = "corner"
        ws_b["B2"] = "y"
        ws_b["A3"] = "s"
        ws_b["B3"] = 2
        del wb[SCHEMA_SHEET]
        wb.save(book)

        with pytest.warns(SchemaRebuiltWarning):
            names = list_tables(book)
        assert names == ["Unique"]

        with pytest.raises(TableNotFoundError):
            Table.read(book, "Dup")
        assert Table.read(book, "Unique").data.rows == [[3]]


class TestRebuildNoOpWhenNoTables:
    def test_fresh_workbook_warns_nothing_and_writes_nothing(self, book, recwarn):
        names = list_tables(book)
        assert names == []
        assert len(recwarn) == 0
        assert SCHEMA_SHEET not in load_workbook(book).sheetnames


class TestRebuildPersistence:
    def test_rebuild_is_not_repeated_on_a_second_call(self, data_sheet, recwarn):
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            list_tables(data_sheet)
        assert SCHEMA_SHEET in load_workbook(data_sheet).sheetnames

        recwarn.clear()
        names = list_tables(data_sheet)
        assert names == ["T"]
        assert len(recwarn) == 0

    def test_table_read_also_triggers_and_persists_rebuild(self, data_sheet):
        Table(data=[[1]], column_headers=["a"], row_labels=["x"], name="T").create(
            data_sheet, sheet="Data"
        )
        _delete_schema_sheet(data_sheet)

        with pytest.warns(SchemaRebuiltWarning):
            Table.read(data_sheet, "T")
        assert SCHEMA_SHEET in load_workbook(data_sheet).sheetnames
