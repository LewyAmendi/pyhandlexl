"""The API-freeze decisions: which errors are raised, and a few things worth keeping cheap."""

from __future__ import annotations

import contextlib
import csv
import inspect
from typing import ClassVar

import pytest
from openpyxl import load_workbook

import pyhandlexl as p
from pyhandlexl import (
    InvalidFileError,
    MalformedTableError,
    PyhandlexlError,
    SheetExistsError,
    SheetNameError,
    Table,
    create_sheet,
    read_sheet,
    rename_sheet,
    write_sheet,
)


class TestTheExceptionHierarchy:
    """Everything the library raises for a problem with the *data or the file* is a
    ``PyhandlexlError``; a wrong *call* raises the standard exception. Where an error was a
    ``ValueError`` before, it still is."""

    LIBRARY_ERRORS: ClassVar[dict[type, tuple[type, ...]]] = {
        p.CellTypeError: (TypeError,),
        p.ColumnTypeError: (TypeError,),
        p.DimensionError: (ValueError,),
        p.FileLockedError: (OSError,),
        p.FileReadOnlyError: (p.FileLockedError, OSError),
        p.InvalidFileError: (ValueError,),
        p.MalformedTableError: (ValueError,),
        p.SheetExistsError: (p.SheetNameError, ValueError),
        p.SheetKindError: (ValueError,),
        p.SheetNameError: (ValueError,),
        p.SheetNotFoundError: (KeyError,),
        p.TableExistsError: (ValueError,),
        p.TableNotFoundError: (KeyError,),
    }

    def test_every_library_error_derives_from_the_base_and_keeps_its_builtin(self):
        for error, builtins in self.LIBRARY_ERRORS.items():
            assert issubclass(error, PyhandlexlError), error
            for builtin in builtins:
                assert issubclass(error, builtin), (error, builtin)

    def test_every_exception_class_the_package_exports_is_covered_here(self):
        exported = {
            obj
            for name in p.__all__
            if inspect.isclass(obj := getattr(p, name))
            and issubclass(obj, Exception)
            and not issubclass(obj, Warning)
            and obj is not PyhandlexlError
        }
        assert exported == set(self.LIBRARY_ERRORS)


class TestSheetExistsError:
    def test_create_sheet_on_a_name_that_is_taken(self, book):
        with pytest.raises(SheetExistsError, match="already exists"):
            create_sheet(book, "Sheet")

    def test_a_name_that_differs_only_by_capitalisation_is_the_same_name(self, book):
        create_sheet(book, "Data")
        with pytest.raises(SheetExistsError, match="case-insensitive"):
            create_sheet(book, "data")

    def test_rename_to_a_taken_name(self, book):
        create_sheet(book, "Other")
        with pytest.raises(SheetExistsError):
            rename_sheet(book, "Sheet", "Other")
        with pytest.raises(SheetExistsError):
            rename_sheet(book, "Sheet", "OTHER")

    def test_write_sheet_to_a_capitalisation_clash(self, book):
        create_sheet(book, "Data")
        with pytest.raises(SheetExistsError):
            write_sheet(book, [[1]], sheet="dAtA")

    def test_it_is_still_a_value_error_and_a_sheet_name_error(self, book):
        for catch in (ValueError, SheetNameError, PyhandlexlError):
            with pytest.raises(catch):
                create_sheet(book, "Sheet")

    def test_the_idempotent_create_pattern(self, book):
        for _ in range(3):
            with contextlib.suppress(SheetExistsError):
                create_sheet(book, "Log")
        assert load_workbook(book).sheetnames.count("Log") == 1

    def test_an_invalid_name_is_not_a_sheet_exists_error(self, book):
        with pytest.raises(SheetNameError) as caught:
            create_sheet(book, "a/b")
        assert not isinstance(caught.value, SheetExistsError)


class TestMalformedTableError:
    @pytest.fixture
    def path(self, book):
        create_sheet(book, "D")
        Table(
            data=[[1, 2], [3, 4]], row_labels=["a", "b"], column_headers=["x", "y"], name="T"
        ).create(book, sheet="D")
        return book

    def blank(self, path, ref):
        wb = load_workbook(path)
        wb["D"][ref] = None
        wb.save(path)

    def test_a_blank_header(self, path):
        self.blank(path, "C2")
        with pytest.raises(MalformedTableError, match="header in cell C2 is blank"):
            Table.read(path, "T")

    def test_a_blank_label(self, path):
        self.blank(path, "A4")
        with pytest.raises(MalformedTableError, match="label in cell A4 is blank"):
            Table.read(path, "T")

    def test_it_is_still_a_value_error_and_a_library_error(self, path):
        self.blank(path, "A3")
        for catch in (ValueError, PyhandlexlError):
            with pytest.raises(catch):
                Table.read(path, "T")

    def test_the_other_reads_still_work(self, path):
        self.blank(path, "A3")
        assert p.list_tables(path) == ["T"]
        assert p.table_info(path, "T").n_rows == 2


class TestAnUnreadableFile:
    def test_a_csv_that_cannot_be_parsed(self, tmp_path, monkeypatch):
        path = tmp_path / "d.csv"
        path.write_text("a,b\n", encoding="utf-8")

        class Failing:
            line_num = 7

            def __init__(self, *args, **kwargs):
                pass

            def __iter__(self):
                raise csv.Error("line contains NUL")

        monkeypatch.setattr(csv, "reader", Failing)
        with pytest.raises(InvalidFileError, match=r"not a readable CSV \(line 7\)"):
            read_sheet(path)

    def test_an_invalid_file_is_a_value_error_too(self, tmp_path):
        path = tmp_path / "junk.xlsx"
        path.write_bytes(b"not a zip")
        for catch in (InvalidFileError, ValueError, PyhandlexlError):
            with pytest.raises(catch):
                read_sheet(path)


class TestCheapLabelAccess:
    @pytest.fixture
    def t(self):
        return Table(
            data=[[1, 2], [3, 4]],
            row_labels=["a", "b"],
            column_headers=["x", "y"],
            corner="c",
            name="T",
        )

    def test_row_labels_column_headers_and_corner(self, t):
        assert t.row_labels == ["a", "b"]
        assert t.column_headers == ["x", "y"]
        assert t.corner == "c"

    def test_they_are_copies(self, t):
        t.row_labels.append("zz")
        t.column_headers.clear()
        assert t.row_labels == ["a", "b"] and t.column_headers == ["x", "y"]

    def test_they_follow_every_edit(self, t):
        t.rename_row("a", "A")
        t.add_column("z", [5, 6])
        t.set_corner("k")
        t.drop_row("b")
        assert (t.row_labels, t.column_headers, t.corner) == (["A"], ["x", "y", "z"], "k")

    def test_they_do_not_build_the_whole_snapshot(self, t, monkeypatch):
        # t.data copies every value (twice, as rows and as columns) on each access; a loop
        # over rows that touches it each time is quadratic, so these must not
        monkeypatch.setattr(Table, "data", property(lambda self: pytest.fail("t.data was built")))
        assert t.row_labels and t.column_headers and t.corner == "c"
        assert "a" in t.row_labels

    def test_they_match_the_snapshot(self, t):
        assert t.row_labels == t.data.row_labels
        assert t.column_headers == t.data.column_headers
        assert t.corner == t.data.corner

    def test_a_table_read_from_a_file(self, book, t):
        create_sheet(book, "D")
        t.create(book, sheet="D")
        back = Table.read(book, "T")
        assert (back.row_labels, back.column_headers, back.corner) == (["a", "b"], ["x", "y"], "c")


class TestTheValidatorsTakeOneArgumentByPosition:
    """Their parameter names are not part of the API."""

    @pytest.mark.parametrize(
        ("function", "keyword"),
        [
            (p.check_cell_value, "value"),
            (p.check_sheet_name, "name"),
            (p.check_sheet_name, "sheet_name"),
            (p.is_valid_xlsx, "path"),
        ],
    )
    def test_a_keyword_is_refused(self, function, keyword):
        with pytest.raises(TypeError):
            function(**{keyword: "x"})

    def test_by_position_they_work(self):
        p.check_cell_value(1)
        p.check_sheet_name("Data")
        assert p.is_valid_xlsx("no such file.xlsx") is False
        with pytest.raises(SheetNameError):
            p.check_sheet_name("a/b")


class TestTheTypeAliasesAreExported:
    def test_they_can_be_imported_and_used_in_annotations(self):
        from typing import get_args

        from pyhandlexl import Orientation, SheetKind

        assert set(get_args(SheetKind)) == {"table", "grid", "empty"}
        assert set(get_args(Orientation)) == {"rows", "columns"}
        assert {"SheetKind", "Orientation"} <= set(p.__all__)
