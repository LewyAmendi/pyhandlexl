"""Tests for named tables: multiple Tables sharing one worksheet."""

from __future__ import annotations

import pytest

from pyhandlexl import (
    CellTypeError,
    SheetNotFoundError,
    Table,
    TableExistsError,
    TableNotFoundError,
    create_sheet,
    delete_table,
    list_sheets,
    list_tables,
    read_sheet,
    write_sheet,
)


@pytest.fixture
def data_sheet(book):
    create_sheet(book, "Data")
    return book


def _sales(name="Sales"):
    return Table(
        data=[[100, 200], [150, 250]],
        column_headers=["North", "South"],
        row_labels=["Q1", "Q2"],
        corner="Metric",
        name=name,
    )


def _inventory(name="Inventory"):
    return Table(
        data=[[10], [20], [30]],
        column_headers=["Units"],
        row_labels=["A", "B", "C"],
        corner="Item",
        name=name,
    )


class TestCreate:
    def test_writes_the_marker_row(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        grid = read_sheet(data_sheet, "Data")
        assert grid[0] == ["TABLE NAME", "Sales"]
        assert grid[1] == ["Metric", "North", "South"]
        assert grid[2] == ["Q1", 100, 200]

    def test_registers_in_list_tables(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        assert list_tables(data_sheet) == ["Sales"]

    def test_second_table_stacks_to_the_right_with_a_gap(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        _inventory().create(data_sheet, sheet="Data")
        grid = read_sheet(data_sheet, "Data", pad=True)
        # Sales occupies cols 1-3 (label+North+South); col 4 must be empty;
        # Inventory starts at col 5.
        assert grid[0][3] is None
        assert grid[0][4:6] == ["TABLE NAME", "Inventory"]

    def test_duplicate_name_raises(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        with pytest.raises(TableExistsError):
            _sales().create(data_sheet, sheet="Data")

    def test_missing_sheet_raises(self, book):
        with pytest.raises(SheetNotFoundError):
            _sales().create(book, sheet="Ghost")

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            _sales().create(tmp_path / "nope.xlsx", sheet="Data")

    def test_unsupported_cell_type_raises(self, data_sheet):
        bad = Table(data=[[[1, 2]]], column_headers=["a"], row_labels=["r"], name="Bad")
        with pytest.raises(CellTypeError):
            bad.create(data_sheet, sheet="Data")


class TestReadNamed:
    def test_round_trip(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        t = Table.read(data_sheet, name="Sales")
        assert t.name == "Sales"
        assert t.data.column_headers == ["North", "South"]
        assert t.data.row_labels == ["Q1", "Q2"]
        assert t.data.corner == "Metric"
        assert t.data.rows == [[100, 200], [150, 250]]

    def test_two_tables_are_independent(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        _inventory().create(data_sheet, sheet="Data")
        sales = Table.read(data_sheet, name="Sales")
        inventory = Table.read(data_sheet, name="Inventory")
        assert sales.data.rows == [[100, 200], [150, 250]]
        assert inventory.data.rows == [[10], [20], [30]]

    def test_unknown_name_raises(self, data_sheet):
        with pytest.raises(TableNotFoundError):
            Table.read(data_sheet, name="Ghost")

    def test_non_string_name_raises_typeerror(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        with pytest.raises(TypeError):
            Table.read(data_sheet, name=123)

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            Table.read(tmp_path / "nope.xlsx", name="Sales")

    def test_self_heals_when_marker_has_moved(self, data_sheet):
        from openpyxl import load_workbook

        _sales().create(data_sheet, sheet="Data")
        wb = load_workbook(data_sheet)
        wb["Data"].insert_cols(1)  # shove everything one column to the right
        wb.save(data_sheet)

        t = Table.read(data_sheet, name="Sales")
        assert t.data.rows == [[100, 200], [150, 250]]

        # the schema should now reflect the healed position
        assert list_tables(data_sheet) == ["Sales"]
        t_again = Table.read(data_sheet, name="Sales")
        assert t_again.data.rows == [[100, 200], [150, 250]]

    def test_raises_when_marker_is_gone(self, data_sheet):
        from openpyxl import load_workbook

        _sales().create(data_sheet, sheet="Data")
        wb = load_workbook(data_sheet)
        wb["Data"]["A1"] = "not a marker anymore"
        wb.save(data_sheet)

        with pytest.raises(TableNotFoundError):
            Table.read(data_sheet, name="Sales")


class TestWriteNamed:
    def test_updates_in_place_without_growth(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        t = Table.read(data_sheet, name="Sales")
        t.set_cell(row="Q1", column="North", value=999)
        t.write(data_sheet)
        assert Table.read(data_sheet, name="Sales").read_cell(row="Q1", column="North") == 999

    def test_sheet_argument_rejected(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        t = Table.read(data_sheet, name="Sales")
        with pytest.raises(TypeError):
            t.write(data_sheet, sheet="Data")

    def test_write_before_create_raises(self, data_sheet):
        t = _sales()
        with pytest.raises(TableNotFoundError):
            t.write(data_sheet)

    def test_write_to_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            _sales().write(tmp_path / "nope.xlsx")

    def test_growth_shifts_the_next_table_right(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        _inventory().create(data_sheet, sheet="Data")

        sales = Table.read(data_sheet, name="Sales")
        sales.add_column("East", [300, 350])
        sales.write(data_sheet)

        assert Table.read(data_sheet, name="Sales").data.column_headers == [
            "North",
            "South",
            "East",
        ]
        # Inventory must still be readable and unchanged after the shift.
        assert Table.read(data_sheet, name="Inventory").data.rows == [[10], [20], [30]]

    def test_growth_with_no_neighbours_does_not_error(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        sales = Table.read(data_sheet, name="Sales")
        sales.add_column("East", [300, 350])
        sales.write(data_sheet)  # should not raise
        assert Table.read(data_sheet, name="Sales").data.column_headers[-1] == "East"

    def test_shrink_clears_stale_cells(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        t = Table.read(data_sheet, name="Sales")
        t.drop_column("South")
        t.write(data_sheet)
        grid = read_sheet(data_sheet, "Data")
        assert grid[1] == ["Metric", "North"]  # no leftover "South" column

    def test_row_growth_never_shifts_anything(self, data_sheet):
        from pyhandlexl import _multi_table as mt
        from pyhandlexl._safety import safe_load

        _sales().create(data_sheet, sheet="Data")
        _inventory().create(data_sheet, sheet="Data")

        wb = safe_load(data_sheet)
        before_col = mt.load_schema(wb)["Inventory"].anchor_col
        wb.close()

        sales = Table.read(data_sheet, name="Sales")
        sales.add_row("Q3", [175, 300])
        sales.write(data_sheet)

        wb = safe_load(data_sheet)
        after_col = mt.load_schema(wb)["Inventory"].anchor_col
        wb.close()

        assert before_col == after_col
        assert Table.read(data_sheet, name="Inventory").data.rows == [[10], [20], [30]]


class TestDeleteTable:
    def test_removes_it_from_list_tables(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        delete_table(data_sheet, "Sales")
        assert list_tables(data_sheet) == []

    def test_clears_its_cells(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        delete_table(data_sheet, "Sales")
        grid = read_sheet(data_sheet, "Data")
        assert grid == []

    def test_subsequent_read_raises(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        delete_table(data_sheet, "Sales")
        with pytest.raises(TableNotFoundError):
            Table.read(data_sheet, name="Sales")

    def test_unknown_name_raises(self, data_sheet):
        with pytest.raises(TableNotFoundError):
            delete_table(data_sheet, "Ghost")

    def test_does_not_disturb_other_tables(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        _inventory().create(data_sheet, sheet="Data")
        delete_table(data_sheet, "Sales")
        assert list_tables(data_sheet) == ["Inventory"]
        assert Table.read(data_sheet, name="Inventory").data.rows == [[10], [20], [30]]

    def test_name_can_be_reused_after_deletion(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        delete_table(data_sheet, "Sales")
        _sales().create(data_sheet, sheet="Data")  # should not raise TableExistsError
        assert list_tables(data_sheet) == ["Sales"]


class TestNamedTableAndGridLayoutCoexist:
    def test_a_plain_grid_sheet_can_sit_alongside_named_tables(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        write_sheet(data_sheet, [["x", "y"]], sheet="Plain")
        assert read_sheet(data_sheet, "Plain") == [["x", "y"]]
        assert list_tables(data_sheet) == ["Sales"]


class TestShow:
    def test_show_all(self, capsys):
        t = Table(
            data=[[1], [2]], column_headers=["v"], row_labels=["a", "b"], corner="#", name="T"
        )
        t.show()
        out = capsys.readouterr().out
        assert "a" in out and "b" in out and "v" in out

    def test_show_rows_limits_output(self, capsys):
        t = Table(
            data=[[i] for i in range(5)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(5)],
            name="T",
        )
        t.show(rows=2)
        lines = capsys.readouterr().out.strip().splitlines()
        assert len(lines) == 3  # header + 2 data rows

    def test_show_head_tail_inserts_divider(self, capsys):
        t = Table(
            data=[[i] for i in range(10)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(10)],
            name="T",
        )
        t.show(head=2, tail=2)
        lines = capsys.readouterr().out.strip().splitlines()
        # header + 2 head rows + divider + 2 tail rows
        assert len(lines) == 6
        assert "..." in lines[3]

    def test_show_head_tail_no_divider_when_it_would_cover_everything(self, capsys):
        t = Table(data=[[1], [2]], column_headers=["v"], row_labels=["a", "b"], name="T")
        t.show(head=2, tail=2)
        lines = capsys.readouterr().out.strip().splitlines()
        assert not any("..." in line for line in lines)

    def test_show_table_with_no_rows_yet(self, capsys):
        # column_headers is mandatory, so there's always at least a header
        # row to show, even before the first add_row()
        Table(data=[], column_headers=["v"], name="T").show()
        out = capsys.readouterr().out
        assert "v" in out

    def test_default_truncates_to_head_and_tail_5(self, capsys):
        t = Table(
            data=[[i] for i in range(12)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(12)],
            name="T",
        )
        t.show()  # no args — should default to head=5, tail=5
        lines = capsys.readouterr().out.strip().splitlines()
        # header + 5 head rows + divider + 5 tail rows
        assert len(lines) == 12
        assert "..." in lines[6]
        assert "r0 " in lines[1]
        assert "r11" in lines[-1]
        assert not any("r5" in line or "r6" in line for line in lines)

    def test_negative_rows_raises(self):
        t = Table(data=[[1]], column_headers=["v"], row_labels=["a"], name="T")
        with pytest.raises(ValueError):
            t.show(rows=-1)
        with pytest.raises(ValueError):
            t.show(head=-1)
        with pytest.raises(ValueError):
            t.show(tail=-1)

    def test_non_int_rows_raises_typeerror(self):
        t = Table(data=[[1]], column_headers=["v"], row_labels=["a"], name="T")
        with pytest.raises(TypeError):
            t.show(head="a")

    def test_default_shows_everything_when_10_rows_or_fewer(self, capsys):
        t = Table(
            data=[[i] for i in range(8)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(8)],
            name="T",
        )
        t.show()
        lines = capsys.readouterr().out.strip().splitlines()
        assert len(lines) == 9  # header + all 8 rows, no divider
        assert not any("..." in line for line in lines)

    def test_head_none_tail_none_shows_everything_explicitly(self, capsys):
        t = Table(
            data=[[i] for i in range(20)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(20)],
            name="T",
        )
        t.show(head=None, tail=None)
        lines = capsys.readouterr().out.strip().splitlines()
        assert len(lines) == 21  # header + all 20 rows

    def test_rows_overrides_the_head_tail_defaults(self, capsys):
        t = Table(
            data=[[i] for i in range(20)],
            column_headers=["v"],
            row_labels=[f"r{i}" for i in range(20)],
            name="T",
        )
        t.show(rows=3)
        lines = capsys.readouterr().out.strip().splitlines()
        assert len(lines) == 4  # header + 3 rows, head/tail defaults ignored


class TestSchemaSheetIsHidden:
    def test_reserved_sheet_is_excluded_from_list_sheets(self, data_sheet):
        _sales().create(data_sheet, sheet="Data")
        assert "_pyhandlexl_tables" not in list_sheets(data_sheet)
