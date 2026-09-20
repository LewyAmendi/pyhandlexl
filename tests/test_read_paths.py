"""Reads open a workbook read-only (parsing only the sheets they touch) — and must give exactly
what a full load gives, whatever the file looks like.

A read-only sheet trusts what the file declares about itself (its size, which rows it lists), and
files in the wild are not always truthful, so this runs every read function over deliberately
awkward workbooks in both modes and demands identical results, errors and warnings.
"""

from __future__ import annotations

import datetime as dt
import re
import shutil
import warnings
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName

import pyhandlexl._safety as safety
from pyhandlexl import (
    InvalidFileError,
    Table,
    TableNotFoundError,
    create_sheet,
    create_workbook,
    list_sheets,
    list_tables,
    read_sheet,
    sheet_exists,
    sheet_kind,
    table_info,
    write_sheet,
)
from pyhandlexl._multi_table import SCHEMA_SHEET

WHEN = dt.datetime(2026, 1, 5, 12, 30, 15)


# ------------------------------------------------------------------ what to compare


def outcome(call):
    """What a call does — its result, or the error it raises — plus every warning it makes."""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        try:
            result = call()
            if isinstance(result, Table):
                result = (result.data, result.style, result.column_types, result.info)
            verdict = ("ok", result)
        except Exception as error:
            verdict = ("raised", type(error).__name__, same_name(str(error)))
    return verdict, sorted((w.category.__name__, same_name(str(w.message))) for w in caught)


def same_name(text: str) -> str:
    """Messages name the file, and the two runs read different copies of it."""
    return re.sub(r"copy-(True|False)", "copy", text)


def read_everything(path: Path) -> dict[str, object]:
    """Every read function, on every sheet and every table the file might hold."""
    out: dict[str, object] = {}
    out["list_sheets"] = outcome(lambda: list_sheets(path))
    sheets = list_sheets(path) if out["list_sheets"][0][0] == "ok" else []  # type: ignore[index]
    for name in [*sheets, "Nope", SCHEMA_SHEET]:
        out[f"sheet_exists {name}"] = outcome(lambda name=name: sheet_exists(path, name))
        out[f"sheet_kind {name}"] = outcome(lambda name=name: sheet_kind(path, name))
        out[f"read_sheet {name}"] = outcome(lambda name=name: read_sheet(path, name))
        out[f"read_sheet pad {name}"] = outcome(lambda name=name: read_sheet(path, name, pad=True))
        out[f"read_sheet columns {name}"] = outcome(
            lambda name=name: read_sheet(path, name, orientation="columns")
        )
    out["read_sheet active"] = outcome(lambda: read_sheet(path))
    out["list_tables"] = outcome(lambda: list_tables(path))
    tables = list_tables(path) if out["list_tables"][0][0] == "ok" else []  # type: ignore[index]
    for name in [*tables, "Ghost"]:
        out[f"Table.read {name}"] = outcome(lambda name=name: Table.read(path, name))
        out[f"table_info {name}"] = outcome(lambda name=name: table_info(path, name))
    return out


def both_ways(path: Path, tmp_path: Path) -> tuple[dict, dict]:
    """The reads over *path*, once with read-only loads and once with full loads — each on its
    own copy, since a read may repair (and so change) the file."""
    results = []
    for read_only in (True, False):
        copy = tmp_path / f"copy-{read_only}.xlsx"
        shutil.copyfile(path, copy)
        previous, safety.READ_ONLY_READS = safety.READ_ONLY_READS, read_only
        try:
            results.append(read_everything(copy))
        finally:
            safety.READ_ONLY_READS = previous
    return results[0], results[1]


def assert_same(path: Path, tmp_path: Path) -> None:
    fast, full = both_ways(path, tmp_path)
    assert fast.keys() == full.keys()
    for key in fast:
        assert fast[key] == full[key], key


# ---------------------------------------------------------- rewriting a saved file


def rewrite(path: Path, part_pattern: str, transform) -> None:
    """Apply *transform* to the text of every part of *path* whose name matches."""
    with zipfile.ZipFile(path) as source:
        parts = {item.filename: source.read(item.filename) for item in source.infolist()}
    for name in parts:
        if re.fullmatch(part_pattern, name):
            parts[name] = transform(parts[name].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in parts.items():
            target.writestr(name, data)


SHEETS = r"xl/worksheets/sheet\d+\.xml"


def declare_dimension(path: Path, ref: str | None) -> None:
    """Make every sheet declare a wrong size — or none at all."""

    def go(text: str) -> str:
        text = re.sub(r"<dimension[^>]*/>", "", text)
        return (
            text
            if ref is None
            else text.replace("<sheetData", f'<dimension ref="{ref}"/><sheetData', 1)
        )

    rewrite(path, SHEETS, go)


# ------------------------------------------------------------------- the workbooks


def tables_and_grids(path: Path) -> None:
    create_workbook(path)
    create_sheet(path, "D")
    create_sheet(path, "Grid")
    create_sheet(path, "Blank")
    Table(
        data=[[1, 2.5, WHEN], [3, None, dt.timedelta(hours=5)], [5, 6, dt.time(1, 2, 3)]],
        row_labels=["r1", "r2", "r3"],
        column_headers=["a", "b", "c"],
        corner="corner",
        name="T",
    ).create(path, sheet="D")
    Table(data=[[7]], row_labels=["x"], column_headers=["y"], name="U").create(path, sheet="D")
    write_sheet(
        path,
        [["h1", "h2"], [1, True], ["=text", None], [dt.date(2026, 1, 1), "é\nx"]],
        sheet="Grid",
    )


def sparse_grid(path: Path) -> None:
    create_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    ws["C5"], ws["E5"], ws["C9"] = "c5", "e5", "c9"
    ws["A1"] = None
    ws["H2"].fill = PatternFill("solid", start_color="FFFF00")  # styled, empty, far to the right
    ws["B12"].fill = PatternFill("solid", start_color="FFFF00")  # styled, empty, below the data
    wb.save(path)


def every_cell_type(path: Path) -> None:
    create_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    values = [
        True,
        False,
        0,
        -7,
        2**60,
        3.14159,
        1e-9,
        1e21,
        "text",
        "",
        "a\r\nb",
        "  padded  ",
        "é中\U0001f600",
        "x" * 1000,
        WHEN,
        dt.date(2026, 3, 4),
        dt.time(1, 2, 3),
        dt.timedelta(days=2, hours=1),
        dt.timedelta(minutes=-90),
        "#DIV/0!",
        "=1+1",
        None,
    ]
    for row, value in enumerate(values, start=1):
        ws.cell(row=row, column=1, value=value)
        ws.cell(row=row, column=2, value=value)
    ws["C1"] = "=SUM(A1:A3)"
    ws["C3"] = "=A3*2"
    wb.save(path)


def mac_epoch(path: Path) -> None:
    create_workbook(path)
    wb = load_workbook(path)
    wb.epoch = CALENDAR_MAC_1904
    wb.active["A1"] = WHEN
    wb.active["A2"] = dt.date(2026, 1, 1)
    wb.active["A3"] = dt.timedelta(hours=30)
    wb.save(path)


def hidden_merged_and_more(path: Path) -> None:
    create_workbook(path)
    wb = load_workbook(path)
    ws = wb.active
    for row in range(1, 6):
        ws.append([row, row * 2, f"v{row}"])
    ws.merge_cells("A1:B2")
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["B"].hidden = True
    ws["C1"].comment = Comment("a note", "me")
    ws["C2"].hyperlink = "https://example.com"
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    chart_data = wb.create_sheet("Data")
    for row in range(1, 5):
        chart_data.append([row, row * row])
    chart = BarChart()
    chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=4))
    chart_sheet = wb.create_chartsheet("Chart")
    chart_sheet.add_chart(chart)
    wb.defined_names["myname"] = DefinedName("myname", attr_text="Data!$A$1")
    wb.save(path)


def odd_sheet_names(path: Path) -> None:
    create_workbook(path)
    for name in ("with space", "it's", "été", "MiXeD", "a.b", "1234", "=eq"):
        create_sheet(path, name)
        write_sheet(path, [[name, 1]], sheet=name)


def rows_deleted_in_excel(path: Path) -> None:
    tables_and_grids(path)
    wb = load_workbook(path)
    wb["D"].delete_rows(4)  # a row of table T
    wb["Grid"].delete_rows(2, 2)
    wb.save(path)


def columns_deleted_in_excel(path: Path) -> None:
    tables_and_grids(path)
    wb = load_workbook(path)
    wb["D"].delete_cols(3)
    wb.save(path)


def rows_inserted_above_a_table(path: Path) -> None:  # the marker moved: reads must repair it
    tables_and_grids(path)
    wb = load_workbook(path)
    wb["D"].insert_rows(1, 2)
    wb.save(path)


def blank_label_and_header(path: Path) -> None:
    tables_and_grids(path)
    wb = load_workbook(path)
    wb["D"]["A4"] = None
    wb["D"]["C2"] = None
    wb.save(path)


def a_sheet_part_missing_from_the_archive(path: Path) -> None:
    tables_and_grids(path)
    with zipfile.ZipFile(path) as source:
        parts = {i.filename: source.read(i.filename) for i in source.infolist()}
    del parts["xl/worksheets/sheet3.xml"]
    with zipfile.ZipFile(path, "w") as target:
        for name, data in parts.items():
            target.writestr(name, data)


def schema_sheet_missing(path: Path) -> None:
    tables_and_grids(path)
    wb = load_workbook(path)
    del wb[SCHEMA_SHEET]
    wb.save(path)


def a_deleted_table(path: Path) -> None:
    from pyhandlexl import delete_table

    tables_and_grids(path)
    delete_table(path, "U")
    delete_table(path, "T")


def only_grids_no_schema(path: Path) -> None:
    create_workbook(path)
    write_sheet(path, [[1, 2], [3, 4]])
    create_sheet(path, "Second")


def empty_workbook(path: Path) -> None:
    create_workbook(path)


def styled_empty_table_cells(path: Path) -> None:
    tables_and_grids(path)
    wb = load_workbook(path)
    for row in range(8, 14):
        wb["D"].cell(row=row, column=2).fill = PatternFill("solid", start_color="00FF00")
    wb.save(path)


BUILDERS = [
    tables_and_grids,
    sparse_grid,
    every_cell_type,
    mac_epoch,
    hidden_merged_and_more,
    odd_sheet_names,
    rows_deleted_in_excel,
    columns_deleted_in_excel,
    rows_inserted_above_a_table,
    blank_label_and_header,
    schema_sheet_missing,
    a_deleted_table,
    only_grids_no_schema,
    empty_workbook,
    styled_empty_table_cells,
]
DIMENSIONS = {
    "as saved": "keep",
    "declared tiny": "A1",
    "declared absent": None,
    "declared huge": "A1:Z400",
}


@pytest.mark.parametrize("dimension", list(DIMENSIONS), ids=list(DIMENSIONS))
@pytest.mark.parametrize("build", BUILDERS, ids=lambda f: f.__name__)
def test_read_only_and_full_loads_agree(build, dimension, tmp_path):
    path = tmp_path / "awkward.xlsx"
    build(path)
    if DIMENSIONS[dimension] != "keep":
        declare_dimension(path, DIMENSIONS[dimension])
    assert_same(path, tmp_path)


# ---------------------------------------------------------- what read-only must not do


class TestReadOnlyReads:
    def test_it_is_what_reads_really_use(self, tmp_path, monkeypatch):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        opened = []
        real = safety.load_workbook
        monkeypatch.setattr(
            safety,
            "load_workbook",
            lambda *a, **k: opened.append(k.get("read_only")) or real(*a, **k),
        )
        Table.read(path, "T")
        read_sheet(path, "Grid")
        list_tables(path)
        table_info(path, "T")
        sheet_kind(path, "Grid")
        list_sheets(path)
        assert opened and all(opened)

    def test_only_the_sheets_that_are_read_are_parsed(self, tmp_path):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        rewrite(
            path,
            r"xl/worksheets/sheet3\.xml",
            lambda text: text.replace("<sheetData>", "<sheetData><<<"),
        )
        # sheet3 is 'Grid': broken XML nobody asked to read
        assert read_sheet(path, "D")[0] != []
        assert Table.read(path, "T").data.row_labels == ["r1", "r2", "r3"]
        assert list_tables(path) == ["T", "U"] and list_sheets(path)[0] == "Sheet"
        with pytest.raises(InvalidFileError):
            read_sheet(path, "Grid")

    def test_a_damaged_sheet_that_is_read_is_an_invalid_file(self, tmp_path):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        rewrite(path, r"xl/worksheets/sheet2\.xml", lambda text: text.replace("</row>", "", 1))
        with pytest.raises(InvalidFileError, match="not a readable"):
            Table.read(path, "T")

    def test_a_missing_table_is_still_table_not_found_not_an_invalid_file(self, tmp_path):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        with pytest.raises(TableNotFoundError, match="no table named 'Ghost'"):
            Table.read(path, "Ghost")

    @pytest.mark.parametrize("content", [b"", b"not a zip", b"PK\x03\x04truncated"])
    def test_a_file_that_is_not_a_workbook_fails_at_once_not_after_retries(
        self, tmp_path, monkeypatch, content
    ):
        slept = []
        monkeypatch.setattr(safety, "sleep", slept.append)
        path = tmp_path / "junk.xlsx"
        path.write_bytes(content)
        for call in (
            lambda: list_sheets(path),
            lambda: read_sheet(path),
            lambda: list_tables(path),
        ):
            with pytest.raises(InvalidFileError):
                call()
        assert slept == []  # BadZipFile is retried as "transient": it must not be, for this

    def test_a_folder_or_a_missing_file(self, tmp_path):
        with pytest.raises(InvalidFileError):
            list_sheets(tmp_path)
        with pytest.raises(FileNotFoundError):
            list_sheets(tmp_path / "nope.xlsx")

    def test_a_wrong_extension_is_an_invalid_file(self, tmp_path):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        wrong = tmp_path / "wb.txt"
        shutil.copyfile(path, wrong)
        with pytest.raises(InvalidFileError):
            read_sheet(wrong)

    def test_a_locked_file_is_a_file_locked_error(self, tmp_path, monkeypatch):
        from pyhandlexl import FileLockedError

        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        monkeypatch.setattr(safety, "sleep", lambda _s: None)

        def locked(*args, **kwargs):
            raise PermissionError("open in Excel")

        monkeypatch.setattr(safety, "load_workbook", locked)
        with pytest.raises(FileLockedError):
            read_sheet(path, "Grid")

    def test_reads_leave_the_file_free_to_replace_or_delete(self, tmp_path):
        # a read-only workbook keeps the archive open until it is closed; on Windows an open
        # file can't be deleted, so every read function must close it
        for read in (
            lambda p: read_sheet(p, "Grid"),
            lambda p: Table.read(p, "T"),
            lambda p: table_info(p, "T"),
            lambda p: list_tables(p),
            lambda p: list_sheets(p),
            lambda p: sheet_kind(p, "Grid"),
            lambda p: sheet_exists(p, "Grid"),
        ):
            path = tmp_path / "wb.xlsx"
            tables_and_grids(path)
            read(path)
            path.unlink()  # raises PermissionError on Windows if a handle was left open

    def test_reads_do_not_change_the_file(self, tmp_path):
        path = tmp_path / "wb.xlsx"
        tables_and_grids(path)
        before = path.read_bytes()
        for read in (
            lambda: read_sheet(path, "Grid"),
            lambda: Table.read(path, "T"),
            lambda: table_info(path, "T"),
            lambda: list_tables(path),
            lambda: sheet_kind(path, "Blank"),
        ):
            read()
        assert path.read_bytes() == before

    def test_an_empty_sheet_read_then_appended_to_lands_on_the_first_row(self, tmp_path):
        # reading must not create cells: opening an empty sheet once made the next append land
        # on row 2
        from pyhandlexl import append_rows

        path = tmp_path / "wb.xlsx"
        create_workbook(path)
        assert sheet_kind(path, "Sheet") == "empty" and read_sheet(path) == []
        append_rows(path, [["a"], ["b"]])
        assert read_sheet(path) == [["a"], ["b"]]

    def test_the_repairing_reads_still_repair_and_save(self, tmp_path):
        moved = tmp_path / "moved.xlsx"
        rows_inserted_above_a_table(moved)
        before = moved.read_bytes()
        assert Table.read(moved, "T").data.row_labels == ["r1", "r2", "r3"]  # found by its marker
        assert moved.read_bytes() != before  # ... and the corrected position was saved
        rebuilt = tmp_path / "rebuilt.xlsx"
        schema_sheet_missing(rebuilt)
        with pytest.warns(UserWarning, match="schema"):
            assert list_tables(rebuilt) == ["T", "U"]
        assert SCHEMA_SHEET in load_workbook(rebuilt).sheetnames  # rebuilt and saved


def test_one_table_from_a_workbook_of_big_sheets_does_not_parse_the_big_sheets(tmp_path):
    path = tmp_path / "big.xlsx"
    create_workbook(path)
    for i in range(3):
        create_sheet(path, f"G{i}")
        write_sheet(path, [[r * c for c in range(8)] for r in range(600)], sheet=f"G{i}")
    create_sheet(path, "Small")
    Table(data=[[1, 2]], row_labels=["a"], column_headers=["x", "y"], name="S").create(
        path, sheet="Small"
    )
    # if a big sheet were parsed this would fail: they are unreadable
    rewrite(
        path,
        r"xl/worksheets/sheet[2-4]\.xml",
        lambda text: text.replace("<sheetData>", "<sheetData><<<"),
    )
    assert Table.read(path, "S").data.rows == [[1, 2]]
    assert table_info(path, "S").n_rows == 1 and list_tables(path) == ["S"]
