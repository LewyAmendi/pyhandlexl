"""Internal support for multiple named tables on one worksheet.

Not part of the public API. A workbook-wide schema sheet (``_SCHEMA_SHEET``)
records where each named table lives. Each table is also self-describing on
the sheet: its very first cell holds the literal marker ``"TABLE NAME"`` and
the cell to its right holds the table's name. The schema is a fast-path cache;
the marker is how a lookup verifies the cache is still correct, and how it
recovers if a table has moved.

Layout of one named table, anchored at (anchor_row, anchor_col):

    anchor_row      TABLE NAME | <name>
    anchor_row + 1  <corner>   | <column headers ...>
    anchor_row + 2..           <row label> | <data ...>

Tables on the same sheet stack left-to-right with exactly one empty column
between them; they always start at row 1.
"""

from __future__ import annotations

import json
import warnings
from collections.abc import Iterator
from contextlib import contextmanager
from contextvars import ContextVar
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from typing import Any, Literal, cast

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pyhandlexl import _safety
from pyhandlexl._merge import Snapshot
from pyhandlexl._safety import damage_is_invalid, safe_load_readonly
from pyhandlexl.column_type import ColumnType
from pyhandlexl.errors import (
    FormulaWarning,
    MalformedTableError,
    SchemaRebuiltWarning,
    SheetKindError,
    SheetNameError,
    TableExistsError,
    TableNotFoundError,
)
from pyhandlexl.style import TableStyle
from pyhandlexl.validate import MAX_CELL_CHARS

SCHEMA_SHEET = "_pyhandlexl_tables"
MARKER = "TABLE NAME"


def is_schema_name(name: str) -> bool:
    """Whether *name* is the reserved schema sheet's name, or a lookalike of it.

    Excel treats sheet names case-insensitively, and openpyxl silently renames a
    would-be duplicate (``_PYHANDLEXL_TABLES`` becomes ``_PYHANDLEXL_TABLES1``) —
    so a lookalike is never harmless: it would either masquerade as the schema
    sheet or push the real one off its exact name. Surrounding whitespace is
    ignored too, so ``"_pyhandlexl_tables "`` can't slip through.

    Every public function that takes a sheet name asks this first, so it is also where a
    name that isn't a string is turned away (as ``SheetNameError``, like ``write_sheet``
    does) instead of failing on ``.strip()``.
    """
    if not isinstance(name, str):
        raise SheetNameError(f"sheet name must be a string, got {type(name).__name__}")
    return name.strip().lower() == SCHEMA_SHEET


def reserved_error(name: str, what: str) -> SheetKindError:
    """The error for using the reserved schema sheet name (or a lookalike) for *what*."""
    if name == SCHEMA_SHEET:
        return SheetKindError(f"{SCHEMA_SHEET!r} is reserved and {what}")
    return SheetKindError(
        f"{name!r} is too close to the reserved sheet name {SCHEMA_SHEET!r} (Excel ignores "
        f"capitalisation, so the two would collide) and {what}"
    )


def case_clash(sheetnames: list[str], name: str) -> str | None:
    """An existing sheet that differs from *name* only by capitalisation, if any.

    Excel (and openpyxl) treat ``Data`` and ``data`` as the same sheet name. A
    creation path must refuse such a name rather than let openpyxl quietly
    invent ``data1`` instead.
    """
    for existing in sheetnames:
        if existing != name and existing.lower() == name.lower():
            return existing
    return None


_SCHEMA_TAB_COLOR = "FF0000"
# warnings.warn depth from rebuild_schema to the user's own line: rebuild_schema (1) <-
# load_schema (2) <- the public function that called it (3) <- the user's code (4).
_CALLER_LEVEL = 4
# A public function that does its work by calling another public one (import_csv_to_xl calls
# write_sheet) adds one frame; it says so here so the warning still names the user's line.
EXTRA_FRAMES: ContextVar[int] = ContextVar("pyhandlexl_extra_frames", default=0)


def caller_level() -> int:
    """The ``stacklevel`` that makes a warning name the user's own line."""
    return _CALLER_LEVEL + EXTRA_FRAMES.get()


_SCHEMA_WARNING = (
    "pyhandlexl-managed — do not edit or delete by hand.\n\n"
    "This sheet tracks every named table's location, size, style, "
    "column-type restrictions, and creation/modification times. If it's "
    "deleted, pyhandlexl automatically rebuilds it next time a table is "
    "read, written, or listed, by scanning the workbook for table markers "
    "— position and size come back exact, but style is only a best-effort "
    "reconstruction; column types are only inferred from what each column "
    "currently holds (so a restriction can be lost, or invented for a "
    "column that was deliberately left unrestricted); and the creation/"
    "modification times cannot be recovered at all (both come back "
    "unknown)."
)
_SCHEMA_HEADER = (
    "name",
    "sheet",
    "anchor_row",
    "anchor_col",
    "n_rows",
    "n_cols",
    "style",
    "column_types",
    "created_at",
    "modified_at",
)


def _to_label(value: object) -> str:
    """Coerce a header/label/corner cell to ``str``; an empty cell becomes ``""``."""
    return "" if value is None else str(value)


@dataclass(frozen=True)
class TableEntry:
    """One row of the schema: where a named table lives, how big it is, and how it looks."""

    name: str
    sheet: str
    anchor_row: int
    anchor_col: int
    n_rows: int  # data rows (not counting the marker or header row)
    n_cols: int  # data columns (not counting the label column)
    style: TableStyle
    column_types: list[ColumnType]  # one per data column, in order
    created_at: datetime | None  # None only if recovered by rebuild_schema
    modified_at: datetime | None  # None only if recovered by rebuild_schema

    @property
    def height(self) -> int:
        """Total rows on the sheet: marker row + header row + data rows."""
        return self.n_rows + 2

    @property
    def width(self) -> int:
        """Total columns on the sheet: label column + data columns."""
        return self.n_cols + 1


def _style_to_json(style: TableStyle) -> str:
    return json.dumps(asdict(style), separators=(",", ":"))


def _style_from_json(value: str) -> TableStyle:
    return TableStyle(**json.loads(value))


# One letter per column type, for tables so wide that the readable JSON list wouldn't fit in a
# single cell (over ~5,000 columns): a cell holds at most MAX_CELL_CHARS characters and a longer
# string is silently truncated, which used to leave a schema that no longer parsed — and locked
# every operation on the whole workbook out with a JSONDecodeError.
_TYPE_LETTER = {
    ColumnType.ANY: "a",
    ColumnType.TEXT: "t",
    ColumnType.NUMBER: "n",
    ColumnType.BOOLEAN: "b",
    ColumnType.DATE: "d",
    ColumnType.TIME: "m",
    ColumnType.DURATION: "u",
}
_LETTER_TYPE = {letter: column_type for column_type, letter in _TYPE_LETTER.items()}
_COMPACT_PREFIX = "~"


def _column_types_to_json(column_types: list[ColumnType]) -> str:
    text = json.dumps([ct.value for ct in column_types], separators=(",", ":"))
    if len(text) <= MAX_CELL_CHARS:
        return text
    return _COMPACT_PREFIX + "".join(_TYPE_LETTER[ct] for ct in column_types)


def _column_types_from_json(value: str) -> list[ColumnType]:
    if value.startswith(_COMPACT_PREFIX):
        return [_LETTER_TYPE[letter] for letter in value[len(_COMPACT_PREFIX) :]]
    return [ColumnType(v) for v in json.loads(value)]


def _datetime_to_str(value: datetime | None) -> str:
    return value.isoformat() if value is not None else ""


def _datetime_from_str(value: str) -> datetime | None:
    return datetime.fromisoformat(value) if value else None


# ---------------------------------------------------------------- the schema


def load_schema(workbook: Workbook) -> dict[str, TableEntry]:
    """Read the schema sheet from an open workbook.

    If the reserved sheet is missing, automatically reconstructs it by
    scanning every other sheet for table markers (see :func:`rebuild_schema`)
    — empty only if that scan also finds nothing, meaning the workbook
    genuinely has no tables yet. Doesn't persist anything itself; the caller
    decides whether to save the reconstruction so the scan isn't repeated.
    """
    if SCHEMA_SHEET not in workbook.sheetnames:
        return rebuild_schema(workbook)
    return _parse_schema(workbook[SCHEMA_SHEET])


@contextmanager
def cheap_schema(path: Any) -> Iterator[tuple[Workbook, dict[str, TableEntry]] | None]:
    """A read-only workbook and its schema, for a read that needs nothing else — or ``None``.

    ``None`` means the read can't be done this way and has to load the whole workbook: the
    schema sheet is missing (so it must be rebuilt, which scans every sheet and saves), or
    reads have been switched to full loads. Anything a read might have to *repair* is left to
    that full path, which is unchanged. The workbook is closed when the block ends, so do
    any fallback after it, not inside.
    """
    if not _safety.READ_ONLY_READS:
        yield None
        return
    workbook = safe_load_readonly(path)
    try:
        with damage_is_invalid(path):
            if SCHEMA_SHEET not in workbook.sheetnames:
                yield None
            else:
                yield workbook, _parse_schema(workbook[SCHEMA_SHEET])
    finally:
        workbook.close()


def all_rows(ws: Any, *, values_only: bool = False, min_row: int | None = None) -> Any:
    """Every row of *ws* (from *min_row*, if given), however the file records its size.

    A read-only sheet trusts the ``<dimension>`` the file declares, and some writers declare
    it wrongly (or not at all), which would silently cut rows off; resetting it makes the
    sheet read to its real end. *min_row* is passed on only when given: on a loaded sheet,
    ``iter_rows()`` of an empty one yields nothing, but naming a row makes openpyxl create
    cells for it — so asking for "row 1" would change the sheet, and the next ``append``
    would land a row too low.
    """
    reset = getattr(ws, "reset_dimensions", None)
    if reset is not None:
        reset()
    if min_row is None:
        return ws.iter_rows(values_only=values_only)
    return ws.iter_rows(min_row=min_row, values_only=values_only)


def _parse_schema(ws: Any) -> dict[str, TableEntry]:
    entries: dict[str, TableEntry] = {}
    for row in all_rows(ws, values_only=True, min_row=2):
        if not row or row[0] is None:
            continue
        # Schemas written by older versions have fewer columns (7 before column
        # types, 8 before creation/modification times) — read what's there and
        # default the rest; the next save upgrades the sheet.
        cells: list[Any] = [*row[:10], *([None] * (10 - len(row[:10])))]
        (
            name,
            sheet,
            anchor_row,
            anchor_col,
            n_rows,
            n_cols,
            style_json,
            column_types_json,
            created_at,
            modified_at,
        ) = cells
        entries[name] = TableEntry(
            name,
            sheet,
            int(anchor_row),
            int(anchor_col),
            int(n_rows),
            int(n_cols),
            _style_from_json(style_json) if style_json else TableStyle.DEFAULT,
            _column_types_from_json(column_types_json)
            if column_types_json
            else [ColumnType.ANY] * int(n_cols),
            _datetime_from_str(created_at),
            _datetime_from_str(modified_at),
        )
    return entries


def save_schema(workbook: Workbook, entries: dict[str, TableEntry]) -> None:
    """Replace the schema sheet's contents in an open workbook.

    Marks the sheet with a red tab color and a warning comment on its
    first cell every time — a visible "don't touch this" for anyone
    browsing the workbook by hand, not just a mention in the docs.
    """
    if SCHEMA_SHEET in workbook.sheetnames:
        ws = workbook[SCHEMA_SHEET]
        ws.delete_rows(1, ws.max_row)
    else:
        clash = case_clash(workbook.sheetnames, SCHEMA_SHEET)
        if clash is not None:  # only a hand-edited file can get here; creating it would
            raise SheetKindError(  # make openpyxl name the schema sheet "..._tables1"
                f"sheet {clash!r} collides with the reserved sheet name {SCHEMA_SHEET!r} "
                "(Excel ignores capitalisation), so the schema sheet can't be created — "
                "rename or remove that sheet"
            )
        ws = workbook.create_sheet(SCHEMA_SHEET)
    ws.sheet_properties.tabColor = _SCHEMA_TAB_COLOR
    ws.append(list(_SCHEMA_HEADER))
    ws["A1"].comment = Comment(_SCHEMA_WARNING, "pyhandlexl")
    for e in entries.values():
        row = [
            e.name,
            e.sheet,
            e.anchor_row,
            e.anchor_col,
            e.n_rows,
            e.n_cols,
            _style_to_json(e.style),
            _column_types_to_json(e.column_types),
            _datetime_to_str(e.created_at),
            _datetime_to_str(e.modified_at),
        ]
        ws.append(row)
        for column, value in enumerate(row, start=1):
            if isinstance(value, str) and value.startswith("="):  # a table named "=x"
                store_text(ws.cell(row=ws.max_row, column=column), value)


# ------------------------------------------------------------- schema rebuild


def _rgb_to_hex(color: Any, default: str) -> str:
    """Best-effort: an openpyxl Color's ``.rgb`` as a plain 6-digit hex string.

    Falls back to *default* for anything that isn't a real ARGB string — a
    theme-palette reference, an indexed color, or no color at all.
    """
    rgb = getattr(color, "rgb", None) if color is not None else None
    if isinstance(rgb, str) and len(rgb) == 8 and all(c in "0123456789abcdefABCDEF" for c in rgb):
        return rgb[2:]  # drop the leading alpha byte
    return default


def _find_markers(workbook: Workbook) -> list[tuple[str, int, int, str]]:
    """Every (sheet, row, col, name) marker found anywhere in the workbook."""
    found = []
    for sheet_name in workbook.sheetnames:
        if sheet_name == SCHEMA_SHEET:
            continue
        ws = workbook[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == MARKER:
                    name = ws.cell(row=cell.row, column=cell.column + 1).value
                    if isinstance(name, str) and name:
                        found.append((sheet_name, cell.row, cell.column, name))
    return found


def _genuine_markers(
    workbook: Workbook, markers: list[tuple[str, int, int, str]]
) -> list[tuple[str, int, int, str]]:
    """Drop cells that read ``TABLE NAME`` but sit inside a table rather than start one.

    A row label (or header, or data cell) that happens to say ``TABLE NAME`` next to a
    non-blank cell looks exactly like a marker. Left in, it becomes a phantom table and
    also cuts the real table's width short, so the real one stops reading. Tables on a
    sheet all start on the same row (row 1, unless someone inserted rows above them all
    by hand), so the markers on that top row are taken as real first; any other marker
    counts only if it lies outside every table already accepted.
    """
    by_sheet: dict[str, list[tuple[str, int, int, str]]] = {}
    for marker in markers:
        by_sheet.setdefault(marker[0], []).append(marker)

    def region(
        ws: Worksheet, row: int, col: int, next_col: int | None
    ) -> tuple[int, int, int, int]:
        n_cols = _infer_width(ws, row, col, next_col)
        n_rows = _infer_height(ws, row, col, n_cols)
        return row, row + n_rows + 1, col, col + n_cols

    kept: set[tuple[str, int, int, str]] = set()
    for sheet_name, found in by_sheet.items():
        ws = workbook[sheet_name]
        top = min(row for _, row, _, _ in found)
        first = sorted((m for m in found if m[1] == top), key=lambda m: m[2])
        rest = sorted((m for m in found if m[1] != top), key=lambda m: (m[1], m[2]))
        regions = []
        for i, (_, row, col, _) in enumerate(first):
            kept.add(first[i])
            next_col = first[i + 1][2] if i + 1 < len(first) else None
            regions.append(region(ws, row, col, next_col))
        for marker in rest:
            _, row, col, _ = marker
            if any(r0 <= row <= r1 and c0 <= col <= c1 for r0, r1, c0, c1 in regions):
                continue
            kept.add(marker)
            regions.append(region(ws, row, col, None))
    return [m for m in markers if m in kept]


def _infer_width(
    ws: Worksheet, anchor_row: int, anchor_col: int, next_anchor_col: int | None
) -> int:
    """Data-column count for the table anchored at (anchor_row, anchor_col).

    Bounded exactly by the next table's anchor column when there is one on
    the same sheet — no guessing needed there. Otherwise (the rightmost
    table on its sheet) the last non-blank header cell is the edge: nothing
    legitimate is ever placed past a table's real content in that row, and
    duplicate headers (including a blank one) are already blocked, so at
    most one blank header cell can appear *within* the real span anyway.
    """
    header_row = anchor_row + 1
    limit = (next_anchor_col - 1) if next_anchor_col is not None else ws.max_column + 1
    last_real = anchor_col
    for col in range(anchor_col + 1, limit):
        if ws.cell(row=header_row, column=col).value is not None:
            last_real = col
    return last_real - anchor_col


def _infer_height(ws: Worksheet, anchor_row: int, anchor_col: int, n_cols: int) -> int:
    """Data-row count below the header row of the table anchored at (anchor_row, anchor_col).

    A row counts as part of the table if *any* cell across its width (label
    plus every data column) is non-blank. A row that's entirely blank is
    ambiguous on its own — but a duplicate blank row label is blocked the
    same as any duplicate label, so at most one such row can legitimately
    exist. One is tolerated by peeking at the row after it before deciding
    the table actually ended before it.
    """
    header_row = anchor_row + 1
    last_real = header_row
    pending_blank: int | None = None
    row = header_row + 1
    while row <= ws.max_row + 1:
        blank = all(
            ws.cell(row=row, column=anchor_col + c).value is None for c in range(n_cols + 1)
        )
        if not blank:
            last_real = row
            pending_blank = None
        elif pending_blank is None:
            pending_blank = row
        else:
            break
        row += 1
    return last_real - header_row


def _infer_style(
    ws: Worksheet, anchor_row: int, anchor_col: int, n_rows: int, n_cols: int
) -> TableStyle:
    """Best-effort TableStyle read back from a table's actual painted cells.

    Geometry (from _infer_width/_infer_height) is exact; this is not — it's
    a reconstruction from whatever the cells currently look like, which can
    drift from what was originally set if formatting was hand-edited.
    """
    header_row = anchor_row + 1
    corner = ws.cell(row=header_row, column=anchor_col)

    header_fill = _rgb_to_hex(corner.fill.fgColor, "") if corner.fill.fill_type == "solid" else ""
    header_font_name = corner.font.name or TableStyle.DEFAULT.header_font_name
    header_font_size = (
        int(corner.font.size) if corner.font.size else TableStyle.DEFAULT.header_font_size
    )
    header_bold = bool(corner.font.bold)
    header_font_color = _rgb_to_hex(corner.font.color, "000000")

    top = corner.border.top
    border_color = _rgb_to_hex(top.color, "000000") if top is not None and top.style else ""

    if n_rows == 0 or n_cols == 0:
        return TableStyle(
            header_font_name=header_font_name,
            header_font_size=header_font_size,
            header_font_color=header_font_color,
            header_bold=header_bold,
            header_fill=header_fill,
            border_color=border_color,
        )

    data_row1 = ws.cell(row=header_row + 1, column=anchor_col + 1)
    data_font_name = data_row1.font.name or TableStyle.DEFAULT.data_font_name
    data_font_size = (
        int(data_row1.font.size) if data_row1.font.size else TableStyle.DEFAULT.data_font_size
    )
    data_font_color = _rgb_to_hex(data_row1.font.color, "000000")

    band_fill = ""
    if n_rows >= 2:
        data_row2 = ws.cell(row=header_row + 2, column=anchor_col + 1)
        fill1 = (
            _rgb_to_hex(data_row1.fill.fgColor, "") if data_row1.fill.fill_type == "solid" else ""
        )
        fill2 = (
            _rgb_to_hex(data_row2.fill.fgColor, "") if data_row2.fill.fill_type == "solid" else ""
        )
        band_fill = (
            fill2 if fill2 and fill2 != fill1 else (fill1 if fill1 and fill1 != fill2 else "")
        )

    return TableStyle(
        header_font_name=header_font_name,
        header_font_size=header_font_size,
        header_font_color=header_font_color,
        header_bold=header_bold,
        header_fill=header_fill,
        data_font_name=data_font_name,
        data_font_size=data_font_size,
        data_font_color=data_font_color,
        band_fill=band_fill,
        border_color=border_color,
    )


_INFERABLE_TYPES = (
    ColumnType.TEXT,
    ColumnType.NUMBER,
    ColumnType.BOOLEAN,
    ColumnType.DATE,
    ColumnType.TIME,
    ColumnType.DURATION,
)


def _infer_column_types(
    ws: Worksheet, anchor_row: int, anchor_col: int, n_rows: int, n_cols: int
) -> list[ColumnType]:
    """Best-effort ColumnType per data column, read back from what it holds.

    A column whose non-blank values all belong to one Excel-native type
    comes back restricted to it; a column that's empty, or mixes types,
    comes back ``ColumnType.ANY``. Every value is exactly one native type
    (``bool`` is never mistaken for a number, ``datetime`` and ``date`` are
    both dates), so at most one restriction can match — there's no
    ambiguity between two restrictions, only between a restriction and
    ``ANY``. That one is real: a column that was deliberately left
    unrestricted but happens to hold a single type is indistinguishable
    from one that was restricted, so it comes back restricted.

    Whatever this returns always admits the data currently in the table,
    so writing it straight back never trips ``ColumnTypeError``.
    """
    first_data_row = anchor_row + 2
    inferred = []
    for j in range(n_cols):
        present = [
            value
            for i in range(n_rows)
            if (value := ws.cell(row=first_data_row + i, column=anchor_col + 1 + j).value)
            is not None
        ]
        matches = (
            [ct for ct in _INFERABLE_TYPES if all(ct.allows(v) for v in present)] if present else []
        )
        inferred.append(matches[0] if len(matches) == 1 else ColumnType.ANY)
    return inferred


def rebuild_schema(workbook: Workbook) -> dict[str, TableEntry]:
    """Reconstruct the schema by scanning every sheet for table markers.

    Used automatically by :func:`load_schema` when the reserved sheet is
    missing. Geometry is exact — see :func:`_infer_width`/:func:`_infer_height`
    for why the marker-based scan can pin it down precisely. A table's style
    is read back from its cells on a best-effort basis (see
    :class:`~pyhandlexl.errors.SchemaRebuiltWarning`). A column's
    :class:`~pyhandlexl.column_type.ColumnType` restriction is *inferred*
    from the values it currently holds (see :func:`_infer_column_types`) —
    nothing in a cell says "this column is restricted," only what happens
    to be in it, so this can lose a restriction (an empty or mixed column
    comes back ``ANY``) or invent one (a column deliberately left
    unrestricted, but holding one type, comes back restricted to it).
    Nothing in a cell records when the table was created or last modified
    either, so both come back ``None`` rather than a guessed value.

    A name claimed by more than one marker can't be safely resolved — that
    table is left out of the result (a warning names it) rather than
    guessing which location is real; every unambiguous table is unaffected.
    Emits nothing, and returns an empty dict, if the scan finds no markers
    at all (the ordinary case: the workbook genuinely has no tables yet).
    """
    markers = _genuine_markers(workbook, _find_markers(workbook))
    if not markers:
        return {}

    by_name: dict[str, list[tuple[str, int, int]]] = {}
    for sheet_name, row, col, name in markers:
        by_name.setdefault(name, []).append((sheet_name, row, col))
    conflicted = {name for name, locations in by_name.items() if len(locations) > 1}
    for name in conflicted:
        warnings.warn(
            f"table name {name!r} was found at more than one location while rebuilding "
            f"the schema ({by_name[name]}) — skipped, since it's not safe to guess which "
            "one is real; it won't be reachable by name until this is resolved by hand",
            SchemaRebuiltWarning,
            stacklevel=caller_level(),
        )

    by_sheet: dict[str, list[tuple[int, int, str]]] = {}
    for sheet_name, row, col, name in markers:
        if name not in conflicted:
            by_sheet.setdefault(sheet_name, []).append((row, col, name))

    entries: dict[str, TableEntry] = {}
    for sheet_name, sheet_markers in by_sheet.items():
        sheet_markers.sort(key=lambda m: m[1])
        ws = workbook[sheet_name]
        for i, (anchor_row, anchor_col, name) in enumerate(sheet_markers):
            next_col = sheet_markers[i + 1][1] if i + 1 < len(sheet_markers) else None
            n_cols = _infer_width(ws, anchor_row, anchor_col, next_col)
            n_rows = _infer_height(ws, anchor_row, anchor_col, n_cols)
            style = _infer_style(ws, anchor_row, anchor_col, n_rows, n_cols)
            column_types = _infer_column_types(ws, anchor_row, anchor_col, n_rows, n_cols)
            entries[name] = TableEntry(
                name,
                sheet_name,
                anchor_row,
                anchor_col,
                n_rows,
                n_cols,
                style,
                column_types,
                created_at=None,
                modified_at=None,
            )

    warnings.warn(
        f"the {SCHEMA_SHEET} schema sheet was missing and has been rebuilt by scanning "
        f"the workbook; found {len(entries)} table(s): {sorted(entries)}. Sizes are "
        "exact; styles are a best-effort reconstruction from the cells themselves and "
        "may not exactly match what was originally set; column types are inferred from "
        "the values each column currently holds, so a restriction can be lost (an empty "
        "or mixed column comes back ColumnType.ANY) or invented (an unrestricted column "
        "holding one type comes back restricted to it); creation and modification times "
        "can't be recovered and come back as None.",
        SchemaRebuiltWarning,
        stacklevel=caller_level(),
    )
    return entries


# ------------------------------------------------------------- region access


def read_block(
    ws: Any, top: int, left: int, height: int, width: int
) -> tuple[list[list[object]], list[str]]:
    """A block of cells as ``(values, formulas)``: the values, and the coordinates of the cells
    that hold a formula (whose value is then the formula's text). One pass, and it works on a
    read-only sheet as well as a loaded one — ``ws.cell()`` would re-read a read-only sheet
    from the top for every cell."""
    values: list[list[object]] = []
    formulas: list[str] = []
    reset = getattr(ws, "reset_dimensions", None)
    if reset is not None:  # a read-only sheet: don't trust the size the file declares
        reset()
    rows = ws.iter_rows(
        min_row=top, max_row=top + height - 1, min_col=left, max_col=left + width - 1
    )
    for i, row in enumerate(rows):
        line: list[object] = []
        for j, cell in enumerate(row):
            line.append(cell.value)
            if cell.data_type == "f":
                formulas.append(f"{get_column_letter(left + j)}{top + i}")
        # Ask for `width` cells and you get them; a short row is padded.
        line.extend([None] * (width - len(line)))
        values.append(line)
    # A read-only sheet stops at the last row it holds, however many were asked for (a row
    # deleted in Excel leaves the block longer than the data): the rest are blank rows.
    values.extend([None] * width for _ in range(height - len(values)))
    return values, formulas


def read_region(ws: Any, top: int, left: int, height: int, width: int) -> list[list[object]]:
    return read_block(ws, top, left, height, width)[0]


def marker_matches(ws: Any, entry: TableEntry) -> bool:
    """Whether *entry*'s recorded position still holds its marker and name."""
    (marker, name), *_ = read_region(ws, entry.anchor_row, entry.anchor_col, 1, 2)
    return bool(marker == MARKER and name == entry.name)


def store_text(cell: Any, text: str) -> None:
    """Store *text* in *cell* as text, even if it starts with ``=``.

    openpyxl turns any string starting with ``=`` into a formula, which Excel would try to
    calculate (showing ``#NAME?``, or rejecting the file, if it isn't a valid one).
    pyhandlexl doesn't support formulas, so text stays text: the cell is stored as a
    string, with Excel's *quote prefix* flag (the invisible apostrophe Excel adds when you
    type ``'=1+1``) so it stays text if someone edits it.
    """
    cell.value = text
    if cell.data_type == "f":
        cell.data_type = "s"
        if cell._style is None:
            cell._style = StyleArray()
        cell._style.quotePrefix = 1


def formula_warning(where: str, refs: list[str]) -> FormulaWarning:
    """The warning for *refs*, the formula cells found in *where*."""
    shown = ", ".join(refs[:5]) + (f" and {len(refs) - 5} more" if len(refs) > 5 else "")
    return FormulaWarning(
        f"{where} holds {len(refs)} formula cell{'s' if len(refs) != 1 else ''} ({shown}). "
        "pyhandlexl doesn't support formulas: they are read as their text, and writing that "
        "text back stores it as text, replacing the formula. Use openpyxl directly to work "
        "with formulas."
    )


def write_region(ws: Worksheet, top: int, left: int, grid: list[list[object]]) -> None:
    """Write *grid* with its top-left cell at (*top*, *left*).

    pyhandlexl doesn't support formulas, so a string is always stored as text, even one that
    starts with ``=`` (see :func:`store_text`).
    """
    # ws.cell(..., value=None) is a no-op in openpyxl (it means "no value given",
    # not "clear it"), so None must be assigned via .value directly.
    for i, row in enumerate(grid):
        for j, value in enumerate(row):
            cell = cast(Any, ws.cell(row=top + i, column=left + j))  # ._style isn't in the stubs
            # A cell that keeps its style across a write may still carry what its *previous*
            # value gave it: a date format left on a number would read back as a date, and a
            # quote prefix left on a formula would stop it being one. Assigning a date, time
            # or duration sets its number format again; store_text sets the quote prefix.
            if cell._style is not None:
                if cell._style.numFmtId:
                    cell._style.numFmtId = 0
                if cell._style.quotePrefix:
                    cell._style.quotePrefix = 0
            if isinstance(value, str) and value.startswith("="):
                store_text(cell, value)
            else:
                cell.value = cast(Any, value)


def clear_region(ws: Worksheet, top: int, left: int, height: int, width: int) -> None:
    # Resetting only .value leaves a painted cell's style record behind, and
    # a cell with an explicit style but no value still counts as "used" —
    # Excel's own dimension tracking works the same way. Reassigning font/
    # fill/border to fresh default objects doesn't help either: openpyxl
    # still allocates a style entry for them. Only clearing the private
    # ._style attribute directly (openpyxl has no public API for this)
    # makes a cell truly indistinguishable from one that was never touched.
    for r in range(top, top + height):
        for c in range(left, left + width):
            cell = cast(Any, ws.cell(row=r, column=c))  # ._style isn't in the stubs
            cell.value = None
            cell._style = None


def clear_for_rewrite(
    ws: Worksheet, entry: TableEntry, height: int, width: int, row_limit: int, col_limit: int
) -> None:
    """Clear the part of a table's block that is about to be rewritten *and* repainted.

    Like :func:`clear_region` over the block of ``height`` x ``width`` cells at *entry*'s
    position, except that the cells :func:`stable_extent` says keep their look — local rows
    below *row_limit* and local columns below *col_limit* — are left as they are, style and
    all. (:func:`write_region` still resets their number formats.) The marker row is
    never among them.
    """
    top, left = entry.anchor_row, entry.anchor_col
    for r in range(top, top + height):
        local_row = r - top - 1  # the marker row is -1, the header row 0
        first = col_limit if 0 <= local_row < row_limit else 0
        for c in range(left + first, left + width):
            cell = cast(Any, ws.cell(row=r, column=c))
            cell.value = None
            cell._style = None


def stable_extent(
    entry: TableEntry, new_n_rows: int, new_n_cols: int, new_style: TableStyle
) -> tuple[int, int]:
    """How much of *entry*'s block keeps exactly the look it has if it becomes the given size.

    Returns ``(row_limit, col_limit)``: the cell at local row ``r`` (the header row is 0) and
    local column ``c`` (the label column is 0) is unchanged when ``r < row_limit`` and
    ``c < col_limit``. Styles belong to positions, not to data — banding depends only on
    a row's parity, and a border only on whether a cell is in the header, the label column,
    the last row or the last column — so a cell's look changes only if the style does, or if
    it was, or now is, on the last row or column. ``(0, 0)`` means repaint everything.
    """
    if entry.style != new_style:
        return 0, 0
    rows = new_n_rows + 1 if new_n_rows == entry.n_rows else min(new_n_rows, entry.n_rows)
    cols = new_n_cols + 1 if new_n_cols == entry.n_cols else min(new_n_cols, entry.n_cols)
    return rows, cols


# ----------------------------------------------------------------- painting


_Weight = Literal["thick", "medium", "thin"]
SheetKind = Literal["table", "grid", "empty"]


def _weights(r: int, c: int, n_rows: int, n_cols: int) -> tuple[_Weight, _Weight, _Weight, _Weight]:
    """The (top, bottom, left, right) line weights of the cell at local position (r, c) in an
    (n_rows x n_cols) data body (r/c include the header row / label column at 0): thick around
    the whole table, medium on the line separating header+labels from data, thin inside."""

    def weight(is_thick: bool, is_medium: bool) -> _Weight:
        return "thick" if is_thick else "medium" if is_medium else "thin"

    return (
        weight(r == 0, r == 1),
        weight(r == n_rows, r == 0),
        weight(c == 0, c == 1),
        weight(c == n_cols, c == 0),
    )


def _border_for(r: int, c: int, n_rows: int, n_cols: int, color: str) -> Border:
    """The Border for the cell at local position (r, c) — see :func:`_weights`."""
    top, bottom, left, right = _weights(r, c, n_rows, n_cols)
    return Border(
        top=Side(style=top, color=color),
        bottom=Side(style=bottom, color=color),
        left=Side(style=left, color=color),
        right=Side(style=right, color=color),
    )


def _paint_border(cell: Cell, r: int, c: int, n_rows: int, n_cols: int, style: TableStyle) -> None:
    cell.border = (
        _border_for(r, c, n_rows, n_cols, style.border_color) if style.border_color else Border()
    )


def _paint_header_like(cell: Cell, style: TableStyle) -> None:
    """Header row and label column share the same look."""
    cell.font = Font(
        name=style.header_font_name,
        size=style.header_font_size,
        bold=style.header_bold,
        color=style.header_font_color,
    )
    cell.fill = (
        PatternFill(start_color=style.header_fill, end_color=style.header_fill, fill_type="solid")
        if style.header_fill
        else PatternFill(fill_type=None)
    )


def _paint_data(cell: Cell, style: TableStyle, banded: bool) -> None:
    cell.font = Font(
        name=style.data_font_name, size=style.data_font_size, color=style.data_font_color
    )
    cell.fill = (
        PatternFill(start_color=style.band_fill, end_color=style.band_fill, fill_type="solid")
        if banded and style.band_fill
        else PatternFill(fill_type=None)
    )


def paint_table(ws: Worksheet, entry: TableEntry, row_limit: int = 0, col_limit: int = 0) -> None:
    """(Re)apply entry.style to the table's header row, label column, and data body.

    Runs the full region unless told which cells to leave alone — even for
    TableStyle.NONE, which is what lets switching to it actively clear previous
    formatting rather than just not adding new formatting on top of stale styling.
    Called after every create()/write() (so a shape change — a new row or column — is
    styled as part of the same pass that assembles it) and after every shift_right (so
    a table moved by a neighbour's growth keeps its look). ``row_limit``/``col_limit``
    (from :func:`stable_extent`) skip the cells that already have the look they'd be
    given — local rows below one and local columns below the other — which the caller
    has left unchanged.

    A table has only a few dozen distinct looks (header/label or data, banded or not, and
    the border variants along its edges), but tens of thousands of cells. Building fresh
    Font/Fill/Border objects for every cell — which openpyxl then hashes to find each one's
    slot in the workbook's style table — was over 90% of the time to create or write a
    table. So each distinct look is painted once the ordinary way, and its finished style
    record is copied onto every other cell that shares it. The cells are always fresh or
    cleared (see :func:`clear_region`), so the record is all a cell's style ever was.
    """
    style = entry.style
    looks: dict[tuple[Any, ...], Any] = {}

    def apply(cell: Any, record: Any) -> None:
        # A date, time or duration was given its number format when its value was written,
        # and text starting with "=" its quote prefix; nothing here paints either, so those
        # two fields stay the cell's own.
        own = cell._style  # cleared cells have none
        number_format = own.numFmtId if own else 0
        quote_prefix = own.quotePrefix if own else 0
        cell._style = copy(record)
        cell._style.numFmtId = number_format
        cell._style.quotePrefix = quote_prefix

    def weights(r: int, c: int) -> tuple[_Weight, _Weight, _Weight, _Weight] | None:
        return _weights(r, c, entry.n_rows, entry.n_cols) if style.border_color else None

    def head(cell: Any, r: int, c: int) -> None:
        key = ("head", weights(r, c))
        record = looks.get(key)
        if record is None:
            _paint_header_like(cell, style)
            _paint_border(cell, r, c, entry.n_rows, entry.n_cols, style)
            looks[key] = copy(cell._style)
        else:
            apply(cell, record)

    def data(cell: Any, r: int, c: int, banded: bool) -> None:
        key = ("data", banded, weights(r, c))
        record = looks.get(key)
        if record is None:
            _paint_data(cell, style, banded)
            _paint_border(cell, r, c, entry.n_rows, entry.n_cols, style)
            looks[key] = copy(cell._style)
        else:
            apply(cell, record)

    header_row = entry.anchor_row + 1
    for r in range(entry.n_rows + 1):
        row = header_row + r
        banded = bool(style.band_fill) and r % 2 == 0
        for c in range(col_limit if r < row_limit else 0, entry.width):
            cell = ws.cell(row=row, column=entry.anchor_col + c)
            if r == 0 or c == 0:
                head(cell, r, c)
            else:
                data(cell, r, c, banded)


# --------------------------------------------------------- marker verify/heal


def verify_or_locate(workbook: Workbook, entry: TableEntry) -> TableEntry:
    """Confirm *entry*'s recorded position still holds its marker.

    If it does not, scan the recorded sheet for the marker/name pair and
    return a corrected entry (position only — size is assumed unchanged).
    Raises TableNotFoundError if the marker cannot be found at all.
    """
    if entry.sheet not in workbook.sheetnames:
        raise TableNotFoundError(f"table {entry.name!r}: sheet {entry.sheet!r} no longer exists")
    ws = workbook[entry.sheet]
    if marker_matches(ws, entry):
        return entry

    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column):
            if ws.cell(r, c).value == MARKER and ws.cell(r, c + 1).value == entry.name:
                return replace(entry, anchor_row=r, anchor_col=c)
    raise TableNotFoundError(
        f"table {entry.name!r} was not found on sheet {entry.sheet!r} "
        "(its marker is missing — it may have been deleted or edited by hand)"
    )


# ------------------------------------------------------------------ lookups


def get_entry(entries: dict[str, TableEntry], name: str) -> TableEntry:
    if not isinstance(name, str):
        raise TypeError(f"name must be str, got {type(name).__name__}: {name!r}")
    try:
        return entries[name]
    except KeyError:
        raise TableNotFoundError(f"no table named {name!r}") from None


def check_not_exists(entries: dict[str, TableEntry], name: str) -> None:
    if name in entries:
        raise TableExistsError(f"table {name!r} already exists")


def drop_sheet_from_entries(entries: dict[str, TableEntry], sheet: str) -> dict[str, TableEntry]:
    """A copy of *entries* with every entry for *sheet* removed."""
    return {name: e for name, e in entries.items() if e.sheet != sheet}


def rename_sheet_in_entries(
    entries: dict[str, TableEntry], old: str, new: str
) -> dict[str, TableEntry]:
    """A copy of *entries* with every entry pointing at *old* moved to *new*."""
    return {name: (replace(e, sheet=new) if e.sheet == old else e) for name, e in entries.items()}


def sheet_kind(workbook: Workbook, entries: dict[str, TableEntry], sheet: str) -> SheetKind:
    """Whether *sheet* currently holds table data, grid data, or neither.

    ``"table"`` if any entry in *entries* points at it; otherwise ``"grid"``
    if it has any cell content; otherwise ``"empty"``.
    """
    if any(e.sheet == sheet for e in entries.values()):
        return "table"
    ws = workbook[sheet]
    if any(value is not None for row in all_rows(ws, values_only=True) for value in row):
        return "grid"
    return "empty"


# ------------------------------------------------------------------ layout


def find_placement(entries: dict[str, TableEntry], sheet: str) -> tuple[int, int]:
    """Where a brand-new table on *sheet* should start: row 1, one column past the rest."""
    same_sheet = [e for e in entries.values() if e.sheet == sheet]
    if not same_sheet:
        return 1, 1
    rightmost_used = max(e.anchor_col + e.width - 1 for e in same_sheet)
    return 1, rightmost_used + 2  # +1 empty separator, +1 to land past it


def tables_to_shift(entries: dict[str, TableEntry], entry: TableEntry) -> list[TableEntry]:
    """Every table on the same sheet sitting to the right of *entry*, left to right."""
    return sorted(
        (e for e in entries.values() if e.sheet == entry.sheet and e.anchor_col > entry.anchor_col),
        key=lambda e: e.anchor_col,
    )


def shift_right(
    workbook: Workbook, entries: dict[str, TableEntry], entry: TableEntry, growth: int
) -> None:
    """Move every table to the right of *entry* on its sheet further right by *growth* columns.

    Reads every affected region into memory before clearing or writing
    anything, so overlapping moves can never clobber each other. Repaints
    each moved table at its new position — its style doesn't change, but
    the styled cells themselves are new ones — and updates *entries* in
    place.
    """
    to_shift = tables_to_shift(entries, entry)
    if not to_shift:
        return
    ws = workbook[entry.sheet]
    contents = [
        (e, read_region(ws, e.anchor_row, e.anchor_col, e.height, e.width)) for e in to_shift
    ]
    for e, _ in contents:
        clear_region(ws, e.anchor_row, e.anchor_col, e.height, e.width)
    for e, grid in contents:
        new_col = e.anchor_col + growth
        write_region(ws, e.anchor_row, new_col, grid)
        moved = replace(e, anchor_col=new_col)
        entries[e.name] = moved
        paint_table(ws, moved)


# ------------------------------------------------- reading a table's state


def read_snapshot(ws: Any, entry: TableEntry) -> Snapshot:
    """The table at *entry*'s (already verified) position, as a file-independent snapshot."""
    return read_table(ws, entry)[0]


def read_table(ws: Any, entry: TableEntry) -> tuple[Snapshot, list[str]]:
    """The table at *entry*'s (already verified) position as ``(snapshot, formulas)``, the
    second being the coordinates of any formula cells found in it."""
    block, formulas = read_block(
        ws, entry.anchor_row + 1, entry.anchor_col, entry.n_rows + 1, entry.n_cols + 1
    )
    headers = [_to_label(h) for h in block[0][1:]]
    labels = [_to_label(row[0]) for row in block[1:]]
    for j, header in enumerate(headers):
        if header == "":
            cell = f"{get_column_letter(entry.anchor_col + 1 + j)}{entry.anchor_row + 1}"
            raise MalformedTableError(
                f"table {entry.name!r} on sheet {entry.sheet!r} can't be read: the column "
                f"header in cell {cell} is blank, and every column needs a header — fill it "
                "in (or delete the column) in Excel"
            )
    for i, label in enumerate(labels):
        if label == "":
            cell = f"{get_column_letter(entry.anchor_col)}{entry.anchor_row + 2 + i}"
            raise MalformedTableError(
                f"table {entry.name!r} on sheet {entry.sheet!r} can't be read: the row "
                f"label in cell {cell} is blank, and every row needs a label — fill it "
                "in (or delete the row) in Excel"
            )
    snapshot = Snapshot(
        row_labels=labels,
        column_headers=headers,
        data=[list(row[1:]) for row in block[1:]],
        corner=_to_label(block[0][0]),
        style=entry.style,
        column_types=list(entry.column_types),
    )
    return snapshot, formulas
